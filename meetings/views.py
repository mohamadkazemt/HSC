from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy
from django.contrib import messages
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from .models import Meeting
from .forms import MeetingForm
from .services import MeetingService
from dashboard.models import Notification
from django.http import HttpResponse
import csv
from django.db.models import Q
from django.db import transaction
from django.conf import settings
import logging
# from jalalidate import JalaliDate as jdate  # امتحان import از jalalidate (بدون آندرلاین)
import jdatetime
from django.utils import timezone
from .sms_utils import send_meeting_cancelled_sms, send_meeting_reminder_sms, send_meeting_updated_sms, send_meeting_deleted_sms, send_meeting_created_sms
from .notifications import (
    notify_meeting_updated,
    notify_meeting_cancelled,
    notify_meeting_deleted,
)
from permissions.utils import permission_required, check_permission
from functools import wraps
from datetime import date


logger = logging.getLogger(__name__)
logger.debug("Logging system is initialized in meetings/views.py")


def _has_named_access(user, view_name):
    if user.is_superuser:
        return True
    permissions = check_permission(user, view_name)
    if isinstance(permissions, dict):
        return any(permissions.values())
    return bool(permissions)


def _can_view_all_meetings(user):
    return _has_named_access(user, 'meeting_view_all')


def _can_approve_meetings(user):
    return _has_named_access(user, 'meeting_approve')


def _visible_meetings_for(user):
    queryset = Meeting.objects.all()
    if _can_view_all_meetings(user):
        return queryset
    return queryset.filter(
        Q(creator=user) |
        Q(approval_status='approved', participants=user)
    ).distinct()


def _can_view_meeting(user, meeting):
    return (
        _can_view_all_meetings(user)
        or meeting.creator_id == user.id
        or (
            meeting.approval_status == 'approved'
            and meeting.participants.filter(pk=user.pk).exists()
        )
    )


def _meeting_recipient_mobiles(meeting):
    mobiles = {
        participant.userprofile.mobile
        for participant in meeting.participants.select_related('userprofile')
        if hasattr(participant, 'userprofile') and participant.userprofile.mobile
    }
    if meeting.manual_numbers:
        mobiles.update(number.strip() for number in meeting.manual_numbers.splitlines() if number.strip())
    if meeting.notify_transport_coordinator:
        mobiles.update(
            coordinator.userprofile.mobile
            for coordinator in MeetingService.get_transport_coordinators().select_related('userprofile')
            if hasattr(coordinator, 'userprofile') and coordinator.userprofile.mobile
        )
    return sorted(mobiles)


def _queue_meeting_sms(meeting, action, *, reason=''):
    from .tasks import send_meeting_event_sms_async

    payload = {
        'action': action,
        'mobiles': _meeting_recipient_mobiles(meeting),
        'meeting_id': meeting.pk,
        'title': meeting.title,
        'date': meeting.date.isoformat(),
        'start_time': meeting.start_time.isoformat(),
        'reason': reason,
    }

    def enqueue():
        try:
            send_meeting_event_sms_async.delay(**payload)
        except Exception:
            logger.exception("Could not enqueue %s SMS for meeting %s", action, meeting.pk)

    transaction.on_commit(enqueue)

class MeetingListView(LoginRequiredMixin, ListView):
    model = Meeting
    template_name = 'meetings/meeting_list.html'
    context_object_name = 'meetings'
    ordering = ['-date', '-start_time']
    paginate_by = 10

    def dispatch(self, request, *args, **kwargs):
        if not check_permission(request.user, "meeting_list"):
            messages.error(request, "شما دسترسی لازم برای مشاهده لیست جلسات را ندارید.")
            return redirect('home')
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        queryset = super().get_queryset()
        
        # اعمال فیلترها
        q = self.request.GET.get('q')
        start_date = self.request.GET.get('start_date')
        end_date = self.request.GET.get('end_date')
        location = self.request.GET.get('location')
        status = self.request.GET.get('status')

        if q:
            queryset = queryset.filter(
                Q(title__icontains=q) |
                Q(description__icontains=q) |
                Q(location__icontains=q) |
                Q(participants__first_name__icontains=q) |
                Q(participants__last_name__icontains=q)
            ).distinct()

        # تبدیل تاریخ شمسی به میلادی و اعمال فیلتر
        try:
            if start_date:
                print(f"تاریخ شروع دریافتی: {start_date}")
                # تبدیل تاریخ شمسی به میلادی
                year, month, day = map(int, start_date.split('/'))
                j_date = jdatetime.date(year, month, day)
                g_date = j_date.togregorian()
                print(f"تاریخ شروع میلادی: {g_date}")
                queryset = queryset.filter(date__gte=g_date)
        except (ValueError, TypeError) as e:
            print(f"خطا در تبدیل تاریخ شروع: {e}")

        try:
            if end_date:
                print(f"تاریخ پایان دریافتی: {end_date}")
                # تبدیل تاریخ شمسی به میلادی
                year, month, day = map(int, end_date.split('/'))
                j_date = jdatetime.date(year, month, day)
                g_date = j_date.togregorian()
                print(f"تاریخ پایان میلادی: {g_date}")
                queryset = queryset.filter(date__lte=g_date)
        except (ValueError, TypeError) as e:
            print(f"خطا در تبدیل تاریخ پایان: {e}")
        
        if location:
            queryset = queryset.filter(location__icontains=location)
        
        if status:
            queryset = queryset.filter(status=status)

        queryset = queryset.filter(
            pk__in=_visible_meetings_for(self.request.user).values('pk')
        )

        return queryset.order_by('-date', '-start_time')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # اضافه کردن آمار به context
        context['total_meetings'] = self.get_queryset().count()
        context['total_participants'] = sum(meeting.participants.count() for meeting in self.get_queryset())
        if context['total_meetings'] > 0:
            context['avg_participants'] = round(context['total_participants'] / context['total_meetings'], 2)
        else:
            context['avg_participants'] = 0
        return context

class MeetingDetailView(LoginRequiredMixin, UserPassesTestMixin, DetailView):
    model = Meeting
    template_name = 'meetings/meeting_detail.html'
    context_object_name = 'meeting'

    def dispatch(self, request, *args, **kwargs):
        if not check_permission(request.user, "meeting_detail"):
            messages.error(request, "شما دسترسی لازم برای مشاهده جزئیات جلسه را ندارید.")
            return redirect('meetings:meeting_list')
        return super().dispatch(request, *args, **kwargs)

    def test_func(self):
        meeting = self.get_object()
        return _can_view_meeting(self.request.user, meeting)

class MeetingCreateView(LoginRequiredMixin, UserPassesTestMixin, CreateView):
    model = Meeting
    form_class = MeetingForm
    template_name = 'meetings/meeting_form.html'
    success_url = reverse_lazy('meetings:meeting_list')

    def dispatch(self, request, *args, **kwargs):
        if not _has_named_access(request.user, "meeting_create"):
            messages.error(request, "شما دسترسی لازم برای ایجاد جلسه را ندارید.")
            return redirect('meetings:meeting_list')
        return super().dispatch(request, *args, **kwargs)

    def test_func(self):
        return _has_named_access(self.request.user, 'meeting_create')

    def form_valid(self, form):
        logger.info("Creating meeting from standard form")
        try:
            meeting = MeetingService.create_meeting(
                title=form.cleaned_data['title'],
                date=form.cleaned_data['date'],
                start_time=form.cleaned_data['start_time'],
                end_time=form.cleaned_data['end_time'],
                creator=self.request.user,
                participants=form.cleaned_data['participants'],
                manual_numbers=form.cleaned_data.get('manual_numbers', ''),
                notify_transport_coordinator=form.cleaned_data.get('notify_transport_coordinator', False),
                location=form.cleaned_data.get('location', ''),
                description=form.cleaned_data.get('description', ''),
            )
            logger.info("Meeting %s created successfully", meeting.id)

            messages.success(self.request, 'جلسه ثبت شد و در انتظار تأیید مدیر است.')
            return redirect('meetings:meeting_list')
        except Exception as e:
            logger.exception("Error creating meeting")
            messages.error(self.request, f'خطا در ایجاد جلسه: {str(e)}')
            return self.form_invalid(form)

    def form_invalid(self, form):
        logger.warning("Form is invalid")
        logger.warning(f"Form errors: {form.errors}")
        logger.warning(f"Form errors detail: {form.errors.as_data()}")
        
        # نمایش خطاهای عمومی
        if form.non_field_errors():
            for error in form.non_field_errors():
                messages.error(self.request, error)
        
        # نمایش خطاهای فیلدها
        for field, errors in form.errors.items():
            if field != '__all__':  # خطاهای عمومی را نادیده می‌گیریم
                for error in errors:
                    messages.error(self.request, f"{form.fields[field].label}: {error}")
        
        return super().form_invalid(form)

class MeetingUpdateView(LoginRequiredMixin, UserPassesTestMixin, UpdateView):
    model = Meeting
    form_class = MeetingForm
    template_name = 'meetings/meeting_form.html'
    success_url = reverse_lazy('meetings:meeting_list')

    def dispatch(self, request, *args, **kwargs):
        if not check_permission(request.user, "meeting_edit"):
            messages.error(request, "شما دسترسی لازم برای ویرایش جلسه را ندارید.")
            return redirect('meetings:meeting_list')
        return super().dispatch(request, *args, **kwargs)

    def test_func(self):
        return self.request.user.has_perm('meetings.change_meeting')

    def form_valid(self, form):
        try:
            with transaction.atomic():
                meeting = form.save()
                if meeting.approval_status == 'approved':
                    notify_meeting_updated(meeting, actor=self.request.user)
                    _queue_meeting_sms(meeting, 'updated')
            messages.success(self.request, 'جلسه با موفقیت بروزرسانی شد.')
            return redirect(self.success_url)
        except Exception:
            logger.exception("Error updating meeting %s", self.object.pk)
            messages.error(self.request, 'خطا در بروزرسانی جلسه. لطفاً دوباره تلاش کنید.')
            return self.form_invalid(form)


class MeetingDeleteView(LoginRequiredMixin, UserPassesTestMixin, DeleteView):
    model = Meeting
    template_name = 'meetings/meeting_confirm_delete.html'
    success_url = reverse_lazy('meetings:meeting_list')

    def dispatch(self, request, *args, **kwargs):
        if not check_permission(request.user, "meeting_delete"):
            messages.error(request, "شما دسترسی لازم برای حذف جلسه را ندارید.")
            return redirect('meetings:meeting_list')
        return super().dispatch(request, *args, **kwargs)

    def test_func(self):
        return self.request.user.has_perm('meetings.delete_meeting')

    def form_valid(self, form):
        meeting = self.object
        with transaction.atomic():
            if meeting.approval_status == 'approved':
                notify_meeting_deleted(meeting, actor=self.request.user)
                _queue_meeting_sms(meeting, 'deleted')
            response = super().form_valid(form)
        messages.success(self.request, 'جلسه با موفقیت حذف شد.')
        return response
@permission_required("meeting_report")
@login_required
def meeting_report(request):
    meetings = Meeting.objects.all().order_by('-date', '-start_time')
    
    # محاسبه آمار
    total_meetings = meetings.count()
    total_participants = sum(meeting.participants.count() for meeting in meetings)
    avg_participants = total_participants / total_meetings if total_meetings > 0 else 0
    
    # گروه‌بندی جلسات بر اساس تاریخ
    meetings_by_date = {}
    for meeting in meetings:
        date = meeting.date
        if date not in meetings_by_date:
            meetings_by_date[date] = []
        meetings_by_date[date].append(meeting)
    
    context = {
        'meetings': meetings,
        'total_meetings': total_meetings,
        'total_participants': total_participants,
        'avg_participants': round(avg_participants, 2),
        'meetings_by_date': meetings_by_date
    }
    
    return render(request, 'meetings/meeting_report.html', context)

@permission_required("meeting_notification")
@login_required
def mark_notification_read(request, notification_id):
    notification = get_object_or_404(Notification, id=notification_id, user=request.user)
    notification.is_read = True
    notification.save()
    return redirect('meetings:meeting_list')

@permission_required("meeting_cancel")
@login_required
def cancel_meeting(request, pk):
    meeting = get_object_or_404(Meeting, pk=pk)

    if request.method == 'POST':
        reason = request.POST.get('cancellation_reason', '').strip()
        if not reason:
            messages.error(request, 'لطفاً دلیل لغو جلسه را وارد کنید.')
            return redirect('meetings:meeting_detail', pk=pk)

        try:
            with transaction.atomic():
                meeting.status = 'cancelled'
                meeting.cancellation_reason = reason
                meeting.cancelled_by = request.user
                meeting.cancelled_at = timezone.now()
                meeting.save(update_fields=[
                    'status',
                    'cancellation_reason',
                    'cancelled_by',
                    'cancelled_at',
                    'updated_at',
                ])
                if meeting.approval_status == 'approved':
                    notify_meeting_cancelled(meeting, reason=reason, actor=request.user)
                    _queue_meeting_sms(meeting, 'cancelled', reason=reason)
            messages.success(request, 'جلسه با موفقیت لغو شد.')
            return redirect('meetings:meeting_list')
        except Exception:
            logger.exception("Error cancelling meeting %s", meeting.pk)
            messages.error(request, 'خطا در لغو جلسه. لطفاً دوباره تلاش کنید.')
            return redirect('meetings:meeting_detail', pk=pk)

    return render(request, 'meetings/cancel_meeting.html', {
        'meeting': meeting,
        'title': 'لغو جلسه',
    })

@permission_required("meeting_delete")
@login_required
def delete_meeting(request, pk):
    meeting = get_object_or_404(Meeting, pk=pk)
    
    if request.method == 'POST':
        if meeting.approval_status != 'approved':
            meeting.delete()
            messages.success(request, 'جلسه حذف شد.')
            return redirect('meetings:meeting_list')
        try:
            print("\n=== شروع حذف جلسه ===")
            print(f"شناسه جلسه: {meeting.id}")
            print(f"عنوان جلسه: {meeting.title}")
            
            # ارسال پیامک حذف به شرکت‌کنندگان
            for participant in meeting.participants.all():
                if hasattr(participant, 'userprofile') and participant.userprofile.mobile:
                    print(f"ارسال پیامک حذف به شماره {participant.userprofile.mobile}")
                    send_meeting_deleted_sms(
                        participant.userprofile.mobile,
                        meeting.id,
                        meeting.title,
                        meeting.date,
                        meeting.start_time
                    )
            
            # ارسال پیامک حذف به شماره‌های دستی
            if meeting.manual_numbers:
                numbers = meeting.manual_numbers.split('\n')
                for number in numbers:
                    if number.strip():
                        print(f"ارسال پیامک حذف به شماره دستی {number.strip()}")
                        send_meeting_deleted_sms(
                            number.strip(),
                            meeting.id,
                            meeting.title,
                            meeting.date,
                            meeting.start_time
                        )
            
            # ارسال پیامک حذف به هماهنگ‌کننده حمل و نقل
            if meeting.notify_transport_coordinator:
                coordinators = MeetingService.get_transport_coordinators()
                for coordinator in coordinators:
                    if hasattr(coordinator, 'userprofile') and coordinator.userprofile.mobile:
                        print(f"ارسال پیامک حذف به هماهنگ‌کننده حمل و نقل {coordinator.userprofile.mobile}")
                        send_meeting_deleted_sms(
                            coordinator.userprofile.mobile,
                            meeting.id,
                            meeting.title,
                            meeting.date,
                            meeting.start_time
                        )
            
            notify_meeting_deleted(meeting, actor=request.user)

            meeting.delete()
            print("جلسه با موفقیت حذف شد")
            messages.success(request, 'جلسه با موفقیت حذف شد.')
            return redirect('meetings:meeting_list')
            
        except Exception as e:
            print(f"خطا در حذف جلسه: {str(e)}")
            import traceback
            print(f"جزئیات خطا: {traceback.format_exc()}")
            messages.error(request, 'خطا در حذف جلسه. لطفاً دوباره تلاش کنید.')
            return redirect('meetings:meeting_detail', pk=pk)
    
    return render(request, 'meetings/meeting_confirm_delete.html', {'meeting': meeting})

@permission_required("meeting_export")
@login_required
def meeting_export(request):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from io import BytesIO
    
    # ایجاد یک فایل اکسل جدید
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "لیست جلسات"
    
    # تنظیم عنوان‌های ستون‌ها
    headers = ['عنوان', 'تاریخ', 'زمان شروع', 'زمان پایان', 'مکان', 'شرکت‌کنندگان', 'وضعیت', 'توضیحات']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal='center')
    
    # دریافت جلسات و مرتب‌سازی
    meetings = Meeting.objects.all().order_by('-date', '-start_time')
    
    # پر کردن داده‌ها
    for row, meeting in enumerate(meetings, 2):
        # تبدیل تاریخ میلادی به شمسی
        j_date = jdatetime.date.fromgregorian(date=meeting.date)
        formatted_date = j_date.strftime("%Y/%m/%d")
        
        # وضعیت جلسه به فارسی
        status_map = {
            'scheduled': 'برنامه‌ریزی شده',
            'cancelled': 'لغو شده',
            'completed': 'برگزار شده'
        }
        
        # لیست شرکت‌کنندگان
        participants = ", ".join([f"{p.get_full_name()}" for p in meeting.participants.all()])
        
        # مقادیر ستون‌ها
        values = [
            meeting.title,
            formatted_date,
            meeting.start_time.strftime("%H:%M"),
            meeting.end_time.strftime("%H:%M"),
            meeting.location,
            participants,
            status_map.get(meeting.status, meeting.status),
            meeting.description
        ]
        
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = value
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # تنظیم عرض ستون‌ها
    for column in ws.columns:
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width if adjusted_width < 50 else 50
    
    # ذخیره فایل
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="meetings.xlsx"'
    
    # ذخیره به بایت‌استریم
    wb.save(response)
    return response

@permission_required("meeting_calendar")
@login_required
def meeting_calendar(request):
    from django.contrib.auth.models import User

    can_create_meeting = _has_named_access(request.user, 'meeting_create')
    context = {
        'users': (
            User.objects.filter(is_active=True)
            .select_related('userprofile')
            .order_by('first_name', 'last_name', 'username')
            if can_create_meeting else User.objects.none()
        ),
        'can_create_meeting': can_create_meeting,
        'can_approve_meetings': _can_approve_meetings(request.user),
        'can_view_all_meetings': _can_view_all_meetings(request.user),
    }
    return render(request, 'meetings/meeting_calendar.html', context)


@login_required
def meeting_view_all_access(request):
    if not _can_view_all_meetings(request.user):
        messages.error(request, 'دسترسی مشاهده همه جلسات را ندارید.')
        return redirect('meetings:meeting_calendar')
    return redirect('meetings:meeting_calendar')


@login_required
def meeting_approval_queue(request):
    if not _can_approve_meetings(request.user):
        messages.error(request, 'دسترسی تأیید جلسات را ندارید.')
        return redirect('meetings:meeting_calendar')
    pending_meetings = Meeting.objects.filter(
        approval_status='pending'
    ).select_related('creator').prefetch_related('participants').order_by('date', 'start_time')
    return render(request, 'meetings/meeting_approval_queue.html', {
        'meetings': pending_meetings,
    })


@login_required
def approve_meeting(request, pk):
    if not _can_approve_meetings(request.user):
        messages.error(request, 'دسترسی تأیید جلسات را ندارید.')
        return redirect('meetings:meeting_calendar')
    if request.method != 'POST':
        return redirect('meetings:meeting_approve')
    try:
        MeetingService.approve_meeting(pk, request.user)
        messages.success(request, 'جلسه تأیید شد و اطلاع‌رسانی‌ها در صف ارسال قرار گرفت.')
    except (Meeting.DoesNotExist, ValueError) as exc:
        messages.error(request, str(exc))
    return redirect('meetings:meeting_approve')


@login_required
def meeting_events_json(request):
    """Return only meetings visible to the current calendar user."""
    from django.http import JsonResponse

    meetings = (
        _visible_meetings_for(request.user)
        .prefetch_related('participants')
        .order_by('date', 'start_time')
    )

    # FullCalendar sends an ISO date range where end is exclusive.
    try:
        if request.GET.get('start'):
            meetings = meetings.filter(date__gte=date.fromisoformat(request.GET['start'][:10]))
        if request.GET.get('end'):
            meetings = meetings.filter(date__lt=date.fromisoformat(request.GET['end'][:10]))
    except ValueError:
        return JsonResponse({'error': 'بازه تاریخ نامعتبر است.'}, status=400)

    events = []
    for meeting in meetings:
        if meeting.approval_status == 'pending':
            color = '#d97706'
        elif meeting.approval_status == 'rejected':
            color = '#6b7280'
        else:
            color = {
                'scheduled': '#3b82f6',
                'cancelled': '#ef4444',
                'completed': '#10b981',
            }.get(meeting.status, '#6b7280')
        events.append({
            'id': meeting.pk,
            'title': meeting.title,
            'start': f"{meeting.date}T{meeting.start_time}",
            'end': f"{meeting.date}T{meeting.end_time}",
            'color': color,
            'location': meeting.location or '',
            'status': meeting.status,
            'status_display': meeting.get_status_display(),
            'approval_status': meeting.approval_status,
            'approval_status_display': meeting.get_approval_status_display(),
            'url': f'/meetings/{meeting.pk}/',
        })
    return JsonResponse(events, safe=False)

@login_required
def meeting_create_ajax(request):
    """ویو AJAX برای ایجاد جلسه"""
    from django.http import JsonResponse
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Method not allowed'}, status=405)
    
    if not _has_named_access(request.user, "meeting_create"):
        return JsonResponse({'success': False, 'error': 'دسترسی لازم برای ایجاد جلسه را ندارید'}, status=403)
    
    form = MeetingForm(request.POST)
    if form.is_valid():
        try:
            meeting = MeetingService.create_meeting(
                title=form.cleaned_data['title'],
                date=form.cleaned_data['date'],
                start_time=form.cleaned_data['start_time'],
                end_time=form.cleaned_data['end_time'],
                creator=request.user,
                participants=form.cleaned_data['participants'],
                manual_numbers=form.cleaned_data.get('manual_numbers', ''),
                notify_transport_coordinator=form.cleaned_data.get('notify_transport_coordinator', False),
                location=form.cleaned_data.get('location', ''),
                description=form.cleaned_data.get('description', '')
            )
            return JsonResponse({
                'success': True,
                'message': 'جلسه ثبت شد و در انتظار تأیید مدیر است.',
                'meeting_id': meeting.id
            })
        except Exception as e:
            logger.error(f"Error creating meeting via AJAX: {str(e)}")
            return JsonResponse({'success': False, 'error': f'خطا در ایجاد جلسه: {str(e)}'}, status=500)
    else:
        errors = {}
        for field, field_errors in form.errors.items():
            errors[field] = field_errors
        return JsonResponse({'success': False, 'errors': errors}, status=400)

@login_required
def meeting_update_ajax(request, pk):
    """AJAX endpoint for updating a meeting."""
    from django.http import JsonResponse

    meeting = get_object_or_404(Meeting, pk=pk)
    if (
        not check_permission(request.user, "meeting_edit")
        or not request.user.has_perm('meetings.change_meeting')
    ):
        return JsonResponse({'success': False, 'error': 'دسترسی لازم برای ویرایش جلسه را ندارید'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'Method not allowed'}, status=405)

    form = MeetingForm(request.POST, instance=meeting)
    if not form.is_valid():
        errors = {field: field_errors for field, field_errors in form.errors.items()}
        return JsonResponse({'success': False, 'errors': errors}, status=400)

    try:
        with transaction.atomic():
            meeting = form.save()
            if meeting.approval_status == 'approved':
                notify_meeting_updated(meeting, actor=request.user)
                _queue_meeting_sms(meeting, 'updated')
        return JsonResponse({
            'success': True,
            'message': 'جلسه با موفقیت بروزرسانی شد.',
            'meeting_id': meeting.id,
        })
    except Exception as exc:
        logger.exception("Error updating meeting %s via AJAX", meeting.pk)
        return JsonResponse({'success': False, 'error': f'خطا در بروزرسانی جلسه: {exc}'}, status=500)

@login_required
def meeting_detail_ajax(request, pk):
    """ویو AJAX برای دریافت جزئیات جلسه"""
    from django.http import JsonResponse
    import jdatetime
    
    meeting = get_object_or_404(
        Meeting.objects.select_related('creator').prefetch_related('participants__userprofile'),
        pk=pk,
    )
    if not check_permission(request.user, "meeting_detail"):
        return JsonResponse({'error': 'دسترسی لازم برای مشاهده جلسه را ندارید'}, status=403)
    if not _can_view_meeting(request.user, meeting):
        return JsonResponse({'error': 'دسترسی لازم برای مشاهده این جلسه را ندارید'}, status=403)
    
    # تبدیل تاریخ به شمسی
    j_date = jdatetime.date.fromgregorian(date=meeting.date)
    
    # تبدیل زمان
    start_time_str = meeting.start_time.strftime('%H:%M')
    end_time_str = meeting.end_time.strftime('%H:%M')
    
    # لیست شرکت‌کنندگان
    participants = []
    for participant in meeting.participants.all():
        try:
            profile = participant.userprofile
            participants.append({
                'id': participant.id,
                'name': participant.get_full_name() or participant.username,
                'personnel_code': profile.personnel_code if hasattr(profile, 'personnel_code') else ''
            })
        except:
            participants.append({
                'id': participant.id,
                'name': participant.get_full_name() or participant.username,
                'personnel_code': ''
            })
    
    return JsonResponse({
        'id': meeting.id,
        'title': meeting.title,
        'date': j_date.strftime('%Y/%m/%d'),
        'date_gregorian': meeting.date.strftime('%Y-%m-%d'),
        'start_time': start_time_str,
        'end_time': end_time_str,
        'location': meeting.location or '',
        'description': meeting.description or '',
        'status': meeting.status,
        'status_display': meeting.get_status_display(),
        'approval_status': meeting.approval_status,
        'approval_status_display': meeting.get_approval_status_display(),
        'participants': participants,
        'manual_numbers': meeting.manual_numbers or '',
        'notify_transport_coordinator': meeting.notify_transport_coordinator,
        'creator': meeting.creator.get_full_name() or meeting.creator.username,
        'created_at': meeting.created_at.strftime('%Y-%m-%d %H:%M'),
        'can_edit': (
            check_permission(request.user, 'meeting_edit')
            and request.user.has_perm('meetings.change_meeting')
        ),
    })

@login_required
def meeting_search(request):
    query = request.GET.get('q', '')
    meetings = Meeting.objects.all()
    
    if query:
        meetings = meetings.filter(
            Q(title__icontains=query) |
            Q(description__icontains=query) |
            Q(location__icontains=query) |
            Q(participants__first_name__icontains=query) |
            Q(participants__last_name__icontains=query)
        ).distinct()
    
    meetings = meetings.order_by('-date', '-start_time')
    return render(request, 'meetings/meeting_list.html', {'meetings': meetings, 'query': query})

@login_required
def meeting_filter(request):
    meetings = Meeting.objects.all()
    
    # فیلتر بر اساس تاریخ
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    if start_date:
        meetings = meetings.filter(date__gte=start_date)
    if end_date:
        meetings = meetings.filter(date__lte=end_date)
    
    # فیلتر بر اساس مکان
    location = request.GET.get('location')
    if location:
        meetings = meetings.filter(location__icontains=location)
    
    # فیلتر بر اساس شرکت‌کننده
    participant = request.GET.get('participant')
    if participant:
        meetings = meetings.filter(participants__id=participant)
    
    meetings = meetings.order_by('-date', '-start_time')
    return render(request, 'meetings/meeting_list.html', {'meetings': meetings})
