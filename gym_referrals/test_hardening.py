from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from unittest.mock import MagicMock, patch

from django.contrib.auth.models import Permission, User
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.test import RequestFactory
from django.core.exceptions import PermissionDenied
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import UserProfile
from .billing import BillingError, GymInvoiceService
from .models import Gym, GymContract, GymInvoice, GymInvoiceAuditLog, Referral
from .presentation import format_jalali
from .services import ReferralService
from . import views


class ProductionHardeningTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user('owner', first_name='Old', last_name='Name')
        self.profile = UserProfile.objects.create(user=self.owner, personnel_code='P-1')
        self.staff = User.objects.create_user('staff', is_staff=True)
        self.gym = Gym.objects.create(name='باشگاه آماده', code='READY')
        today = timezone.localdate()
        self.contract = GymContract.objects.create(gym=self.gym, start_date=today.replace(day=1), end_date=today + timedelta(days=90), price_per_referral=Decimal('100.00'))
        self.referral = ReferralService.create_referral(requester=self.owner, beneficiary_type='EMPLOYEE', beneficiary_id=self.profile.pk, gym_id=self.gym.pk, source='WEB')[0]
        self.year, self.month = today.year, today.month

    def generate(self, actor=None):
        return GymInvoiceService.generate_invoice(contract_id=self.contract.pk, year=self.year, month=self.month, actor=actor or self.staff)

    def grant(self, codename):
        self.staff.user_permissions.add(Permission.objects.get(content_type__app_label='gym_referrals', codename=codename))
        self.staff = User.objects.get(pk=self.staff.pk)

    def test_zero_and_negative_tariff_are_rejected(self):
        for value in (Decimal('0.00'), Decimal('-1.00')):
            self.contract.price_per_referral = value
            self.contract.save(update_fields=('price_per_referral',))
            with self.assertRaises(BillingError) as caught:
                self.generate()
            self.assertEqual(caught.exception.code, 'INVALID_CONTRACT_TARIFF')

    def test_adjustment_requires_reason_and_is_audited(self):
        invoice = self.generate()
        with self.assertRaises(BillingError):
            GymInvoiceService.set_adjustment(invoice.pk, Decimal('10.00'), self.staff, '')
        GymInvoiceService.set_adjustment(invoice.pk, Decimal('10.00'), self.staff, 'اصلاح قرارداد')
        audit = GymInvoiceAuditLog.objects.get(invoice=invoice, action='ADJUSTED')
        self.assertEqual(audit.metadata['old_amount'], '0.00')
        self.assertEqual(audit.metadata['new_amount'], '10.00')

    def test_non_positive_total_cannot_finalize(self):
        invoice = self.generate()
        GymInvoiceService.set_adjustment(invoice.pk, Decimal('-100.00'), self.staff, 'صفر کردن')
        with self.assertRaises(BillingError) as caught:
            GymInvoiceService.transition(invoice.pk, GymInvoice.Status.FINALIZED, self.staff)
        self.assertEqual(caught.exception.code, 'INVOICE_TOTAL_NOT_POSITIVE')

    def test_staff_without_permission_cannot_generate(self):
        request = RequestFactory().post('/invoices/', {'contract': self.contract.pk, 'year': self.year, 'month': self.month})
        request.user = self.staff
        with self.assertRaises(PermissionDenied):
            views.invoice_list(request)

    def test_generate_permission_allows_generation(self):
        self.grant('view_gyminvoice'); self.grant('generate_gym_invoice')
        self.client.force_login(self.staff)
        response = self.client.post(reverse('gym_referrals:invoice_list'), {'contract': self.contract.pk, 'year': self.year, 'month': self.month})
        self.assertEqual(response.status_code, 302)

    def test_approve_requires_specific_permission(self):
        invoice = self.generate()
        GymInvoiceService.transition(invoice.pk, 'FINALIZED', self.staff)
        url = reverse('gym_referrals:invoice_transition', args=(invoice.pk, 'APPROVED'))
        request = RequestFactory().post(url); request.user = self.staff
        with self.assertRaises(PermissionDenied):
            views.invoice_transition(request, invoice.pk, 'APPROVED')
        self.grant('approve_gym_invoice')
        self.client.force_login(self.staff)
        self.assertEqual(self.client.post(url).status_code, 302)

    def test_excel_permission_and_snapshot(self):
        invoice = self.generate()
        url = reverse('gym_referrals:invoice_export_excel', args=(invoice.pk,))
        request = RequestFactory().get(url); request.user = self.staff
        with self.assertRaises(PermissionDenied):
            views.invoice_export_excel(request, invoice.pk)
        self.grant('export_gym_invoice')
        self.client.force_login(self.staff)
        self.owner.first_name = 'New'; self.owner.last_name = 'Person'; self.owner.save()
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        values = [cell.value for row in workbook.active.iter_rows() for cell in row]
        self.assertIn('Old Name', values)
        self.assertNotIn('New Person', values)
        self.assertIn(Decimal('100.00'), values)

    def test_jalali_formatter(self):
        self.assertEqual(format_jalali(timezone.datetime(2024, 3, 20).date()), '1403/01/01')

    @patch('qrcode.make')
    def test_qr_uses_public_token_not_database_id(self, make):
        from permissions.models import UserPermission
        UserPermission.objects.create(user=self.owner, view_name='gym_referral_history', can_view=True)
        image = MagicMock(); make.return_value = image
        image.save.side_effect = lambda buffer, format: buffer.write(b'png')
        self.client.force_login(self.owner)
        response = self.client.get(reverse('gym_referrals:referral_qr', args=(self.referral.referral_number,)))
        self.assertEqual(response.status_code, 200)
        payload = make.call_args.args[0]
        self.assertIn(str(self.referral.public_token), payload)
        self.assertNotIn(f'/referrals/{self.referral.pk}/', payload)

    def test_readiness_command(self):
        self.contract.price_per_referral = Decimal('0.00'); self.contract.save(update_fields=('price_per_referral',))
        with self.assertRaises(CommandError):
            call_command('check_gym_billing', stdout=StringIO(), stderr=StringIO())
        self.contract.price_per_referral = Decimal('100.00'); self.contract.save(update_fields=('price_per_referral',))
        call_command('check_gym_billing', stdout=StringIO(), stderr=StringIO())
