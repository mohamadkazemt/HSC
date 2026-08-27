import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from asgiref.sync import async_to_sync
from django.contrib.auth.models import Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.core.exceptions import PermissionDenied
from django.test import TestCase
from django.test import RequestFactory
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import Dependent, UserProfile
from permissions.models import UserPermission
from permissions.utils import get_all_views_with_labels
from rubika_bot.models import RubikaUser
from rubika_bot.tasks import notify_gym_operators_new_referral
from .access import has_gym_admin_access
from .billing import BillingError, GymInvoiceService
from .models import Gym, GymContract, GymInvoice, GymOperator, Referral, ReferralUsage
from .rubika import RubikaGymOperatorAdapter
from .services import ReferralError, ReferralService, create_gym_account, is_referral_active, reset_gym_account_password


class ReferralServiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("ali", password="pass", first_name="علی", last_name="احمدی")
        self.profile = UserProfile.objects.create(user=self.user, personnel_code="100")
        self.dep_a = Dependent.objects.create(personnel=self.profile, first_name="سارا", last_name="احمدی", national_code="001", relationship="فرزند")
        self.dep_b = Dependent.objects.create(personnel=self.profile, first_name="رضا", last_name="احمدی", national_code="002", relationship="فرزند")
        self.gym_a = Gym.objects.create(name="باشگاه الف", code="A")
        self.gym_b = Gym.objects.create(name="باشگاه ب", code="B")
        today = timezone.localdate()
        for gym in (self.gym_a, self.gym_b):
            GymContract.objects.create(gym=gym, start_date=today - timedelta(days=1), end_date=today + timedelta(days=365))

    def create(self, beneficiary_type=None, beneficiary_id=None, gym=None, **kwargs):
        return ReferralService.create_referral(
            requester=self.user, beneficiary_type=beneficiary_type or Referral.BeneficiaryType.EMPLOYEE,
            beneficiary_id=beneficiary_id or self.profile.pk, gym_id=(gym or self.gym_a).pk,
            source=kwargs.pop("source", Referral.Source.WEB), **kwargs,
        )[0]

    def test_person_without_active_referral_can_create(self):
        referral = self.create()
        self.assertTrue(is_referral_active(referral))

    def test_active_referral_blocks_same_and_other_gym(self):
        self.create()
        for gym in (self.gym_a, self.gym_b):
            with self.assertRaises(ReferralError) as caught:
                self.create(gym=gym)
            self.assertEqual(caught.exception.code, "ACTIVE_REFERRAL_ALREADY_EXISTS")

    def test_beneficiaries_are_independent(self):
        self.create()
        self.create(Referral.BeneficiaryType.DEPENDENT, self.dep_a.pk, self.gym_b)
        self.create(Referral.BeneficiaryType.DEPENDENT, self.dep_b.pk, self.gym_a)
        self.assertEqual(Referral.objects.filter(status=Referral.Status.ACTIVE).count(), 3)

    def test_expired_and_cancelled_do_not_block(self):
        old = self.create(
            valid_from=timezone.localdate() - timedelta(days=10),
            valid_until=timezone.localdate() - timedelta(days=1),
        )
        new = self.create(gym=self.gym_b)
        old.refresh_from_db()
        self.assertEqual(old.status, Referral.Status.EXPIRED)
        ReferralService.cancel(new, self.user, "test")
        self.create(gym=self.gym_a)

    def test_active_legacy_blocks_digital(self):
        self.create(source=Referral.Source.PHYSICAL_LEGACY, legacy_letter_no="L-1")
        with self.assertRaises(ReferralError) as caught:
            self.create(gym=self.gym_b)
        self.assertEqual(caught.exception.code, "ACTIVE_REFERRAL_ALREADY_EXISTS")

    def test_idempotent_retry_returns_same_referral(self):
        first = self.create(idempotency_key="request-1")
        second = self.create(idempotency_key="request-1")
        self.assertEqual(first.pk, second.pk)

    def test_database_constraint_rejects_second_active(self):
        first = self.create()
        first.pk = None
        first.referral_number = "GYM-X-2"
        first.public_token = None
        with self.assertRaises(IntegrityError), transaction.atomic():
            first.save(force_insert=True)

    def test_gym_mismatch_redeem_is_rejected(self):
        referral = self.create()
        with self.assertRaises(ReferralError) as caught:
            ReferralService.redeem(referral.public_token, self.user, self.gym_b)
        self.assertEqual(caught.exception.code, "REFERRAL_GYM_MISMATCH")

    def test_other_users_dependent_is_forbidden(self):
        other = User.objects.create_user("other")
        UserProfile.objects.create(user=other)
        with self.assertRaises(ReferralError) as caught:
            ReferralService.create_referral(requester=other, beneficiary_type="DEPENDENT", beneficiary_id=self.dep_a.pk, gym_id=self.gym_a.pk, source="WEB")
        self.assertEqual(caught.exception.code, "FORBIDDEN")


class ReferralApiTests(ReferralServiceTests):
    def setUp(self):
        super().setUp()
        UserPermission.objects.create(user=self.user, view_name="gym_referral_create", can_view=True)
        UserPermission.objects.create(user=self.user, view_name="gym_referral_history", can_view=True)

    def test_create_api_conflict_contains_existing_referral(self):
        self.client.force_login(self.user)
        url = reverse("gym_referrals:api_create")
        payload = {"beneficiaryType": "EMPLOYEE", "beneficiaryId": self.profile.pk, "gymId": self.gym_a.pk}
        self.assertEqual(self.client.post(url, json.dumps(payload), content_type="application/json").status_code, 201)
        response = self.client.post(url, json.dumps(payload), content_type="application/json")
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["code"], "ACTIVE_REFERRAL_ALREADY_EXISTS")
        self.assertEqual(response.json()["existingReferral"]["gymName"], self.gym_a.name)

    def test_public_verify_uses_token(self):
        referral = self.create()
        response = self.client.get(reverse("gym_referrals:api_verify", args=(referral.public_token,)))
        self.assertEqual(response.status_code, 200)
        payload = response.json()["referral"]
        self.assertNotIn("employeeAccountId", payload)
        self.assertNotIn("beneficiaryName", payload)
        self.assertNotIn("publicToken", payload)
        referral.refresh_from_db()
        self.assertIsNone(referral.verified_at)
        self.assertFalse(referral.audit_logs.filter(action="VERIFIED").exists())

    def test_public_verify_reports_expired_without_writing(self):
        referral = self.create()
        Referral.objects.filter(pk=referral.pk).update(
            valid_from=timezone.localdate() - timedelta(days=10),
            valid_until=timezone.localdate() - timedelta(days=1),
        )
        response = self.client.get(reverse("gym_referrals:api_verify", args=(referral.public_token,)))
        payload = response.json()
        self.assertFalse(payload["valid"])
        self.assertEqual(payload["referral"]["status"], Referral.Status.EXPIRED)
        referral.refresh_from_db()
        self.assertEqual(referral.status, Referral.Status.ACTIVE)


class BillingReleaseTests(ReferralServiceTests):
    def setUp(self):
        super().setUp()
        self.contract = self.gym_a.contracts.get()
        self.contract.price_per_referral = Decimal("125000.00")
        self.contract.billing_policy = GymContract.BillingPolicy.ISSUED
        self.contract.save(update_fields=("price_per_referral", "billing_policy"))

    def invoice(self):
        referral = self.create()
        invoice = GymInvoiceService.generate_invoice(
            contract_id=self.contract.pk, year=referral.issue_date.year,
            month=referral.issue_date.month, actor=self.user,
        )
        return referral, invoice

    def test_invoice_lifecycle_rejects_skipped_transition(self):
        _, invoice = self.invoice()
        with self.assertRaises(BillingError):
            GymInvoiceService.transition(invoice.pk, GymInvoice.Status.PAID, self.user)
        invoice.refresh_from_db()
        self.assertEqual(invoice.status, GymInvoice.Status.DRAFT)

    def test_finalized_invoice_snapshots_do_not_follow_master_data(self):
        referral, invoice = self.invoice()
        item = invoice.items.get()
        GymInvoiceService.transition(invoice.pk, GymInvoice.Status.FINALIZED, self.user)
        self.contract.price_per_referral = Decimal("999999.00")
        self.contract.save(update_fields=("price_per_referral",))
        referral.beneficiary_full_name_snapshot = "نام جدید"
        referral.save(update_fields=("beneficiary_full_name_snapshot",))
        item.refresh_from_db()
        self.assertEqual(item.unit_price, Decimal("125000.00"))
        self.assertNotEqual(item.beneficiary_name_snapshot, "نام جدید")

    def test_duplicate_invoice_is_rejected_without_duplicate_items(self):
        _, invoice = self.invoice()
        with self.assertRaises(BillingError) as caught:
            GymInvoiceService.generate_invoice(
                contract_id=self.contract.pk, year=invoice.period_year,
                month=invoice.period_month, actor=self.user,
            )
        self.assertEqual(caught.exception.code, "INVOICE_ALREADY_EXISTS")
        self.assertEqual(GymInvoice.objects.count(), 1)
        self.assertEqual(invoice.items.count(), 1)

    def test_contract_boundaries_are_inclusive_for_issued_policy(self):
        today = timezone.localdate()
        self.contract.start_date = today
        self.contract.end_date = today
        self.contract.save(update_fields=("start_date", "end_date"))
        referral = self.create(valid_from=today, valid_until=today, issue_date=today)
        invoice = GymInvoiceService.generate_invoice(
            contract_id=self.contract.pk, year=today.year, month=today.month, actor=self.user
        )
        self.assertEqual(invoice.items.get().referral_id, referral.pk)

    def test_used_policy_uses_redeemed_timestamp(self):
        self.contract.billing_policy = GymContract.BillingPolicy.USED
        self.contract.save(update_fields=("billing_policy",))
        referral = self.create()
        ReferralService.redeem(referral.public_token, self.user, self.gym_a)
        invoice = GymInvoiceService.generate_invoice(
            contract_id=self.contract.pk, year=timezone.localdate().year,
            month=timezone.localdate().month, actor=self.user,
        )
        self.assertEqual(invoice.items.get().referral_usage_id, ReferralUsage.objects.get().pk)

    def test_excel_export_sanitizes_formula_prefixes(self):
        self.user.first_name = '=HYPERLINK("https://example.invalid")'
        self.user.last_name = ""
        self.user.save(update_fields=("first_name", "last_name"))
        _, invoice = self.invoice()
        self.user.user_permissions.add(Permission.objects.get(codename="export_gym_invoice"))
        self.client.force_login(self.user)
        response = self.client.get(reverse("gym_referrals:invoice_export_excel", args=(invoice.pk,)))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        value = workbook.active.cell(13, 2).value
        self.assertTrue(value.startswith("'="), value)

    def test_financial_transition_requires_explicit_permission(self):
        _, invoice = self.invoice()
        from .views import invoice_transition
        request = RequestFactory().post("/financial-transition/")
        request.user = self.user
        with self.assertRaises(PermissionDenied):
            invoice_transition(request, invoice.pk, GymInvoice.Status.FINALIZED)
        invoice.refresh_from_db()
        self.assertEqual(invoice.status, GymInvoice.Status.DRAFT)

    def test_large_invoice_generation_and_excel_export(self):
        today = timezone.localdate()
        rows = [
            Referral(
                referral_number=f"GYM-LARGE-{index:04d}", employee=self.user,
                beneficiary_type=Referral.BeneficiaryType.DEPENDENT,
                beneficiary_account_id=100000 + index,
                beneficiary_key=f"DEPENDENT:{100000 + index}",
                relation_snapshot="فرزند", employee_full_name_snapshot="پرسنل آزمون",
                beneficiary_full_name_snapshot=f"ذی‌نفع {index}", personnel_no_snapshot="100",
                gym=self.gym_a, issue_date=today, valid_from=today,
                valid_until=today + timedelta(days=30), status=Referral.Status.ACTIVE,
                source=Referral.Source.WEB, created_by=self.user,
            )
            for index in range(1000)
        ]
        Referral.objects.bulk_create(rows)
        invoice = GymInvoiceService.generate_invoice(
            contract_id=self.contract.pk, year=today.year, month=today.month, actor=self.user
        )
        self.assertEqual(invoice.item_count, 1000)
        self.user.user_permissions.add(Permission.objects.get(codename="export_gym_invoice"))
        self.client.force_login(self.user)
        response = self.client.get(reverse("gym_referrals:invoice_export_excel", args=(invoice.pk,)))
        self.assertEqual(response.status_code, 200)
        self.assertGreater(len(response.content), 10000)


class LegacyImportTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_superuser("admin", password="x", first_name="مدیر", last_name="سامانه")
        self.profile = UserProfile.objects.create(user=self.staff, personnel_code="900")
        self.dependents = [
            Dependent.objects.create(personnel=self.profile, first_name=f"فرزند{index}", last_name="سامانه",
                                     national_code=f"N-{index}", relationship="فرزند")
            for index in range(1, 7)
        ]
        self.gym = Gym.objects.create(name="باشگاه واردات", code="IMP")
        today = timezone.localdate()
        GymContract.objects.create(gym=self.gym, start_date=today - timedelta(days=1), end_date=today + timedelta(days=365))
        self.client.force_login(self.staff)

    def _csv(self, rows):
        header = "کد پرسنلی,کد ملی تحت تکفل,شناسه باشگاه,شماره معرفی‌نامه,تاریخ صدور,تاریخ شروع,تاریخ پایان\n"
        return SimpleUploadedFile("legacy.csv", (header + "\n".join(rows)).encode("utf-8-sig"), content_type="text/csv")

    def _row(self, index):
        today = timezone.localdate()
        dependent = self.dependents[(index - 1) % len(self.dependents)]
        fields = [self.profile.personnel_code, dependent.national_code, str(self.gym.pk),
                  f"L-{index}", today.isoformat(), (today - timedelta(days=30)).isoformat(), today.isoformat()]
        return ",".join(fields)

    def _employee_row(self, index):
        today = timezone.localdate()
        fields = [self.profile.personnel_code, "", str(self.gym.pk),
                  f"E-{index}", today.isoformat(), (today - timedelta(days=30)).isoformat(), today.isoformat()]
        return ",".join(fields)

    def test_oversized_file_is_rejected(self):
        upload = SimpleUploadedFile("legacy.csv", b"a" * (2 * 1024 * 1024 + 1), content_type="text/csv")
        response = self.client.post(reverse("gym_referrals:legacy_import"), {"file": upload})
        self.assertContains(response, "حجم فایل")
        self.assertEqual(Referral.objects.count(), 0)

    def test_import_stops_at_row_cap(self):
        rows = [self._row(index) for index in range(1, 6)]
        with patch("gym_referrals.views.MAX_LEGACY_IMPORT_ROWS", 3):
            response = self.client.post(reverse("gym_referrals:legacy_import"), {"file": self._csv(rows)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Referral.objects.count(), 3)
        summary = response.context["summary"]
        self.assertEqual(summary["total"], 3)
        self.assertTrue(summary["truncated"])

    def test_duplicate_letter_is_idempotent_and_bad_row_is_reported(self):
        rows = [self._row(1)]
        first = self.client.post(reverse("gym_referrals:legacy_import"), {"file": self._csv(rows)})
        summary = first.context["summary"]
        self.assertEqual(summary["successful"], 1)
        self.assertFalse(summary["truncated"])
        second = self.client.post(reverse("gym_referrals:legacy_import"), {"file": self._csv(rows)})
        self.assertEqual(second.context["summary"]["successful"], 1)
        self.assertIn("تکراری", second.context["report"][0]["detail"])
        self.assertEqual(Referral.objects.count(), 1)

    def test_employee_row_without_dependent_code_creates_employee_referral(self):
        response = self.client.post(reverse("gym_referrals:legacy_import"), {"file": self._csv([self._employee_row(1)])})
        summary = response.context["summary"]
        self.assertEqual(summary["successful"], 1)
        referral = Referral.objects.get()
        self.assertEqual(referral.beneficiary_type, "EMPLOYEE")
        self.assertEqual(referral.employee, self.staff)


class PermissionsIntegrationTests(TestCase):
    """Access to gym pages is controlled through the central permissions app."""

    def setUp(self):
        self.employee = User.objects.create_user("worker", password="pass")
        self.profile = UserProfile.objects.create(user=self.employee, personnel_code="700")
        self.gym = Gym.objects.create(name="باشگاه دسترسی", code="ACC")
        today = timezone.localdate()
        GymContract.objects.create(gym=self.gym, start_date=today - timedelta(days=1), end_date=today + timedelta(days=365))

    def grant(self, view_name, user=None):
        UserPermission.objects.create(user=user or self.employee, view_name=view_name, can_view=True)

    def test_gym_views_are_registered_in_permissions_ui(self):
        registered = {view["name"] for view in get_all_views_with_labels() if view.get("app_label") == "gym_referrals"}
        for expected in ("gym_referral_create", "gym_referral_history", "gym_referral_management",
                         "gym_referral_usage", "gym_manage", "gym_contract_manage",
                         "gym_operator_manage", "gym_invoice_view"):
            self.assertIn(expected, registered)

    def test_regular_user_without_grant_is_denied(self):
        self.client.force_login(self.employee)
        for url in (reverse("gym_referrals:create"), reverse("gym_referrals:history")):
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 403)

    def test_dynamic_grant_allows_only_granted_views(self):
        self.grant("gym_referral_create")
        self.client.force_login(self.employee)
        self.assertEqual(self.client.get(reverse("gym_referrals:create")).status_code, 200)
        self.assertEqual(self.client.get(reverse("gym_referrals:history")).status_code, 403)

    def test_part_level_style_grant_via_user_permission_api(self):
        self.grant("gym_referral_create")
        self.client.force_login(self.employee)
        response = self.client.post(
            reverse("gym_referrals:api_create"),
            data=json.dumps({"beneficiaryType": "EMPLOYEE", "beneficiaryId": self.profile.pk, "gymId": self.gym.pk}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)

    def test_superuser_bypasses_dynamic_gate(self):
        admin = User.objects.create_superuser("root", password="x")
        self.assertTrue(has_gym_admin_access(admin))
        self.client.force_login(admin)
        self.assertEqual(self.client.get(reverse("gym_referrals:create")).status_code, 200)

    def test_model_permission_holder_passes_dual_path(self):
        staff = User.objects.create_user("staffy", password="pass", is_staff=True)
        staff.user_permissions.add(Permission.objects.get(codename="add_gymcontract"))
        staff = User.objects.get(pk=staff.pk)
        self.assertTrue(has_gym_admin_access(staff))
        self.assertFalse(UserPermission.objects.exists())

    def test_public_verify_stays_anonymous(self):
        employee = self.employee
        profile = self.profile
        referral = ReferralService.create_referral(
            requester=employee, beneficiary_type="EMPLOYEE", beneficiary_id=profile.pk,
            gym_id=self.gym.pk, source="WEB",
        )[0]
        self.client.logout()
        response = self.client.get(reverse("gym_referrals:api_verify", args=(referral.public_token,)))
        self.assertEqual(response.status_code, 200)


class GymAccountAndPortalTests(TestCase):
    """Gym login accounts are auto-created and power the gym portal."""

    def setUp(self):
        self.employee = User.objects.create_user("worker", password="pass")
        self.profile = UserProfile.objects.create(user=self.employee, personnel_code="900")
        self.gym = Gym.objects.create(name="باشگاه حساب", code="ACCT")
        today = timezone.localdate()
        GymContract.objects.create(gym=self.gym, start_date=today - timedelta(days=1), end_date=today + timedelta(days=365))

    def make_referral(self):
        return ReferralService.create_referral(
            requester=self.employee, beneficiary_type="EMPLOYEE", beneficiary_id=self.profile.pk,
            gym_id=self.gym.pk, source="WEB",
        )[0]

    def test_create_gym_account_links_user_and_operator(self):
        user, raw = create_gym_account(self.gym)
        self.gym.refresh_from_db()
        self.assertTrue(user.check_password(raw))
        self.assertTrue(self.gym.account_user_id, user.pk)
        self.assertTrue(self.gym.operators.filter(user=user, is_active=True).exists())
        self.assertTrue(user.username.startswith("gym_"))

    def test_gym_usernames_are_unique(self):
        gym2 = Gym.objects.create(name="باشگاه حساب ۲", code="ACCT2")
        _, _ = create_gym_account(self.gym)
        user2, _ = create_gym_account(gym2)
        self.assertNotEqual(user2.username, "gym_acct")
        self.assertFalse(User.objects.filter(username="gym_acct").count() > 1)

    def test_reset_creates_missing_account_and_changes_password(self):
        user, raw = reset_gym_account_password(self.gym)
        self.gym.refresh_from_db()
        self.assertTrue(user.check_password(raw))
        self.assertEqual(self.gym.account_user_id, user.pk)

    def test_portal_accessible_only_by_operators(self):
        unrelated = User.objects.create_user("outsider", password="pass")
        self.client.force_login(unrelated)
        self.assertEqual(self.client.get(reverse("gym_referrals:gym_portal")).status_code, 403)

        operator = User.objects.create_user("gymman", password="pass")
        GymOperator.objects.create(gym=self.gym, user=operator, is_active=True)
        self.client.force_login(operator)
        self.assertEqual(self.client.get(reverse("gym_referrals:gym_portal")).status_code, 200)

        admin = User.objects.create_superuser("root", password="x")
        self.client.force_login(admin)
        self.assertEqual(self.client.get(reverse("gym_referrals:gym_portal")).status_code, 200)

    def test_portal_redeem_registers_usage(self):
        referral = self.make_referral()
        operator = User.objects.create_user("gymman", password="pass")
        GymOperator.objects.create(gym=self.gym, user=operator, is_active=True)
        self.client.force_login(operator)
        response = self.client.post(reverse("gym_referrals:gym_portal_redeem"), {"gym": self.gym.pk, "token": str(referral.public_token)})
        self.assertNotEqual(response.status_code, 403)
        referral.refresh_from_db()
        self.assertIsNotNone(referral.usage)
        self.assertEqual(referral.status, Referral.Status.USED)

    def test_portal_redeem_rejects_foreign_token(self):
        other = Gym.objects.create(name="باشگاه دیگر", code="OTHER")
        today = timezone.localdate()
        GymContract.objects.create(gym=other, start_date=today - timedelta(days=1), end_date=today + timedelta(days=365))
        foreign = ReferralService.create_referral(
            requester=self.employee, beneficiary_type="EMPLOYEE", beneficiary_id=self.profile.pk,
            gym_id=other.pk, source="WEB",
        )[0]
        operator = User.objects.create_user("gymman", password="pass")
        GymOperator.objects.create(gym=self.gym, user=operator, is_active=True)
        self.client.force_login(operator)
        self.client.post(
            reverse("gym_referrals:gym_portal_redeem"),
            {"gym": self.gym.pk, "token": str(foreign.public_token)},
        )
        foreign.refresh_from_db()
        self.assertIsNone(getattr(foreign, "usage", None))
        self.assertEqual(foreign.status, Referral.Status.ACTIVE)


class RubikaOperatorFlowTests(TestCase):
    """Rubika-side operator panel: review, verify and redeem referrals."""

    def setUp(self):
        self.employee = User.objects.create_user("op_worker", password="pass")
        self.profile = UserProfile.objects.create(user=self.employee, personnel_code="911")
        self.gym = Gym.objects.create(name="باشگاه روبیکایی", code="RBK1")
        self.other = Gym.objects.create(name="باشگاه دیگر", code="RBK2")
        self.inactive = Gym.objects.create(name="باشگاه غیرفعال", code="RBKX", is_active=False)
        today = timezone.localdate()
        for gym in (self.gym, self.other):
            GymContract.objects.create(gym=gym, start_date=today - timedelta(days=1), end_date=today + timedelta(days=365))

        self.operator = User.objects.create_user("op_rbk", password="pass")
        self.rubika_user = RubikaUser.objects.create(chat_id="op-chat-1", user=self.operator)
        for gym in (self.gym, self.other, self.inactive):
            GymOperator.objects.create(gym=gym, user=self.operator, is_active=True)

        self.referral = ReferralService.create_referral(
            requester=self.employee, beneficiary_type="EMPLOYEE", beneficiary_id=self.profile.pk,
            gym_id=self.gym.pk, source=Referral.Source.RUBIKA,
        )[0]

    def test_find_referral_by_token_and_number(self):
        self.assertEqual(ReferralService.find_referral(str(self.referral.public_token)).pk, self.referral.pk)
        self.assertEqual(ReferralService.find_referral(self.referral.referral_number).pk, self.referral.pk)
        self.assertIsNone(ReferralService.find_referral("UNDEFINED"))

    def test_operator_gyms_excludes_inactive(self):
        gyms = async_to_sync(RubikaGymOperatorAdapter.operator_gyms)(self.rubika_user)
        ids = {g["id"] for g in gyms}
        self.assertEqual(ids, {self.gym.pk, self.other.pk})

    def test_unlinked_rubika_user_is_rejected(self):
        foreign = RubikaUser.objects.create(chat_id="op-chat-2")
        for method in (RubikaGymOperatorAdapter.list_referrals, RubikaGymOperatorAdapter.verify_code, RubikaGymOperatorAdapter.redeem_code):
            with self.assertRaises(PermissionError):
                async_to_sync(method)(foreign, self.gym.pk, self.referral.referral_number)

    def test_operator_without_gym_access_is_rejected(self):
        outsider = User.objects.create_user("op_outsider", password="pass")
        rubika_outsider = RubikaUser.objects.create(chat_id="op-chat-3", user=outsider)
        GymOperator.objects.create(gym=self.other, user=outsider, is_active=True)
        with self.assertRaises(PermissionError):
            async_to_sync(RubikaGymOperatorAdapter.list_referrals)(rubika_outsider, self.gym.pk)

    def test_verify_code_reports_status(self):
        info = async_to_sync(RubikaGymOperatorAdapter.verify_code)(
            self.rubika_user, self.gym.pk, self.referral.referral_number,
        )
        self.assertEqual(info["referral_number"], self.referral.referral_number)
        self.assertTrue(info["is_valid"])

    def test_verify_accepts_public_token(self):
        info = async_to_sync(RubikaGymOperatorAdapter.verify_code)(
            self.rubika_user, self.gym.pk, str(self.referral.public_token),
        )
        self.assertEqual(info["referral_number"], self.referral.referral_number)

    def test_verify_foreign_gym_code_is_rejected(self):
        with self.assertRaises(ReferralError) as caught:
            async_to_sync(RubikaGymOperatorAdapter.verify_code)(
                self.rubika_user, self.other.pk, self.referral.referral_number,
            )
        self.assertEqual(caught.exception.code, "REFERRAL_GYM_MISMATCH")

    def test_verify_unknown_code_is_rejected(self):
        with self.assertRaises(ReferralError) as caught:
            async_to_sync(RubikaGymOperatorAdapter.verify_code)(self.rubika_user, self.gym.pk, "UNKNOWN")
        self.assertEqual(caught.exception.code, "REFERRAL_NOT_FOUND")

    def test_redeem_registers_usage(self):
        referral = async_to_sync(RubikaGymOperatorAdapter.redeem_code)(
            self.rubika_user, self.gym.pk, self.referral.referral_number,
        )
        referral.refresh_from_db()
        self.assertEqual(referral.status, Referral.Status.USED)
        self.assertIsNotNone(referral.redeemed_at)
        self.assertEqual(referral.usage.gym_id, self.gym.pk)
        self.assertEqual(referral.usage.redeemed_by_id, self.operator.pk)

    def test_redeem_twice_is_rejected(self):
        async_to_sync(RubikaGymOperatorAdapter.redeem_code)(self.rubika_user, self.gym.pk, self.referral.referral_number)
        with self.assertRaises(ReferralError) as caught:
            async_to_sync(RubikaGymOperatorAdapter.redeem_code)(self.rubika_user, self.gym.pk, self.referral.referral_number)
        self.assertEqual(caught.exception.code, "REFERRAL_ALREADY_USED")

    def test_redeem_foreign_gym_is_rejected(self):
        with self.assertRaises(ReferralError) as caught:
            async_to_sync(RubikaGymOperatorAdapter.redeem_code)(
                self.rubika_user, self.other.pk, self.referral.referral_number,
            )
        self.assertEqual(caught.exception.code, "REFERRAL_GYM_MISMATCH")

    def test_new_referral_notifies_linked_operators(self):
        with patch("rubika_bot.tasks.send_rubika_message") as sender:
            sent = notify_gym_operators_new_referral.run(self.referral.pk)
        self.assertEqual(sent, 1)
        sender.delay.assert_called_once()
        chat_id, text = sender.delay.call_args[0]
        self.assertEqual(chat_id, "op-chat-1")
        self.assertIn(self.gym.name, text)
        self.assertIn(self.referral.referral_number, text)
