import csv
import io
from datetime import date
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.db import IntegrityError
from django.db.models import Q
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST

from accounts.models import Dependent, UserProfile
from permissions.utils import check_permission
from .access import gym_access_required
from .billing import BillingError, GymInvoiceService
from .forms import GymContractForm, GymForm, GymOperatorForm, InvoiceAdjustmentForm, InvoiceGenerateForm, LegacyImportForm, ReferralCancelForm, ReferralCreateForm
from .models import Gym, GymContract, GymInvoice, GymOperator, Referral, ReferralAuditLog, ReferralUsage
from .services import ReferralError, ReferralService


def _spreadsheet_safe(value):
    """Prevent user-controlled text from being interpreted as a formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@permission_required("gym_referrals.view_gym", raise_exception=True)
@login_required
@gym_access_required("gym_manage")
def gym_dashboard(request):
    today = timezone.localdate()
    gyms = Gym.objects.all()
    active_gyms = gyms.filter(is_active=True)
    contracts = GymContract.objects.filter(is_active=True, start_date__lte=today, end_date__gte=today)
    operators = GymOperator.objects.filter(is_active=True)
    active_referrals = Referral.objects.filter(
        status=Referral.Status.ACTIVE, valid_from__lte=today, valid_until__gte=today
    )
    recent_referrals = Referral.objects.select_related("employee", "gym").order_by("-created_at")[:5]
    
    stats = {
        "total_gyms": gyms.count(),
        "active_gyms": active_gyms.count(),
        "total_contracts": GymContract.objects.count(),
        "active_contracts": contracts.count(),
        "total_operators": GymOperator.objects.count(),
        "active_operators": operators.count(),
        "total_referrals": Referral.objects.count(),
        "active_referrals": active_referrals.count(),
        "used_referrals": Referral.objects.filter(status=Referral.Status.USED).count(),
        "expired_referrals": Referral.objects.filter(status=Referral.Status.EXPIRED).count(),
    }
    
    onboarding = {
        "has_gyms": gyms.exists(),
        "has_contracts": GymContract.objects.exists(),
        "has_operators": GymOperator.objects.exists(),
        "can_create_referral": gyms.filter(is_active=True, contracts__is_active=True).exists(),
    }
    onboarding["setup_complete"] = all(onboarding.values())
    
    return render(request, "gym_referrals/dashboard.html", {
        "stats": stats,
        "onboarding": onboarding,
        "recent_referrals": recent_referrals,
        "today": today,
    })


def _decorate_contract(contract, today):
    if not contract.is_active:
        contract.ui_status = "غیرفعال"
        contract.ui_status_class = "bg-gray-100 text-gray-700"
    elif contract.start_date > today:
        contract.ui_status = "آتی"
        contract.ui_status_class = "bg-blue-100 text-blue-700"
    elif contract.end_date < today:
        contract.ui_status = "منقضی"
        contract.ui_status_class = "bg-amber-100 text-amber-700"
    else:
        contract.ui_status = "فعال"
        contract.ui_status_class = "bg-emerald-100 text-emerald-700"
    reason_labels = {"INACTIVE_CONTRACT": "قرارداد غیرفعال است.", "INVALID_CONTRACT_TARIFF": "تعرفه تعیین نشده است.", "INVALID_BILLING_POLICY": "روش صورتحساب معتبر نیست."}
    contract.ui_readiness_errors = [reason_labels.get(code, code) for code in contract.billing_readiness_errors()]
    return contract


def _page(request, queryset, size=25):
    return Paginator(queryset, size).get_page(request.GET.get("page"))


def _filter_referrals(request, queryset):
    q = request.GET.get("q", "").strip()
    if q:
        queryset = queryset.filter(Q(referral_number__icontains=q) | Q(employee_full_name_snapshot__icontains=q) | Q(beneficiary_full_name_snapshot__icontains=q))
    gym_id = request.GET.get("gym", "")
    status = request.GET.get("status", "")
    source = request.GET.get("source", "")
    if gym_id.isdigit():
        queryset = queryset.filter(gym_id=gym_id)
    if status in Referral.Status.values:
        queryset = queryset.filter(status=status)
    if source in Referral.Source.values:
        queryset = queryset.filter(source=source)
    return queryset


@login_required
@gym_access_required("gym_referral_create")
@require_http_methods(["GET", "POST"])
def create_page(request):
    form = ReferralCreateForm(request.POST or None, user=request.user)
    if request.method == "POST" and form.is_valid():
        try:
            referral, _ = ReferralService.create_referral(
                requester=request.user, beneficiary_type=form.cleaned_data["beneficiary_type"],
                beneficiary_id=form.cleaned_data["beneficiary_id"], gym_id=form.cleaned_data["gym"].pk,
                source=Referral.Source.WEB, idempotency_key=request.POST.get("idempotency_key"),
            )
            messages.success(request, f"معرفی‌نامه {referral.referral_number} صادر شد.")
            return redirect(f'{reverse("gym_referrals:history")}?issued={referral.referral_number}')
        except ReferralError as exc:
            form.add_error(None, exc.message)
    today = timezone.localdate()
    active_referrals = Referral.objects.filter(
        employee=request.user, status=Referral.Status.ACTIVE,
        valid_from__lte=today, valid_until__gte=today,
    ).select_related("gym")
    available_gyms = Gym.objects.filter(
        is_active=True, contracts__is_active=True,
        contracts__start_date__lte=today, contracts__end_date__gte=today,
    ).distinct().count()
    return render(request, "gym_referrals/create.html", {
        "form": form, "active_referrals": active_referrals,
        "available_gyms": available_gyms,
    })


@permission_required("gym_referrals.view_gym", raise_exception=True)
@gym_access_required("gym_manage")
def gym_list(request):
    today = timezone.localdate()
    gyms = list(Gym.objects.prefetch_related("contracts"))
    for gym in gyms:
        gym.current_contract = next((item for item in gym.contracts.all() if item.is_valid_on(today)), None)
    return render(request, "gym_referrals/gym_list.html", {"gyms": gyms, "today": today})


@permission_required("gym_referrals.add_gym", raise_exception=True)
@gym_access_required("gym_manage")
@require_http_methods(["GET", "POST"])
def gym_create(request):
    form = GymForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        gym = form.save()
        messages.success(request, "باشگاه با موفقیت ایجاد شد.")
        return redirect("gym_referrals:gym_detail", gym.pk)
    return render(request, "gym_referrals/gym_form.html", {"form": form, "page_title": "افزودن باشگاه"})


@permission_required("gym_referrals.change_gym", raise_exception=True)
@gym_access_required("gym_manage")
@require_http_methods(["GET", "POST"])
def gym_edit(request, gym_id):
    gym = get_object_or_404(Gym, pk=gym_id)
    form = GymForm(request.POST or None, instance=gym)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "اطلاعات باشگاه به‌روز شد.")
        return redirect("gym_referrals:gym_detail", gym.pk)
    return render(request, "gym_referrals/gym_form.html", {"form": form, "gym": gym, "page_title": "ویرایش باشگاه"})


@permission_required("gym_referrals.view_gym", raise_exception=True)
@gym_access_required("gym_manage")
def gym_detail(request, gym_id):
    gym = get_object_or_404(Gym, pk=gym_id)
    contracts = list(gym.contracts.all().order_by("-start_date"))
    today = timezone.localdate()
    for contract in contracts:
        _decorate_contract(contract, today)
    current_contract = next((item for item in contracts if item.is_valid_on(today)), None)
    operators = gym.operators.select_related("user").order_by("user__first_name", "user__last_name")
    active_referral_count = gym.referrals.filter(
        status=Referral.Status.ACTIVE, valid_from__lte=today, valid_until__gte=today
    ).count()
    return render(request, "gym_referrals/gym_detail.html", {
        "gym": gym, "contracts": contracts, "current_contract": current_contract,
        "operators": operators, "active_referral_count": active_referral_count, "today": today,
    })


@permission_required("gym_referrals.view_gymcontract", raise_exception=True)
@gym_access_required("gym_contract_manage")
def contract_list(request):
    today = timezone.localdate()
    contracts = [_decorate_contract(item, today) for item in GymContract.objects.select_related("gym")]
    return render(request, "gym_referrals/contract_list.html", {"contracts": contracts, "today": today})


@permission_required("gym_referrals.view_gymcontract", raise_exception=True)
@gym_access_required("gym_contract_manage")
def contract_detail(request, contract_id):
    contract = get_object_or_404(GymContract.objects.select_related("gym"), pk=contract_id)
    _decorate_contract(contract, timezone.localdate())
    return render(request, "gym_referrals/contract_detail.html", {"contract": contract})


@permission_required("gym_referrals.add_gymcontract", raise_exception=True)
@gym_access_required("gym_contract_manage")
@require_http_methods(["GET", "POST"])
def contract_create(request):
    initial = {"gym": request.GET.get("gym")} if request.GET.get("gym") else None
    form = GymContractForm(request.POST or None, initial=initial)
    if request.method == "POST" and form.is_valid():
        contract = form.save()
        messages.success(request, "قرارداد باشگاه ایجاد شد.")
        return redirect("gym_referrals:gym_detail", contract.gym_id)
    return render(request, "gym_referrals/contract_form.html", {"form": form, "page_title": "افزودن قرارداد"})


@permission_required("gym_referrals.add_gymcontract", raise_exception=True)
@gym_access_required("gym_contract_manage")
@require_http_methods(["GET", "POST"])
def gym_contract_create(request, gym_id):
    gym = get_object_or_404(Gym, pk=gym_id)
    form = GymContractForm(request.POST or None, gym=gym)
    if request.method == "POST" and form.is_valid():
        contract = form.save(commit=False)
        contract.gym = gym
        contract.save()
        messages.success(request, "قرارداد باشگاه با موفقیت ثبت شد.")
        return redirect("gym_referrals:gym_detail", gym.pk)
    return render(request, "gym_referrals/contract_form.html", {
        "form": form, "gym": gym, "page_title": "افزودن قرارداد",
        "cancel_url": reverse("gym_referrals:gym_detail", args=(gym.pk,)),
    })


@permission_required("gym_referrals.change_gymcontract", raise_exception=True)
@gym_access_required("gym_contract_manage")
@require_http_methods(["GET", "POST"])
def contract_edit(request, contract_id):
    contract = get_object_or_404(GymContract, pk=contract_id)
    form = GymContractForm(request.POST or None, instance=contract)
    if request.method == "POST" and form.is_valid():
        contract = form.save()
        messages.success(request, "قرارداد باشگاه به‌روز شد.")
        return redirect("gym_referrals:gym_detail", contract.gym_id)
    return render(request, "gym_referrals/contract_form.html", {
        "form": form, "contract": contract, "page_title": "ویرایش قرارداد",
        "cancel_url": reverse("gym_referrals:gym_detail", args=(contract.gym_id,)),
    })


@login_required
@gym_access_required("gym_referral_history")
def history_page(request):
    referrals = Referral.objects.filter(employee=request.user).select_related("gym")
    issued = None
    issued_number = request.GET.get("issued")
    if issued_number:
        issued = referrals.filter(referral_number=issued_number).first()
    return render(request, "gym_referrals/history.html", {"referrals": referrals, "today": timezone.localdate(), "issued": issued})


@permission_required("gym_referrals.view_referral", raise_exception=True)
def referral_admin_list(request):
    queryset = _filter_referrals(request, Referral.objects.select_related("employee", "gym").all())
    return render(request, "gym_referrals/referral_admin_list.html", {"page_obj": _page(request, queryset), "gyms": Gym.objects.all(), "statuses": Referral.Status.choices, "sources": Referral.Source.choices})


@login_required
@gym_access_required("gym_referral_history")
def referral_detail(request, referral_number):
    queryset = Referral.objects.select_related("employee", "gym", "cancelled_by").prefetch_related("audit_logs__actor")
    if not (request.user.has_perm("gym_referrals.view_referral") or any(check_permission(request.user, "gym_referral_management").values())):
        queryset = queryset.filter(employee=request.user)
    referral = get_object_or_404(queryset, referral_number=referral_number)
    return render(request, "gym_referrals/referral_detail.html", {"referral": referral})


@permission_required("gym_referrals.view_gymoperator", raise_exception=True)
@gym_access_required("gym_operator_manage")
def operator_list(request):
    operators = GymOperator.objects.select_related("gym", "user")
    gym_id = request.GET.get("gym")
    if gym_id and gym_id.isdigit(): operators = operators.filter(gym_id=gym_id)
    return render(request, "gym_referrals/operator_list.html", {"operators": operators, "gyms": Gym.objects.all()})


@permission_required("gym_referrals.add_gymoperator", raise_exception=True)
@gym_access_required("gym_operator_manage")
@require_http_methods(["GET", "POST"])
def operator_create(request, gym_id=None):
    gym = get_object_or_404(Gym, pk=gym_id) if gym_id else None
    form = GymOperatorForm(request.POST or None, gym=gym)
    if request.method == "POST" and form.is_valid():
        try:
            form.save(); messages.success(request, "اپراتور باشگاه ثبت شد.")
            return redirect("gym_referrals:operator_list")
        except IntegrityError:
            form.add_error("user", "این کاربر قبلاً برای این باشگاه ثبت شده است.")
    return render(request, "gym_referrals/operator_form.html", {"form": form, "gym": gym})


@permission_required("gym_referrals.delete_gymoperator", raise_exception=True)
@gym_access_required("gym_operator_manage")
@require_POST
def operator_delete(request, operator_id):
    get_object_or_404(GymOperator, pk=operator_id).delete()
    messages.success(request, "ارتباط اپراتور با باشگاه حذف شد.")
    return redirect("gym_referrals:operator_list")


@permission_required("gym_referrals.view_referralusage", raise_exception=True)
@gym_access_required("gym_referral_usage")
def usage_list(request):
    usages = ReferralUsage.objects.select_related("referral__employee", "referral__gym", "redeemed_by")
    gym_id, q = request.GET.get("gym"), request.GET.get("q", "").strip()
    if gym_id and gym_id.isdigit(): usages = usages.filter(gym_id=gym_id)
    if q: usages = usages.filter(Q(referral__referral_number__icontains=q) | Q(referral__beneficiary_full_name_snapshot__icontains=q) | Q(referral__employee_full_name_snapshot__icontains=q))
    return render(request, "gym_referrals/usage_list.html", {"page_obj": _page(request, usages.order_by("-redeemed_at")), "gyms": Gym.objects.all()})


@permission_required("gym_referrals.view_referral", raise_exception=True)
@gym_access_required("gym_referral_management")
def legacy_list(request):
    queryset = Referral.objects.filter(source=Referral.Source.PHYSICAL_LEGACY).select_related("gym", "employee")
    return render(request, "gym_referrals/legacy_list.html", {"page_obj": _page(request, queryset)})


@permission_required("gym_referrals.view_referral", raise_exception=True)
@gym_access_required("gym_referral_management")
def reports_page(request):
    today = timezone.localdate()
    stats = {"active": Referral.objects.filter(status=Referral.Status.ACTIVE, valid_until__gte=today).count(), "issued": Referral.objects.count(), "usage": ReferralUsage.objects.count(), "expired": Referral.objects.filter(status=Referral.Status.EXPIRED).count(), "gyms": Gym.objects.filter(is_active=True).count(), "contracts": GymContract.objects.filter(is_active=True, start_date__lte=today, end_date__gte=today).count()}
    finance = None
    if request.user.has_perm("gym_referrals.view_gyminvoice"):
        finance = {key.lower(): GymInvoice.objects.filter(status=key).count() for key in ("DRAFT", "APPROVED", "PAID")}
    return render(request, "gym_referrals/reports.html", {"stats": stats, "finance": finance})


@login_required
@gym_access_required("gym_referral_history")
@require_POST
def cancel_page(request, referral_number):
    referral = get_object_or_404(Referral, referral_number=referral_number, employee=request.user)
    form = ReferralCancelForm(request.POST)
    if form.is_valid():
        try:
            ReferralService.cancel(referral, request.user, form.cleaned_data["reason"])
            messages.success(request, "معرفی‌نامه لغو شد.")
        except ReferralError as exc:
            messages.error(request, exc.message)
    return redirect("gym_referrals:history")


MAX_LEGACY_IMPORT_ROWS = 500


@permission_required("gym_referrals.add_referral", raise_exception=True)
@gym_access_required("gym_referral_management")
@require_http_methods(["GET", "POST"])
def legacy_import(request):
    form, report, truncated = LegacyImportForm(request.POST or None, request.FILES or None), [], False
    if request.method == "POST" and form.is_valid():
        uploaded_file = form.cleaned_data["file"]
        file_name = uploaded_file.name.lower()
        try:
            if file_name.endswith(".csv"):
                rows = _parse_csv(uploaded_file)
            else:
                rows = _parse_excel(uploaded_file)
            for line, row in enumerate(rows, 2):
                if line > MAX_LEGACY_IMPORT_ROWS + 1:
                    truncated = True
                    break
                try:
                    profile = UserProfile.objects.select_related("user").get(personnel_code=row["personnel_no"])
                    employee = profile.user
                    dn_code = row.get("dependent_national_code", "").strip()
                    if dn_code:
                        btype = "DEPENDENT"
                        dependent = profile.dependents.get(national_code=dn_code)
                        bid = dependent.pk
                    else:
                        btype = "EMPLOYEE"
                        bid = profile.pk
                    issue_date = date.fromisoformat(row["issue_date"])
                    valid_from = date.fromisoformat(row["valid_from"])
                    valid_until = date.fromisoformat(row["valid_until"])
                    referral, created = ReferralService.create_referral(
                        requester=employee, beneficiary_type=btype, beneficiary_id=bid, gym_id=int(row["gym_id"]),
                        source=Referral.Source.PHYSICAL_LEGACY, idempotency_key=f"legacy:{row['legacy_letter_no']}",
                        issue_date=issue_date, valid_from=valid_from, valid_until=valid_until,
                        legacy_letter_no=row["legacy_letter_no"], notes=row.get("notes", ""), actor=request.user,
                    )
                    report.append({"line": line, "ok": True, "detail": referral.referral_number if created else "تکراری"})
                except (KeyError, ValueError, ReferralError, IntegrityError, UserProfile.DoesNotExist, Dependent.DoesNotExist) as exc:
                    report.append({"line": line, "ok": False, "detail": getattr(exc, "message", str(exc))})
        except UnicodeDecodeError:
            form.add_error("file", "فایل باید UTF-8 باشد.")
        except ImportError:
            form.add_error("file", "برای خواندن فایل Excel، کتابخانه openpyxl نصب نیست.")
        except Exception as exc:
            form.add_error("file", f"خطا در خواندن فایل: {exc}")
    summary = {
        "total": len(report),
        "successful": sum(1 for item in report if item["ok"]),
        "failed": sum(1 for item in report if not item["ok"]),
        "truncated": truncated,
        "max_rows": MAX_LEGACY_IMPORT_ROWS,
    }
    if truncated:
        messages.warning(request, f"پردازش پس از {MAX_LEGACY_IMPORT_ROWS} ردیف متوقف شد؛ فایل را به بخش‌های کوچک‌تر تقسیم کنید.")
    return render(request, "gym_referrals/import.html", {"form": form, "report": report, "summary": summary})


def _parse_csv(uploaded_file):
    rows_dict = csv.DictReader(io.StringIO(uploaded_file.read().decode("utf-8-sig")))
    return [_normalize_row(row) for row in rows_dict]


def _parse_excel(uploaded_file):
    from openpyxl import load_workbook
    wb = load_workbook(uploaded_file, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers_raw = next(rows_iter, None)
    if not headers_raw:
        return []
    headers = [str(h).strip().lower().replace(" ", "_").replace("\u200c", "_") for h in headers_raw]
    result = []
    for row_values in rows_iter:
        if all(v is None for v in row_values):
            continue
        row_dict = {}
        for idx, header in enumerate(headers):
            val = row_values[idx] if idx < len(row_values) else None
            if val is None:
                row_dict[header] = ""
            elif isinstance(val, float) and val == int(val):
                row_dict[header] = str(int(val))
            else:
                row_dict[header] = str(val).strip()
        result.append(_normalize_row(row_dict))
    return result


def _normalize_row(row):
    mapping = {
        "personnel_no": "personnel_no",
        "personnel_code": "personnel_no",
        "personnel_number": "personnel_no",
        "کد پرسنلی": "personnel_no",
        "کد_پرسنلی": "personnel_no",
        "dependent_national_code": "dependent_national_code",
        "dependent_code": "dependent_national_code",
        "dnational_code": "dependent_national_code",
        "کد ملی تحت تکفل": "dependent_national_code",
        "کد_ملی_تحت_تکفل": "dependent_national_code",
        "کد ملی": "dependent_national_code",
        "کد_ملی": "dependent_national_code",
        "gym_id": "gym_id",
        "gym": "gym_id",
        "شناسه باشگاه": "gym_id",
        "شناسه_باشگاه": "gym_id",
        "legacy_letter_no": "legacy_letter_no",
        "letter_no": "legacy_letter_no",
        "letter_number": "legacy_letter_no",
        "شماره معرفی‌نامه": "legacy_letter_no",
        "شماره_معرفی_نامه": "legacy_letter_no",
        "issue_date": "issue_date",
        "تاریخ صدور": "issue_date",
        "تاریخ_صدور": "issue_date",
        "valid_from": "valid_from",
        "تاریخ شروع": "valid_from",
        "تاریخ_شروع": "valid_from",
        "valid_until": "valid_until",
        "تاریخ پایان": "valid_until",
        "تاریخ_پایان": "valid_until",
        "notes": "notes",
        "description": "notes",
        "توضیحات": "notes",
    }
    normalized = {}
    for key, value in row.items():
        clean_key = key.strip().lower().replace(" ", "_").replace("\u200c", "_")
        canonical = mapping.get(clean_key, clean_key)
        normalized[canonical] = str(value).strip() if value is not None else ""
    return normalized


@login_required
@gym_access_required("gym_referral_management")
def legacy_import_sample(request):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "نمونه ورود اطلاعات"
    ws.sheet_view.rightToLeft = True

    headers = [
        ("کد پرسنلی", "کد پرسنلی کارمند"),
        ("کد ملی تحت تکفل", "برای صدور معرفی‌نامه همسر/فرزند؛ برای خود کارمند خالی بماند"),
        ("شناسه باشگاه", "شناسه باشگاه در سامانه"),
        ("شماره معرفی‌نامه", "شماره معرفی‌نامه فیزیکی"),
        ("تاریخ صدور", "فرمت YYYY-MM-DD"),
        ("تاریخ شروع", "فرمت YYYY-MM-DD"),
        ("تاریخ پایان", "فرمت YYYY-MM-DD"),
        ("توضیحات", "اختیاری"),
    ]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="Tahoma", bold=True, color="FFFFFF", size=11)
    cell_font = Font(name="Tahoma", size=10)
    example_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    col_widths = [18, 30, 20, 26, 18, 18, 18, 26]
    for col_idx, (field_name, description) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=field_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        cell.comment = Comment(description, "سامانه باشگاه‌ها")

        ws.column_dimensions[cell.column_letter].width = col_widths[col_idx - 1]

    example_rows = [
        ["110001", "", "1", "PH-1403-001", "1403-06-15", "1403-06-15", "1403-12-29", "معرفی‌نامه کاغذی قدیمی"],
        ["110001", "0012345678", "1", "PH-1403-002", "1403-07-01", "1403-07-01", "1403-12-29", "معرفی همسر"],
    ]
    for row_idx, row_values in enumerate(example_rows, 2):
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = example_fill
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    ws.auto_filter.ref = "A1:H3"
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 110
    guide_text = [
        "راهنمای ورود اطلاعات",
        "",
        "1) ستون‌ها باید فارسی باشد: کد پرسنلی / کد ملی تحت تکفل / شناسه باشگاه / شماره معرفی‌نامه / تاریخ صدور / تاریخ شروع / تاریخ پایان / توضیحات",
        "2) برای صدور معرفی‌نامه برای خود کارمند: فقط کد پرسنلی را پر کنید و ستون کد ملی تحت تکفل را خالی بگذارید.",
        "3) برای همسر یا فرزند: کد پرسنلی کارمند و کد ملی فرد تحت تکفل را وارد کنید.",
        "4) تاریخ‌ها به فرمت YYYY-MM-DD باشد.",
        "5) کد پرسنلی و کد ملی افراد باید قبلاً در سامانه ثبت شده باشند.",
        "6) ردیف‌های نمونه در برگه اول را با اطلاعات واقعی جایگزین کنید.",
    ]
    for idx, text in enumerate(guide_text, 1):
        cell = guide.cell(row=idx, column=1, value=text)
        if idx == 1:
            cell.font = Font(name="Tahoma", bold=True, size=13)
        else:
            cell.font = Font(name="Tahoma", size=11)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="legacy_import_sample.xlsx"'
    return response


@permission_required("gym_referrals.view_gyminvoice", raise_exception=True)
@gym_access_required("gym_invoice_view")
@require_http_methods(["GET", "POST"])
def invoice_list(request):
    form = InvoiceGenerateForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        if not request.user.has_perm("gym_referrals.generate_gym_invoice"):
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
        try:
            invoice = GymInvoiceService.generate_invoice(
                contract_id=form.cleaned_data["contract"].pk, year=form.cleaned_data["year"],
                month=form.cleaned_data["month"], actor=request.user,
            )
            messages.success(request, "صورتحساب ایجاد شد.")
            return redirect("gym_referrals:invoice_detail", invoice.pk)
        except BillingError as exc:
            form.add_error(None, exc.message)
    invoices = GymInvoice.objects.select_related("gym", "gym_contract", "created_by")
    return render(request, "gym_referrals/invoice_list.html", {"form": form, "invoices": invoices})


@permission_required("gym_referrals.view_gyminvoice", raise_exception=True)
@gym_access_required("gym_invoice_view")
@require_http_methods(["GET", "POST"])
def invoice_detail(request, invoice_id):
    invoice = get_object_or_404(GymInvoice.objects.select_related("gym", "gym_contract"), pk=invoice_id)
    form = InvoiceAdjustmentForm(request.POST or None, initial={"adjustment_amount": invoice.adjustment_amount, "notes": invoice.notes})
    if request.method == "POST" and request.POST.get("action") == "adjust" and form.is_valid():
        if not request.user.has_perm("gym_referrals.adjust_gym_invoice"):
            from django.core.exceptions import PermissionDenied
            raise PermissionDenied
        try:
            invoice = GymInvoiceService.set_adjustment(
                invoice.pk, form.cleaned_data["adjustment_amount"], request.user, form.cleaned_data["notes"]
            )
            messages.success(request, "مبالغ صورتحساب به‌روز شد.")
            return redirect("gym_referrals:invoice_detail", invoice.pk)
        except BillingError as exc:
            form.add_error(None, exc.message)
    return render(request, "gym_referrals/invoice_detail.html", {"invoice": invoice, "items": invoice.items.select_related("referral", "referral_usage"), "form": form})


@login_required
@gym_access_required("gym_invoice_view")
@require_POST
def invoice_transition(request, invoice_id, status):
    required = {"FINALIZED": "gym_referrals.finalize_gym_invoice", "APPROVED": "gym_referrals.approve_gym_invoice", "PAID": "gym_referrals.mark_gym_invoice_paid", "CANCELLED": "gym_referrals.cancel_gym_invoice"}.get(status.upper())
    if not required or not request.user.has_perm(required):
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    try:
        GymInvoiceService.transition(invoice_id, status.upper(), request.user)
        messages.success(request, "وضعیت صورتحساب تغییر کرد.")
    except (GymInvoice.DoesNotExist, BillingError) as exc:
        messages.error(request, getattr(exc, "message", "صورتحساب یافت نشد."))
    return redirect("gym_referrals:invoice_detail", invoice_id)


@permission_required("gym_referrals.export_gym_invoice", raise_exception=True)
@gym_access_required("gym_invoice_view")
def invoice_export_excel(request, invoice_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from .presentation import format_jalali
    invoice = get_object_or_404(GymInvoice.objects.select_related("gym", "gym_contract"), pk=invoice_id)
    workbook, sheet = Workbook(), None
    sheet = workbook.active
    sheet.title = "صورتحساب"
    header = [("شماره صورتحساب", invoice.invoice_number), ("باشگاه", invoice.gym.name), ("قرارداد", invoice.gym_contract_id), ("دوره", f"{format_jalali(invoice.period_start)} تا {format_jalali(invoice.period_end)}"), ("مبنای محاسبه", invoice.get_billing_policy_display()), ("وضعیت", invoice.get_status_display()), ("تعداد", invoice.item_count), ("جمع", invoice.subtotal), ("تعدیل", invoice.adjustment_amount), ("مبلغ نهایی", invoice.total_amount)]
    for row, (label, value) in enumerate(header, 1):
        sheet.cell(row, 1, label).font = Font(bold=True)
        sheet.cell(row, 2, _spreadsheet_safe(value))
    table_row = len(header) + 2
    columns = ["ردیف", "نام پرسنل", "شماره پرسنلی", "نام ذی‌نفع", "نسبت", "شماره معرفی", "منبع", "تاریخ صدور", "تاریخ استفاده", "قیمت واحد", "تعداد", "جمع ردیف"]
    for col, value in enumerate(columns, 1):
        sheet.cell(table_row, col, value).font = Font(bold=True)
        sheet.cell(table_row, col).alignment = Alignment(horizontal="center")
    for index, item in enumerate(invoice.items.all(), 1):
        values = [index, item.employee_name_snapshot, item.personnel_no_snapshot, item.beneficiary_name_snapshot, item.relation_snapshot, item.referral_number, item.get_referral_source_display(), format_jalali(item.issue_date), format_jalali(item.usage_date), item.unit_price, item.quantity, item.line_total]
        for col, value in enumerate(values, 1):
            sheet.cell(table_row + index, col, _spreadsheet_safe(value))
    sheet.sheet_view.rightToLeft = True
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{invoice.invoice_number}.xlsx"'
    return response


@login_required
@gym_access_required("gym_referral_history")
def referral_qr(request, referral_number):
    import qrcode
    referral = get_object_or_404(Referral, referral_number=referral_number, employee=request.user)
    verification_url = request.build_absolute_uri(reverse("gym_referrals:api_verify", args=(referral.public_token,)))
    image = qrcode.make(verification_url)
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return HttpResponse(buffer.getvalue(), content_type="image/png")
