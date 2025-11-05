# views.py
from django.contrib import messages
from django.shortcuts import render, redirect
from django.http import HttpResponseForbidden, JsonResponse, HttpResponse
from permissions.utils import permission_required
from .forms import ReportForm, ReportFilterForm, persian_to_english_numbers
from django.contrib.auth.decorators import login_required
from django.contrib.auth import authenticate, login as auth_login
from .models import Report
from django.utils import timezone
from django.db.models import Q, Count
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from .utils import get_current_user_shift_and_group, SHIFT_PATTERN
from django.shortcuts import render, redirect, get_object_or_404
from .models import Contractor, Employee, Vehicle  # Vehicle را هم ایمپورت کنید
from django.contrib.auth.decorators import user_passes_test
from dashboard.utils import log_user_activity  # اضافه کردن ایمپورت
from django.urls import reverse  # برای ساخت URL
import json
import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
import logging
import re
from shift_manager.utils import get_shift_for_date
import datetime
from datetime import date, timedelta

logger = logging.getLogger(__name__)

@permission_required("create_report")
@login_required
def create_report(request):
    logger.info("Entering create_report view")
    
    # ثبت فعالیت مشاهده فرم ایجاد گزارش
    if request.method == 'GET':
        log_user_activity(
            user=request.user,
            activity_type='view',
            description='مشاهده فرم ایجاد گزارش جدید',
            related_model='Report',
            related_object_id=None,
            url=reverse('contractor_management:create_report'),
            request=request
        )
    
    if request.method == 'POST':
        post_data = request.POST.copy()
        logger.info(f"داده‌های POST اصلی: {post_data}")

        gdate = None # Initialize gdate
        report_date_input = post_data.get('reeport_date') or post_data.get('report_date')
        if report_date_input:
            try:
                # تبدیل اعداد فارسی به انگلیسی
                report_date = persian_to_english_numbers(report_date_input.strip())
                # حذف کاراکترهای اضافی
                report_date = re.sub(r'[^0-9/]', '', report_date)
                year, month, day = map(int, report_date.split('/'))
                
                # تصحیح سال دو رقمی
                if year < 100:
                    year += 1400
                
                # ساخت تاریخ شمسی
                jdate = jdatetime.date(year, month, day)
                gdate = jdate.togregorian() # Assign the Gregorian date object
                post_data['report_date'] = gdate.strftime('%Y-%m-%d')
                
                # دریافت شیفت کاری بر اساس تاریخ و گروه کاربر
                if hasattr(request.user, 'userprofile'):
                    shifts = get_shift_for_date(gdate, request.user.userprofile)
                    user_shift = shifts.get('user_group_shift')
                    if user_shift:
                        post_data['shift'] = user_shift
                        logger.info(f"شیفت کاری تشخیص داده شده: {user_shift}")
                    else:
                        logger.warning("شیفت کاری برای کاربر تشخیص داده نشد")
                
                logger.info(f"تاریخ میلادی محاسبه شده: {gdate}, Type: {type(gdate)}") # Log the calculated Gregorian date
                logger.info(f"تاریخ رشته‌ای برای فرم: {post_data['report_date']}")
            except (ValueError, IndexError, AttributeError) as e:
                logger.error(f"خطا در تبدیل تاریخ: {str(e)}")
                error_msg = 'لطفاً تاریخ را به فرمت صحیح وارد کنید (مثال: 1402/12/29)'
                
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': False,
                        'errors': {'report_date': [error_msg]},
                        'message': error_msg
                    }, status=400)
                
                messages.error(request, error_msg)
                form = ReportForm()
                return render(request, 'contractor_management/report_form.html', {'form': form})
        else:
             error_msg = 'فیلد تاریخ کارکرد الزامی است.'
             
             if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                 return JsonResponse({
                     'success': False,
                     'errors': {'report_date': [error_msg]},
                     'message': error_msg
                 }, status=400)
             
             messages.error(request, error_msg)
             return render(request, 'contractor_management/report_form.html', {'form': ReportForm(post_data)})

        logger.info(f"داده‌های POST نهایی برای فرم: {post_data}")
        form = ReportForm(post_data)
        if form.is_valid():
            logger.info("فرم معتبر است")
            report = form.save(commit=False)
            report.user = request.user
            report.report_datetime = timezone.now()

            # --- شروع اصلاح ---
            # اطمینان از اینکه تاریخ میلادی استاندارد ذخیره می‌شود
            # حتی اگر فرم به اشتباه jdatetime.date برگردانده باشد
            if gdate and isinstance(gdate, datetime.date):
                 report.report_date = gdate
                 logger.info(f"Report.report_date explicitly set to Gregorian date: {report.report_date}, Type: {type(report.report_date)}")
            else:
                 # اگر gdate محاسبه نشده یا نوع آن اشتباه است، سعی کن از فرم بگیری
                 # اما ممکن است هنوز مشکل داشته باشد.
                 logger.warning(f"gdate was not a valid datetime.date. Using value from form: {report.report_date}, Type: {type(report.report_date)}")
                 # اگر نوع تاریخ فرم همچنان jdatetime.date است، خطا بده
                 if isinstance(report.report_date, jdatetime.date):
                     logger.error("Form provided jdatetime.date, cannot save correctly.")
                     messages.error(request, "خطای داخلی در پردازش تاریخ رخ داده است. لطفا با پشتیبانی تماس بگیرید.")
                     return render(request, 'contractor_management/report_form.html', {'form': form})
            # --- پایان اصلاح ---

            # اضافه کردن پیمانکار بر اساس خودرو
            selected_vehicle = form.cleaned_data.get('vehicle')
            logger.info(f"خودروی انتخاب شده: {selected_vehicle}")
            
            if selected_vehicle:
                report.contractor = selected_vehicle.contractor
                logger.info(f"پیمانکار تنظیم شده: {report.contractor}")

            user_shift, user_group = get_current_user_shift_and_group(report.user)
            logger.info(f"شیفت و گروه کاربر: شیفت={user_shift}, گروه={user_group}")

            if user_group:
                report.group = user_group
            else:
                # اگر گروه کاری تشخیص داده نشد، از پروفایل کاربر استفاده کنیم
                if hasattr(request.user, 'userprofile') and request.user.userprofile.group:
                    report.group = request.user.userprofile.group
                    logger.info(f"گروه کاری از پروفایل کاربر استخراج شد: {report.group}")
                else:
                    logger.warning("گروه کاری نامعتبر است")
                    error_msg = 'امکان ثبت گزارش در این بازه زمانی وجود ندارد.'
                    
                    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                        return JsonResponse({
                            'success': False,
                            'message': error_msg
                        }, status=400)
                    
                    messages.error(request, error_msg)
                    return render(request, 'contractor_management/report_form.html', {
                        'form': form
                    })

            try:
                logger.info("در حال ذخیره گزارش...")
                
                # بررسی شیفت کاری کاربر (با تاریخ میلادی صحیح)
                if hasattr(request.user, 'userprofile') and report.report_date:
                     try:
                        shifts = get_shift_for_date(report.report_date, request.user.userprofile)
                        user_shift = shifts.get('user_group_shift', '')
                        
                        if 'OFF' in user_shift:
                            logger.warning(f"کاربر در تاریخ {report.report_date} در شیفت {user_shift} (OFF) است")
                            # نمایش تاریخ شمسی در پیام خطا
                            jdate_display = jdatetime.date.fromgregorian(date=report.report_date).strftime('%Y/%m/%d')
                            error_msg = f'شما در تاریخ {jdate_display} در شیفت {user_shift} (OFF) هستید و امکان ثبت گزارش ندارید.'
                            
                            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                                return JsonResponse({
                                    'success': False,
                                    'message': error_msg
                                }, status=400)
                            
                            messages.error(request, error_msg)
                            return render(request, 'contractor_management/report_form.html', {
                                'form': form
                            })
                     except Exception as e:
                         logger.error(f"Error getting shift for date {report.report_date}: {e}")
                         # Handle error appropriately, maybe show a generic error message

                # بررسی وجود گزارش قبلی برای این خودرو در این تاریخ و شیفت (با تاریخ میلادی صحیح)
                if report.report_date:
                    existing_report = Report.objects.filter(
                        vehicle=report.vehicle,
                        report_date=report.report_date, # Use the correct Gregorian date
                        shift=report.shift,
                        user=request.user
                    ).first()
                    
                    if existing_report:
                        logger.warning(f"گزارش قبلی برای این خودرو در تاریخ {report.report_date} و شیفت {report.shift} توسط همین کاربر وجود دارد")
                        # نمایش تاریخ شمسی در پیام خطا
                        jdate_display = jdatetime.date.fromgregorian(date=report.report_date).strftime('%Y/%m/%d')
                        error_msg = f'شما قبلاً برای خودرو {report.vehicle} در تاریخ {jdate_display} و شیفت {report.shift} گزارش ثبت کرده‌اید. لطفاً از تاریخ یا شیفت دیگری استفاده کنید.'
                        
                        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                            return JsonResponse({
                                'success': False,
                                'message': error_msg
                            }, status=400)
                        
                        messages.error(request, error_msg)
                        return render(request, 'contractor_management/report_form.html', {
                            'form': form
                        })
                
                logger.info(f"Saving Report - Report Date before save: {report.report_date}, Type: {type(report.report_date)}")
                report.save()
                logger.info(f"گزارش با موفقیت ذخیره شد. شناسه: {report.id}")
                
                # ثبت فعالیت ایجاد گزارش
                log_user_activity(
                    user=request.user,
                    activity_type='create',
                    description=f'ثبت گزارش جدید برای خودرو {report.vehicle.license_plate} متعلق به پیمانکار {report.contractor.company_name}',
                    related_model='Report',
                    related_object_id=report.id,
                    url=reverse('contractor_management:report_detail', args=[report.id]),
                    request=request
                )
                
                # بررسی AJAX request
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': True,
                        'message': 'گزارش با موفقیت ثبت شد',
                        'redirect': reverse('contractor_management:create_report')
                    })
                
                messages.success(request, 'گزارش با موفقیت ثبت شد.')
                return redirect('contractor_management:create_report')
            except Exception as e:
                error_message = str(e)
                if "UNIQUE constraint failed" in error_message:
                    logger.error(f"خطا در ذخیره گزارش (محدودیت یکتا): {error_message}")
                    error_msg = f'شما قبلاً برای خودرو {report.vehicle} در تاریخ {report.report_date} و شیفت {report.shift} گزارش ثبت کرده‌اید. لطفاً از تاریخ یا شیفت دیگری استفاده کنید.'
                else:
                    logger.error(f"خطای غیرمنتظره در ذخیره گزارش: {error_message}")
                    error_msg = 'خطایی در ثبت گزارش رخ داد. لطفاً دوباره تلاش کنید.'
                
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': False,
                        'message': error_msg
                    }, status=400)
                
                messages.error(request, error_msg)
        else:
            logger.warning(f"خطاهای اعتبارسنجی فرم: {form.errors}")
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors,
                    'message': 'لطفا خطاهای زیر را اصلاح کنید'
                }, status=400)
            
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"خطا در فیلد {form[field].label} : {error}")
    else:
        form = ReportForm()
    return render(request, 'contractor_management/report_form.html', {'form': form})


@permission_required("all_reports")
@login_required
def all_reports(request):
    """نمایش همه گزارش‌های کارکرد خودروها"""
    logger.info("Entering all_reports view")
    
    reports = Report.objects.all().order_by('-report_datetime')
    vehicles = Vehicle.objects.all()
    
    # محاسبه وضعیت هر خودرو بر اساس آخرین گزارش
    vehicle_statuses = {}
    for vehicle in vehicles:
        latest_report = Report.objects.filter(vehicle=vehicle).order_by('-report_datetime', '-id').first()
        
        # تعیین وضعیت خودرو بر اساس آخرین گزارش
        vehicle_status, status_class = get_vehicle_status_display(latest_report)
        
        vehicle_statuses[vehicle.id] = {
            'status': vehicle_status,
            'status_class': status_class,
            'latest_report': latest_report
        }
    
    # ثبت فعالیت مشاهده لیست گزارش‌ها
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده لیست گزارش‌های پیمانکاران',
        related_model='Report',
        related_object_id=None,
        url=reverse('contractor_management:all_reports'),
        request=request
    )
    
    form = ReportFilterForm(request.GET or None)
    query = Q()
    
    if request.method == 'GET' and any(request.GET.values()):
        logger.info(f"Applying filters from GET request: {request.GET}")
        if form.is_valid():
            start_date = form.cleaned_data.get('start_date')
            end_date = form.cleaned_data.get('end_date')
            contractor = form.cleaned_data.get('contractor')
            vehicle = form.cleaned_data.get('vehicle')
            shift = form.cleaned_data.get('shift')
            group = form.cleaned_data.get('group')
            
            if start_date:
                try:
                    # تبدیل تاریخ شمسی به میلادی
                    logger.debug(f"Original start_date string: {start_date}")
                    # تبدیل اعداد فارسی به انگلیسی
                    start_date_en = persian_to_english_numbers(start_date.strip())
                    # حذف کاراکترهای اضافی
                    start_date_en = re.sub(r'[^0-9/]', '', start_date_en)
                    year, month, day = map(int, start_date_en.split('/'))
                    if year < 100: year += 1400 # تصحیح سال دو رقمی
                    start_date_j = jdatetime.date(year, month, day)
                    start_date_g = start_date_j.togregorian()
                    query &= Q(report_date__gte=start_date_g)
                    logger.info(f"Applied start_date filter: report_date >= {start_date_g}")
                except (ValueError, IndexError, AttributeError) as e:
                    logger.warning(f"Invalid start_date format: {start_date}. Error: {e}")
                    messages.error(request, f'فرمت تاریخ شروع ({start_date}) نامعتبر است. لطفا از فرمت YYYY/MM/DD استفاده کنید.')

            if end_date:
                try:
                    # تبدیل تاریخ شمسی به میلادی
                    logger.debug(f"Original end_date string: {end_date}")
                    end_date_en = persian_to_english_numbers(end_date.strip())
                    end_date_en = re.sub(r'[^0-9/]', '', end_date_en)
                    year, month, day = map(int, end_date_en.split('/'))
                    if year < 100: year += 1400 # تصحیح سال دو رقمی
                    end_date_j = jdatetime.date(year, month, day)
                    end_date_g = end_date_j.togregorian()
                    query &= Q(report_date__lte=end_date_g)
                    logger.info(f"Applied end_date filter: report_date <= {end_date_g}")
                except (ValueError, IndexError, AttributeError) as e:
                    logger.warning(f"Invalid end_date format: {end_date}. Error: {e}")
                    messages.error(request, f'فرمت تاریخ پایان ({end_date}) نامعتبر است. لطفا از فرمت YYYY/MM/DD استفاده کنید.')

            if contractor:
                query &= Q(contractor=contractor)
                vehicles = vehicles.filter(contractor=contractor)
                logger.info(f"Applied contractor filter: {contractor}")

            if vehicle:
                query &= Q(vehicle=vehicle)
                logger.info(f"Applied vehicle filter: {vehicle}")

            if shift:
                query &= Q(shift=shift)
                logger.info(f"Applied shift filter: {shift}")

            if group:
                query &= Q(group=group)
                logger.info(f"Applied group filter: {group}")

            # اگر فیلتری اعمال شده، فعالیت جستجو را هم ثبت کنیم
            if any([start_date, end_date, contractor, vehicle, shift, group]):
                log_user_activity(
                    user=request.user,
                    activity_type='view',
                    description='جستجو در همه گزارش‌های پیمانکاران',
                    related_model='Report',
                    related_object_id=None,
                    url=request.get_full_path(),
                    request=request
                )
        else:
            logger.warning(f"Filter form is invalid: {form.errors}")

    reports = reports.filter(query)
    logger.info(f"Total reports after filtering: {reports.count()}")

    # --- اضافه کردن لاگ برای بررسی مقادیر تاریخ ---
    if reports.exists():
        logger.info("Logging report_date values for the first few reports:")
        for report in reports[:3]: # لاگ ۳ گزارش اول
            logger.info(f"Report ID: {report.id}, Report Date: {report.report_date}, Type: {type(report.report_date)}")
            try:
                # تلاش برای تبدیل و لاگ تاریخ شمسی
                jdate = jdatetime.date.fromgregorian(date=report.report_date)
                logger.info(f"Report ID: {report.id}, Converted Shamsi Date: {jdate.strftime('%Y/%m/%d')}")
            except Exception as e:
                logger.error(f"Error converting date for Report ID {report.id}: {e}")
    else:
        logger.info("No reports found after filtering.")
    # --- پایان بخش لاگ ---

    paginator = Paginator(reports, 10)  # Show 10 reports per page

    page = request.GET.get('page')
    try:
        reports_page = paginator.page(page)
    except PageNotAnInteger:
        reports_page = paginator.page(1)
    except EmptyPage:
        reports_page = paginator.page(paginator.num_pages)
    
    # محاسبه تعداد خودروهای فعال، نیمه فعال و غیرفعال
    active_count = sum(1 for status in vehicle_statuses.values() if status['status'] == 'فعال')
    partial_count = sum(1 for status in vehicle_statuses.values() if status['status'] == 'نیمه فعال')
    inactive_count = sum(1 for status in vehicle_statuses.values() if status['status'] == 'غیرفعال')
    unknown_count = sum(1 for status in vehicle_statuses.values() if status['status'] == 'نامشخص')
    
    context = {
        'reports': reports_page,
        'vehicles': vehicles,
        'vehicle_statuses': vehicle_statuses,
        'form': form,
        'title': 'گزارش‌های کارکرد خودروها',
        'active_count': active_count,
        'partial_count': partial_count,
        'inactive_count': inactive_count,
        'unknown_count': unknown_count
    }
    logger.info("Rendering all_reports.html template.")
    return render(request, 'contractor_management/reports/all_reports.html', context)


def get_contractors_ajax(request):
    # ثبت فعالیت دریافت لیست پیمانکاران
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='دریافت لیست پیمانکاران از طریق AJAX',
        related_model='Contractor',
        related_object_id=None,
        url=None,
        request=request
    )
    
    contractors = []
    search_term = request.GET.get('term', '')
    
    if search_term:
        # ثبت فعالیت جستجوی پیمانکاران
        log_user_activity(
            user=request.user,
            activity_type='view',
            description=f'جستجوی پیمانکار با عبارت "{search_term}"',
            related_model='Contractor',
            related_object_id=None,
            url=None,
            request=request
        )
        
        contractors = Contractor.objects.filter(
            Q(company_name__icontains=search_term)
        ).values("id", "company_name")
    else:
        contractors = Contractor.objects.values("id", "company_name")
    
    return JsonResponse(list(contractors), safe=False)


def get_contractor_employees_ajax(request):
    # ثبت فعالیت دریافت لیست کارکنان پیمانکار
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='دریافت لیست کارکنان پیمانکار از طریق AJAX',
        related_model='Employee',
        related_object_id=None,
        url=None,
        request=request
    )
    
    employees = []
    search_term = request.GET.get('term', '')
    
    if search_term:
        # ثبت فعالیت جستجوی کارکنان پیمانکار
        log_user_activity(
            user=request.user,
            activity_type='view',
            description=f'جستجوی کارمند پیمانکار با عبارت "{search_term}"',
            related_model='Employee',
            related_object_id=None,
            url=None,
            request=request
        )
        
        employees = Employee.objects.filter(
            Q(first_name__icontains=search_term) |
            Q(last_name__icontains=search_term)
        ).values("id", "first_name", "last_name")
    else:
        employees = Employee.objects.values("id", "first_name", "last_name")
    
    formatted_employees = []
    for employee in employees:
        formatted_employees.append({
            "id": employee['id'],
            "name": f"{employee['first_name']} {employee['last_name']}"
        })
    
    return JsonResponse(list(formatted_employees), safe=False)


def get_all_vehicles_ajax(request):
    # ثبت فعالیت دریافت لیست خودروها
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='دریافت لیست تمام خودروهای پیمانکاران از طریق AJAX',
        related_model='Vehicle',
        related_object_id=None,
        url=None,
        request=request
    )
    
    vehicles = Vehicle.objects.all().values('id', 'driver_name', 'license_plate', 'contractor__company_name')
    vehicle_list = [
        {
            'id': v['id'],
            'text': f"{v['driver_name']} ({v['license_plate']}) - {v['contractor__company_name']}"
        } for v in vehicles
    ]
    
    return JsonResponse(vehicle_list, safe=False)


@login_required
def report_list(request):
    # ثبت فعالیت مشاهده لیست گزارش‌ها
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده لیست گزارش‌های پیمانکاران',
        related_model='Report',
        related_object_id=None,
        url=reverse('contractor_management:report_list'),
        request=request
    )
    
    reports = Report.objects.all()
    query = Q()
    
    if request.method == 'POST':
        form = ReportFilterForm(request.POST)
        if form.is_valid():
            start_date = form.cleaned_data.get('start_date')
            end_date = form.cleaned_data.get('end_date')
            contractor = form.cleaned_data.get('contractor')
            vehicle = form.cleaned_data.get('vehicle')
            shift = form.cleaned_data.get('shift')
            group = form.cleaned_data.get('group')
            if start_date and end_date:
                query &= Q(report_datetime__date__range=[start_date, end_date])
            elif start_date:
                query &= Q(report_datetime__date__gte=start_date)
            elif end_date:
                query &= Q(report_datetime__date__lte=end_date)
            if contractor:
                query &= Q(contractor__company_name=contractor)
            if vehicle:
                query &= Q(vehicle__license_plate=vehicle)
            if shift:
                query &= Q(shift=shift)
            if group:
                query &= Q(group=group)
                
            # ثبت فعالیت جستجو در گزارش‌ها
            log_user_activity(
                user=request.user,
                activity_type='view',
                description='جستجو در لیست گزارش‌های پیمانکاران',
                related_model='Report',
                related_object_id=None,
                url=request.get_full_path(),
                request=request
            )
    else:
        form = ReportFilterForm()

    reports = reports.filter(query)
    paginator = Paginator(reports, 10)  # Show 10 reports per page

    page = request.GET.get('page')
    try:
        reports = paginator.page(page)
    except PageNotAnInteger:
        reports = paginator.page(1)
    except EmptyPage:
        reports = paginator.page(paginator.num_pages)
    return render(request, 'contractor_management/reports/all_reports.html', {
        'reports': reports,
        'form': form,
    })


@login_required
def report_detail(request, pk):
    report = get_object_or_404(Report, pk=pk)
    
    # ثبت فعالیت مشاهده جزئیات گزارش
    log_user_activity(
        user=request.user,
        activity_type='view',
        description=f'مشاهده جزئیات گزارش خودرو {report.vehicle.license_plate}',
        related_model='Report',
        related_object_id=report.id,
        url=reverse('contractor_management:report_detail', args=[pk]),
        request=request
    )
    
    context = {
        'report': report,
        'title': f'جزئیات گزارش کارکرد خودرو {report.vehicle.license_plate}'
    }
    
    return render(request, 'contractor_management/reports/report_detail.html', context)


@login_required
def edit_report(request, pk):
    report = get_object_or_404(Report, pk=pk)
    
    # ثبت فعالیت مشاهده فرم ویرایش گزارش
    if request.method == 'GET':
        log_user_activity(
            user=request.user,
            activity_type='view',
            description=f'مشاهده فرم ویرایش گزارش خودرو {report.vehicle.license_plate}',
            related_model='Report',
            related_object_id=report.id,
            url=reverse('contractor_management:edit_report', args=[pk]),
            request=request
        )
    
    if request.method == 'POST':
        form = ReportForm(request.POST, instance=report)
        if form.is_valid():
            form.save()
            
            # ثبت فعالیت ویرایش گزارش
            log_user_activity(
                user=request.user,
                activity_type='update',
                description=f'ویرایش گزارش خودرو {report.vehicle.license_plate}',
                related_model='Report',
                related_object_id=report.id,
                url=reverse('contractor_management:report_detail', args=[pk]),
                request=request
            )
            
            messages.success(request, 'گزارش با موفقیت ویرایش شد.')
            return redirect('contractor_management:report_detail', pk=pk)
    else:
        form = ReportForm(instance=report)
        
    return render(request, 'contractor_management/edit_report.html', {'form': form, 'report': report})


@login_required
def delete_report(request, pk):
    report = get_object_or_404(Report, pk=pk)
    
    # ثبت فعالیت مشاهده صفحه حذف گزارش
    if request.method == 'GET':
        log_user_activity(
            user=request.user,
            activity_type='view',
            description=f'مشاهده صفحه حذف گزارش خودرو {report.vehicle.license_plate}',
            related_model='Report',
            related_object_id=report.id,
            url=reverse('contractor_management:delete_report', args=[pk]),
            request=request
        )
    
    if request.method == 'POST':
        vehicle_info = f'{report.vehicle.license_plate}'
        report_id = report.id
        report.delete()
        
        # ثبت فعالیت حذف گزارش
        log_user_activity(
            user=request.user,
            activity_type='delete',
            description=f'حذف گزارش خودرو {vehicle_info}',
            related_model='Report',
            related_object_id=report_id,
            url=None,
            request=request
        )
        
        messages.success(request, 'گزارش با موفقیت حذف شد.')
        return redirect('contractor_management:report_list')
    
    return render(request, 'contractor_management/delete_report.html', {'report': report})


@login_required
def contractor_list(request):
    # ثبت فعالیت مشاهده لیست پیمانکاران
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده لیست پیمانکاران',
        related_model='Contractor',
        related_object_id=None,
        url=reverse('contractor_management:contractor_list'),
        request=request
    )
    
    contractors = Contractor.objects.all()
    return render(request, 'contractor_management/contractor_list.html', {'contractors': contractors})


@login_required
def contractor_detail(request, pk):
    contractor = get_object_or_404(Contractor, pk=pk)
    
    # ثبت فعالیت مشاهده جزئیات پیمانکار
    log_user_activity(
        user=request.user,
        activity_type='view',
        description=f'مشاهده جزئیات پیمانکار {contractor.company_name}',
        related_model='Contractor',
        related_object_id=contractor.id,
        url=reverse('contractor_management:contractor_detail', args=[pk]),
        request=request
    )
    
    return render(request, 'contractor_management/contractor_detail.html', {'contractor': contractor})


@login_required
def employee_list(request):
    # ثبت فعالیت مشاهده لیست کارکنان
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده لیست کارکنان پیمانکاران',
        related_model='Employee',
        related_object_id=None,
        url=reverse('contractor_management:employee_list'),
        request=request
    )
    
    employees = Employee.objects.all()
    return render(request, 'contractor_management/employee_list.html', {'employees': employees})


@login_required
def employee_detail(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    
    # ثبت فعالیت مشاهده جزئیات کارمند
    log_user_activity(
        user=request.user,
        activity_type='view',
        description=f'مشاهده جزئیات کارمند {employee.first_name} {employee.last_name}',
        related_model='Employee',
        related_object_id=employee.id,
        url=reverse('contractor_management:employee_detail', args=[pk]),
        request=request
    )
    
    return render(request, 'contractor_management/employee_detail.html', {'employee': employee})


@login_required
def vehicle_list(request):
    # ثبت فعالیت مشاهده لیست خودروها
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده لیست خودروهای پیمانکاران',
        related_model='Vehicle',
        related_object_id=None,
        url=reverse('contractor_management:vehicle_list'),
        request=request
    )
    
    vehicles = Vehicle.objects.all()
    return render(request, 'contractor_management/vehicle_list.html', {'vehicles': vehicles})


@login_required
def vehicle_detail(request, pk):
    vehicle = get_object_or_404(Vehicle, pk=pk)
    
    # بررسی آخرین گزارش خودرو برای تعیین وضعیت فعال/غیرفعال
    latest_report = Report.objects.filter(vehicle=vehicle).order_by('-report_datetime').first()
    
    # تعیین وضعیت خودرو
    if latest_report:
        if latest_report.status == 'inactive':
            vehicle_status = 'غیرفعال'
            status_class = 'danger'
        elif latest_report.status == 'partial':
            vehicle_status = 'نیمه فعال'
            status_class = 'warning'
        else:
            vehicle_status = 'فعال'
            status_class = 'success'
        
        # محاسبه زمان آخرین گزارش
        last_report_time = latest_report.report_datetime
    else:
        vehicle_status = 'نامشخص'
        status_class = 'secondary'
        last_report_time = None
    
    # ثبت فعالیت مشاهده جزئیات خودرو
    log_user_activity(
        user=request.user,
        activity_type='view',
        description=f'مشاهده جزئیات خودرو {vehicle.license_plate}',
        related_model='Vehicle',
        related_object_id=vehicle.id,
        url=reverse('contractor_management:vehicle_detail', args=[pk]),
        request=request
    )
    
    # گزارش‌های اخیر خودرو (۵ گزارش آخر)
    recent_reports = Report.objects.filter(vehicle=vehicle).order_by('-report_datetime')[:5]
    
    context = {
        'vehicle': vehicle,
        'vehicle_status': vehicle_status,
        'status_class': status_class,
        'last_report_time': last_report_time,
        'recent_reports': recent_reports,
        'title': f'جزئیات خودرو {vehicle.license_plate}'
    }
    
    return render(request, 'contractor_management/vehicle_detail.html', context)


@login_required
def vehicle_reports_pdf(request, vehicle_id):
    """خروجی PDF گزارش‌های خودرو با WeasyPrint"""
    from weasyprint import HTML, CSS
    from django.template.loader import render_to_string
    from core.models import SiteSettings
    import jdatetime
    
    vehicle = get_object_or_404(Vehicle, pk=vehicle_id)
    
    # دریافت گزارش‌ها با فیلترها
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    shift = request.GET.get('shift')
    group = request.GET.get('group')
    
    reports_query = Report.objects.filter(vehicle=vehicle).select_related('user').order_by('-report_date', '-report_datetime')
    reports_query = apply_date_filters(reports_query, start_date, end_date)
    
    if shift:
        reports_query = reports_query.filter(shift=shift)
    if group:
        reports_query = reports_query.filter(group=group)
    
    # دریافت تنظیمات سایت
    site_settings = SiteSettings.objects.first()
    
    # تاریخ جاری به شمسی
    current_date_jalali = jdatetime.datetime.now().strftime('%Y/%m/%d - %H:%M')
    
    # Render HTML template
    html_string = render_to_string('contractor_management/reports/vehicle_reports_pdf.html', {
        'vehicle': vehicle,
        'reports': reports_query,
        'site_settings': site_settings,
        'current_date': current_date_jalali,
        'request': request
    })
    
    # Generate PDF
    html = HTML(string=html_string, base_url=request.build_absolute_uri())
    pdf = html.write_pdf()
    
    # Return PDF response
    response = HttpResponse(pdf, content_type='application/pdf')
    filename = f'vehicle_reports_{vehicle.license_plate}_{timezone.now().strftime("%Y%m%d")}.pdf'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response


@login_required
def vehicle_reports(request, vehicle_id):
    """نمایش گزارش‌های کارکرد یک خودروی خاص"""
    
    vehicle = get_object_or_404(Vehicle, pk=vehicle_id)
    
    # محاسبه وضعیت خودرو بر اساس آخرین گزارش
    latest_report = Report.objects.filter(vehicle=vehicle).order_by('-report_datetime', '-id').first()
    
    # تعیین وضعیت خودرو
    vehicle_status, status_class, last_report_time = get_vehicle_status(latest_report)
    
    # دریافت گزارش‌های فیلتر شده
    reports = get_filtered_reports(request, vehicle)
    
    # فرم فیلتر
    form = ReportFilterForm(request.GET or None)
    
    # ثبت فعالیت مشاهده گزارش‌های خودرو
    log_user_activity(
        user=request.user,
        activity_type='view',
        description=f'مشاهده گزارش‌های کارکرد خودرو {vehicle.license_plate}',
        related_model='Vehicle',
        related_object_id=vehicle.id,
        url=request.get_full_path(),
        request=request
    )
    
    context = {
        'vehicle': vehicle,
        'vehicle_status': vehicle_status,
        'status_class': status_class,
        'last_report_time': last_report_time,
        'latest_report': latest_report,
        'reports': reports,
        'form': form,
        'title': f'گزارش‌های کارکرد خودرو {vehicle.license_plate}'
    }
    
    return render(request, 'contractor_management/reports/vehicle_reports.html', context)


def get_vehicle_status(latest_report):
    """محاسبه وضعیت خودرو بر اساس آخرین گزارش"""
    if latest_report:
        if latest_report.status == 'inactive':
            vehicle_status = 'غیرفعال'
            status_class = 'danger'
        elif latest_report.status == 'partial':
            vehicle_status = 'نیمه فعال'
            status_class = 'warning'
        elif latest_report.status == 'full':
            vehicle_status = 'فعال'
            status_class = 'success'
        else:
            # حالت پیش‌فرض اگر مقدار status شناخته شده نباشد
            vehicle_status = f'نامشخص ({latest_report.status})'
            status_class = 'secondary'
        
        last_report_time = latest_report.report_datetime
    else:
        vehicle_status = 'نامشخص'
        status_class = 'secondary'
        last_report_time = None
    
    return vehicle_status, status_class, last_report_time


def get_filtered_reports(request, vehicle):
    """دریافت گزارش‌های فیلتر شده"""
    # فیلترهای جستجو
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    shift = request.GET.get('shift')
    group = request.GET.get('group')
    
    # پایه کوئری
    reports_query = Report.objects.filter(vehicle=vehicle).order_by('-report_datetime')
    
    # اعمال فیلترها
    reports_query = apply_date_filters(reports_query, start_date, end_date)
    
    if shift:
        reports_query = reports_query.filter(shift=shift)
    
    if group:
        reports_query = reports_query.filter(group=group)
    
    # صفحه‌بندی
    paginator = Paginator(reports_query, 10)  # 10 گزارش در هر صفحه
    page = request.GET.get('page')
    
    try:
        reports = paginator.page(page)
    except PageNotAnInteger:
        reports = paginator.page(1)
    except EmptyPage:
        reports = paginator.page(paginator.num_pages)
    
    return reports


def apply_date_filters(reports_query, start_date, end_date):
    """اعمال فیلترهای تاریخ به کوئری"""
    logger.debug(f"Applying date filters - Start: {start_date}, End: {end_date}")
    if start_date:
        try:
            # تبدیل تاریخ شمسی به میلادی
            start_date_en = persian_to_english_numbers(start_date.strip())
            start_date_en = re.sub(r'[^0-9/]', '', start_date_en)
            year, month, day = map(int, start_date_en.split('/'))
            if year < 100: year += 1400
            start_date_j = jdatetime.date(year, month, day)
            start_date_g = start_date_j.togregorian()
            # فیلتر بر اساس report_date (برای گزارش‌های جدید) یا report_datetime (برای گزارش‌های قدیمی)
            reports_query = reports_query.filter(
                Q(report_date__gte=start_date_g) | 
                Q(report_date__isnull=True, report_datetime__date__gte=start_date_g)
            )
            logger.info(f"Applied start_date filter: report_date >= {start_date_g}")
        except (ValueError, IndexError, AttributeError) as e:
            logger.warning(f"Error converting start_date: {start_date}. Error: {e}")
            pass

    if end_date:
        try:
            # تبدیل تاریخ شمسی به میلادی
            end_date_en = persian_to_english_numbers(end_date.strip())
            end_date_en = re.sub(r'[^0-9/]', '', end_date_en)
            year, month, day = map(int, end_date_en.split('/'))
            if year < 100: year += 1400
            end_date_j = jdatetime.date(year, month, day)
            end_date_g = end_date_j.togregorian()
            # فیلتر بر اساس report_date (برای گزارش‌های جدید) یا report_datetime (برای گزارش‌های قدیمی)
            reports_query = reports_query.filter(
                Q(report_date__lte=end_date_g) | 
                Q(report_date__isnull=True, report_datetime__date__lte=end_date_g)
            )
            logger.info(f"Applied end_date filter: report_date <= {end_date_g}")
        except (ValueError, IndexError, AttributeError) as e:
            logger.warning(f"Error converting end_date: {end_date}. Error: {e}")
            pass

    return reports_query


def get_vehicle_status_display(latest_report):
    """
    تعیین وضعیت خودرو و کلاس وضعیت (badge) بر اساس آخرین گزارش.
    """
    if latest_report:
        if latest_report.status == 'inactive':
            vehicle_status = 'غیرفعال'
            status_class = 'danger'
        elif latest_report.status == 'partial':
            vehicle_status = 'نیمه فعال'
            status_class = 'warning'
        elif latest_report.status == 'full':
            vehicle_status = 'فعال'
            status_class = 'success'
        else:
            vehicle_status = 'نامشخص'
            status_class = 'secondary'
    else:
        vehicle_status = 'نامشخص'
        status_class = 'secondary'
    
    return vehicle_status, status_class


def export_reports_to_excel(request, vehicle_id=None):
    """خروجی اکسل گزارش‌ها (لیست کلی و فیلترشده) با سربرگ نام و آرم شرکت"""
    # Lazy imports for optional deps
    from core.models import SiteSettings
    try:
        from openpyxl.drawing.image import Image as XLImage
    except Exception:
        XLImage = None

    # Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش‌های کارکرد خودرو"

    # RTL
    ws.sheet_properties.rightToLeft = True

    # Base queryset: all reports
    reports_query = Report.objects.select_related('user', 'vehicle', 'contractor').all()

    # Vehicle-specific filter (single vehicle export)
    vehicle = None
    if vehicle_id:
        reports_query = reports_query.filter(vehicle_id=vehicle_id)
        vehicle = Vehicle.objects.get(id=vehicle_id)

    # Filters from querystring (same as list page)
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    contractor_id = request.GET.get('contractor')
    vehicle_filter_id = request.GET.get('vehicle')
    shift = request.GET.get('shift')
    group = request.GET.get('group')

    if start_date or end_date:
        reports_query = apply_date_filters(reports_query, start_date, end_date)
    if contractor_id:
        reports_query = reports_query.filter(contractor_id=contractor_id)
    if vehicle_filter_id:
        reports_query = reports_query.filter(vehicle_id=vehicle_filter_id)
    if shift:
        reports_query = reports_query.filter(shift=shift)
    if group:
        reports_query = reports_query.filter(group=group)

    reports_query = reports_query.order_by('-report_date', '-report_datetime')

    # Styles
    title_font = Font(name='B Nazanin', size=20, bold=True)
    subtitle_font = Font(name='B Nazanin', size=14, bold=True)
    info_font = Font(name='B Nazanin', size=12)
    header_font = Font(name='B Nazanin', size=12, bold=True)
    data_font = Font(name='B Nazanin', size=11)

    header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    title_fill = PatternFill(start_color='E8E8E8', end_color='E8E8E8', fill_type='solid')
    info_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')

    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Site settings (logo + name)
    site = SiteSettings.objects.first()
    site_name = site.site_name if site and site.site_name else 'شرکت'
    title_text = f"گزارش‌های کارکرد خودرو - {site_name}"

    # Title row (merge across A:K)
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = center_alignment
    ws.merge_cells('A1:K1')
    ws.row_dimensions[1].height = 45

    # Logo (if available)
    if site and getattr(site, 'company_logo', None) and hasattr(site.company_logo, 'path') and XLImage:
        try:
            logo_img = XLImage(site.company_logo.path)
            logo_img.height = 60
            # Anchor at A1 (top-right in RTL appearance)
            ws.add_image(logo_img, 'A1')
        except Exception:
            pass

    # If single-vehicle export, add info rows
    current_row = 2
    if vehicle:
        info_cell1 = ws.cell(row=current_row, column=1, value=f'نام پیمانکار: {vehicle.contractor.company_name}')
        info_cell1.font = subtitle_font
        info_cell1.fill = info_fill
        info_cell1.alignment = center_alignment
        ws.merge_cells(f'A{current_row}:K{current_row}')
        ws.row_dimensions[current_row].height = 30
        current_row += 1

        cell_vtype = ws.cell(row=current_row, column=1, value=f'نوع خودرو: {vehicle.vehicle_type}')
        cell_vtype.font = info_font; cell_vtype.fill = info_fill; cell_vtype.alignment = center_alignment
        ws.merge_cells(f'A{current_row}:D{current_row}')
        cell_plate = ws.cell(row=current_row, column=5, value=f'پلاک: {vehicle.license_plate}')
        cell_plate.font = info_font; cell_plate.fill = info_fill; cell_plate.alignment = center_alignment
        ws.merge_cells(f'E{current_row}:G{current_row}')
        cell_phone = ws.cell(row=current_row, column=8, value=f'شماره تماس مدیر: {vehicle.contractor.manager_phone}')
        cell_phone.font = info_font; cell_phone.fill = info_fill; cell_phone.alignment = center_alignment
        ws.merge_cells(f'H{current_row}:K{current_row}')
        ws.row_dimensions[current_row].height = 28
        current_row += 1

    # Headers (include vehicle/contractor and report_date)
    headers = [
        'شماره گزارش',
        'خودرو / پیمانکار',
        'تاریخ گزارش',
        'تاریخ و زمان ثبت',
        'ثبت کننده',
        'شیفت کاری',
        'گروه کاری',
        'وضعیت کارکرد',
        'ساعت شروع توقف',
        'ساعت پایان توقف',
        'توضیحات'
    ]

    header_row = current_row if vehicle else 2
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    ws.row_dimensions[header_row].height = 28

    # Column widths A-K
    column_widths = {
        'A': 12,
        'B': 35,
        'C': 14,
        'D': 22,
        'E': 20,
        'F': 15,
        'G': 10,
        'H': 16,
        'I': 16,
        'J': 16,
        'K': 45,
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # Data rows
    data_start_row = header_row + 1
    for row_idx, report in enumerate(reports_query, data_start_row):
        ws.row_dimensions[row_idx].height = 22
        # Vehicle cell content
        vehicle_text = f"{report.vehicle.driver_name} ({report.vehicle.license_plate}) - {report.contractor.company_name}"
        # Jalali dates
        report_date_j = ''
        if report.report_date:
            try:
                report_date_j = jdatetime.date.fromgregorian(date=report.report_date).strftime('%Y/%m/%d')
            except Exception:
                report_date_j = str(report.report_date)
        report_dt_j = ''
        if report.report_datetime:
            try:
                report_dt_j = jdatetime.datetime.fromgregorian(datetime=report.report_datetime).strftime('%Y/%m/%d %H:%M')
            except Exception:
                report_dt_j = report.report_datetime.strftime('%Y/%m/%d %H:%M')

        ws.cell(row=row_idx, column=1, value=report.id)
        ws.cell(row=row_idx, column=2, value=vehicle_text)
        ws.cell(row=row_idx, column=3, value=report_date_j)
        ws.cell(row=row_idx, column=4, value=report_dt_j)
        ws.cell(row=row_idx, column=5, value=report.user.get_full_name() or report.user.username)
        ws.cell(row=row_idx, column=6, value=report.shift or '')
        ws.cell(row=row_idx, column=7, value=report.group or '')
        ws.cell(row=row_idx, column=8, value=dict(Report.STATUS_CHOICES).get(report.status, report.status))
        ws.cell(row=row_idx, column=9, value=str(report.stop_start_time) if report.stop_start_time else '')
        ws.cell(row=row_idx, column=10, value=str(report.stop_end_time) if report.stop_end_time else '')
        ws.cell(row=row_idx, column=11, value=report.description or '')

        for c in range(1, 12):
            cell = ws.cell(row=row_idx, column=c)
            cell.font = data_font
            cell.alignment = center_alignment
            thin_border = Side(border_style="thin", color="000000")
            cell.border = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)

    # Borders for title and header
    thin_border = Side(border_style="thin", color="000000")
    border_style = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)
    ws['A1'].border = border_style
    if vehicle:
        ws[f'A{header_row-2}'].border = border_style  # contractor row
        ws[f'A{header_row-1}'].border = border_style  # vehicle info row
    for c in range(1, 12):
        ws.cell(row=header_row, column=c).border = border_style

    # Filename
    filename = f"reports_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    if vehicle:
        filename = f"reports_{vehicle.license_plate}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    # Response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)

    # Log export
    log_user_activity(
        user=request.user,
        activity_type='export',
        description='دریافت خروجی اکسل گزارش‌ها',
        related_model='Report',
        related_object_id=None,
        url=request.get_full_path(),
        request=request
    )
    return response


# =====================================
# New Views for Contractor Management
# =====================================

@permission_required("contractor_dashboard")
@login_required
def manage_contractor_user(request, pk):
    contractor = get_object_or_404(Contractor, pk=pk)
    form = ContractorUserManagementForm(request.POST or None)

    if request.method == 'POST':
        action = request.POST.get('action')
        from django.contrib.auth.models import User
        if action == 'create':
            username = form.cleaned_data.get('username') if form.is_valid() else None
            if not username:
                messages.error(request, 'نام کاربری الزامی است')
                return render(request, 'contractor_management/auth/manage_contractor_user.html', {'contractor': contractor, 'form': form})
            if User.objects.filter(username=username).exists():
                messages.error(request, 'این نام کاربری قبلاً گرفته شده است')
                return render(request, 'contractor_management/auth/manage_contractor_user.html', {'contractor': contractor, 'form': form})
            user = User.objects.create_user(username=username)
            user.set_password(form.cleaned_data['password'])
            user.first_name = contractor.manager_name.split(' ')[0] if contractor.manager_name else ''
            user.last_name = 'Contractor'
            user.save()
            contractor.user = user
            contractor.save(update_fields=['user'])
            messages.success(request, f'حساب کاربری ایجاد و لینک شد: {user.username}')
            return redirect('contractor_management:data_management')

        elif action == 'link' and form.is_valid():
            existing_user = form.cleaned_data.get('existing_user')
            if not existing_user:
                messages.error(request, 'کاربر موجود را انتخاب کنید')
                return render(request, 'contractor_management/auth/manage_contractor_user.html', {'contractor': contractor, 'form': form})
            if hasattr(existing_user, 'contractor_profile') and existing_user.contractor_profile:
                messages.error(request, 'این کاربر قبلاً به پیمانکار دیگری لینک شده است')
                return render(request, 'contractor_management/auth/manage_contractor_user.html', {'contractor': contractor, 'form': form})
            contractor.user = existing_user
            contractor.save(update_fields=['user'])
            messages.success(request, f'کاربر {existing_user.username} لینک شد')
            return redirect('contractor_management:data_management')

        elif action == 'change_password' and contractor.user and form.is_valid():
            contractor.user.set_password(form.cleaned_data['password'])
            contractor.user.save()
            messages.success(request, 'رمز عبور با موفقیت تغییر کرد')
            return redirect('contractor_management:data_management')

        elif action == 'unlink':
            contractor.user = None
            contractor.save(update_fields=['user'])
            messages.success(request, 'لینک کاربر از پیمانکار حذف شد')
            return redirect('contractor_management:data_management')

    return render(request, 'contractor_management/auth/manage_contractor_user.html', {
        'contractor': contractor,
        'form': form,
    })


@permission_required("contractor_dashboard")
@login_required
def reset_employee_password(request, pk):
    if request.method != 'POST':
        return HttpResponseForbidden('Method not allowed')
    employee = get_object_or_404(Employee, pk=pk)
    from django.contrib.auth.models import User
    import secrets
    if employee.user:
        user = employee.user
    else:
        # اگر حساب کاربری ندارد، ایجاد نمی‌کنیم (فقط ریست)
        messages.error(request, 'برای این پرسنل حساب کاربری وجود ندارد')
        return redirect('contractor_management:data_management')
    new_password = secrets.token_urlsafe(8)
    user.set_password(new_password)
    user.save()
    messages.success(request, f'رمز عبور کاربر {user.username} بازنشانی شد. رمز موقت: {new_password}')
    return redirect('contractor_management:data_management')

@login_required
def manage_employees(request):
    """مدیریت پرسنل توسط مدیر پیمانکار: مشاهده لیست و افزودن پرسنل جدید"""
    # فقط کاربرانی که پروفایل پیمانکاری دارند
    if not hasattr(request.user, 'contractor_profile'):
        return HttpResponseForbidden('دسترسی غیرمجاز')

    contractor = request.user.contractor_profile
    employees = Employee.objects.filter(contractor=contractor).order_by('first_name', 'last_name')
    temp_password = None
    open_modal = False

    if request.method == 'POST':
        form = ContractorEmployeeForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data
            # ساخت کاربر
            from django.contrib.auth.models import User
            import secrets
            base_username = data['national_id']
            username = base_username
            i = 1
            while User.objects.filter(username=username).exists():
                username = f"{base_username}{i}"
                i += 1
            temp_password = secrets.token_urlsafe(8)
            user = User.objects.create_user(
                username=username,
                password=temp_password,
                first_name=data.get('first_name',''),
                last_name=data.get('last_name','')
            )
            # ساخت پرسنل
            employee = form.save(commit=False)
            employee.contractor = contractor
            employee.user = user
            employee.save()
            messages.success(request, f'پرسنل جدید ایجاد شد. نام کاربری: {username} | رمز موقت: {temp_password}')
            return redirect('contractor_management:manage_employees')
        else:
            messages.error(request, 'فرم نامعتبر است. لطفاً خطاها را بررسی کنید.')
            open_modal = True
    else:
        form = ContractorEmployeeForm()

    return render(request, 'contractor_management/auth/manage_employees.html', {
        'employees': employees,
        'form': form,
        'temp_password': temp_password,
        'contractor': contractor,
        'open_modal': open_modal,
    })


@login_required
def contractor_employee_edit(request, pk):
    if not hasattr(request.user, 'contractor_profile'):
        return HttpResponseForbidden('دسترسی غیرمجاز')
    contractor = request.user.contractor_profile
    employee = get_object_or_404(Employee, pk=pk, contractor=contractor)

    if request.method == 'POST':
        form = ContractorEmployeeForm(request.POST, instance=employee)
        if form.is_valid():
            form.save()
            messages.success(request, 'پرسنل با موفقیت ویرایش شد')
            return redirect('contractor_management:manage_employees')
    else:
        form = ContractorEmployeeForm(instance=employee)

    return render(request, 'contractor_management/auth/contractor_employee_edit.html', {
        'form': form,
        'employee': employee,
        'contractor': contractor,
    })


@login_required
def contractor_employee_delete(request, pk):
    if not hasattr(request.user, 'contractor_profile'):
        return HttpResponseForbidden('دسترسی غیرمجاز')
    contractor = request.user.contractor_profile
    employee = get_object_or_404(Employee, pk=pk, contractor=contractor)

    if request.method == 'POST':
        name = f"{employee.first_name} {employee.last_name}"
        employee.delete()
        messages.success(request, f'پرسنل {name} حذف شد')
        return redirect('contractor_management:manage_employees')

    return render(request, 'contractor_management/confirm_delete.html', {
        'object': employee,
        'object_type': 'پرسنل',
        'object_name': f"{employee.first_name} {employee.last_name}",
        'cancel_url': reverse('contractor_management:manage_employees'),
    })


@login_required
def contractor_employee_reset_password(request, pk):
    if not hasattr(request.user, 'contractor_profile'):
        return HttpResponseForbidden('دسترسی غیرمجاز')
    contractor = request.user.contractor_profile
    employee = get_object_or_404(Employee, pk=pk, contractor=contractor)

    if request.method != 'POST':
        return HttpResponseForbidden('Method not allowed')

    from django.contrib.auth.models import User
    import secrets

    # Ensure user exists
    if employee.user:
        user = employee.user
    else:
        base_username = employee.national_id or f"emp{employee.pk}"
        username = base_username
        i = 1
        while User.objects.filter(username=username).exists():
            username = f"{base_username}{i}"
            i += 1
        user = User.objects.create_user(
            username=username,
            first_name=employee.first_name or '',
            last_name=employee.last_name or ''
        )
        employee.user = user
        employee.save(update_fields=['user'])

    temp_password = secrets.token_urlsafe(8)
    user.set_password(temp_password)
    user.save()

    messages.success(request, f"رمز عبور بازنشانی شد. نام کاربری: {user.username} | رمز موقت: {temp_password}")
    return redirect('contractor_management:manage_employees')


@login_required
def contractor_employee_credentials(request, pk):
    if not hasattr(request.user, 'contractor_profile'):
        return HttpResponseForbidden('دسترسی غیرمجاز')
    contractor = request.user.contractor_profile
    employee = get_object_or_404(Employee, pk=pk, contractor=contractor)
    # Generate new temp password and show printable page
    from django.contrib.auth.models import User
    import secrets
    user = employee.user
    if not user:
        # create account if missing
        base_username = employee.national_id or f"emp{employee.pk}"
        username = base_username
        i = 1
        while User.objects.filter(username=username).exists():
            username = f"{base_username}{i}"
            i += 1
        user = User.objects.create_user(username=username, first_name=employee.first_name or '', last_name=employee.last_name or '')
        employee.user = user
        employee.save(update_fields=['user'])
    temp_password = secrets.token_urlsafe(8)
    user.set_password(temp_password)
    user.save()
    context = {
        'contractor': contractor,
        'employee': employee,
        'username': user.username,
        'password': temp_password,
    }
    return render(request, 'contractor_management/auth/employee_credentials.html', context)

@permission_required("contractor_dashboard")
@login_required
def data_management(request):
    """صفحه مدیریت اطلاعات پایه (پیمانکاران، کارکنان، خودروها)"""
    
    # دریافت همه اطلاعات
    contractors = Contractor.objects.all().order_by('company_name')
    employees = Employee.objects.select_related('contractor').all().order_by('first_name', 'last_name')
    vehicles = Vehicle.objects.select_related('contractor').all().order_by('license_plate')
    
    # ثبت فعالیت مشاهده صفحه مدیریت اطلاعات پایه
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده صفحه مدیریت اطلاعات پایه (پیمانکاران، کارکنان، خودروها)',
        related_model='Contractor',
        related_object_id=None,
        url=reverse('contractor_management:data_management'),
        request=request
    )
    
    context = {
        'contractors': contractors,
        'employees': employees,
        'vehicles': vehicles,
    }
    
    return render(request, 'contractor_management/data_management.html', context)


@permission_required("contractor_dashboard")
@login_required
def contractor_dashboard(request):
    """داشبورد اصلی مدیریت پیمانکاران"""
    today = datetime.datetime.now().date()
    thirty_days_later = today + timedelta(days=30)
    
    # محاسبه آمار کلی
    total_contractors = Contractor.objects.count()
    active_employees = Employee.objects.count()
    active_vehicles = Vehicle.objects.count()
    
    # یافتن پیمانکاران با مدارک منقضی شده
    expired_contractors = []
    for contractor in Contractor.objects.all():
        expired_docs = []
        
        # بررسی بیمه مسئولیت
        if contractor.liability_insurance_expiry:
            if contractor.liability_insurance_expiry <= thirty_days_later:
                status = 'expired' if contractor.liability_insurance_expiry < today else 'warning'
                expired_docs.append({
                    'name': 'بیمه مسئولیت مدنی',
                    'expiry_date': contractor.liability_insurance_expiry,
                    'status': status
                })
        
        # بررسی بیمه آتش‌سوزی
        if contractor.fire_insurance_expiry:
            if contractor.fire_insurance_expiry <= thirty_days_later:
                status = 'expired' if contractor.fire_insurance_expiry < today else 'warning'
                expired_docs.append({
                    'name': 'بیمه آتش‌سوزی',
                    'expiry_date': contractor.fire_insurance_expiry,
                    'status': status
                })
        
        if expired_docs:
            expired_contractors.append({
                'id': contractor.id,
                'company_name': contractor.company_name,
                'expired_documents': expired_docs
            })
    
    # محاسبه تعداد خودروها بر اساس دسته‌بندی
    vehicle_counts = Vehicle.objects.values('vehicle_category').annotate(count=Count('id'))
    mining_vehicles = next((item['count'] for item in vehicle_counts if item['vehicle_category'] == 'mining'), 0)
    light_vehicles = next((item['count'] for item in vehicle_counts if item['vehicle_category'] == 'light'), 0)
    transportation_vehicles = next((item['count'] for item in vehicle_counts if item['vehicle_category'] == 'transportation'), 0)
    
    context = {
        'total_contractors': total_contractors,
        'active_employees': active_employees,
        'active_vehicles': active_vehicles,
        'expired_documents': len(expired_contractors),
        'contractors_with_expired_docs': expired_contractors,
        'mining_vehicles': mining_vehicles,
        'light_vehicles': light_vehicles,
        'transportation_vehicles': transportation_vehicles,
    }
    
    return render(request, 'contractor_management/dashboard/contractor_dashboard.html', context)


def contractor_login(request):
    """صفحه ورود اختصاصی پیمانکاران و پرسنل"""
    
    # اگر کاربر قبلاً لاگین کرده، هدایت به پورتال
    if request.user.is_authenticated:
        if hasattr(request.user, 'employee_profile') or hasattr(request.user, 'contractor_profile'):
            return redirect('contractor_management:contractor_portal')
    
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        
        user = authenticate(request, username=username, password=password)
        
        if user is not None:
            # جلوگیری از ورود پرسنل اورژانس
            if user.groups.filter(name__in=['EmergencyManager', 'EmergencyDoctor', 'EmergencyNurse']).exists():
                messages.error(request, 'لطفاً از پورتال اورژانس وارد شوید.')
                return render(request, 'contractor_management/auth/contractor_login.html')
            
            # بررسی اینکه کاربر پیمانکار یا کارمند است
            if hasattr(user, 'employee_profile') or hasattr(user, 'contractor_profile'):
                auth_login(request, user)
                messages.success(request, f'خوش آمدید {user.get_full_name()}')
                return redirect('contractor_management:contractor_portal')
            else:
                messages.error(request, 'شما مجوز ورود به این بخش را ندارید.')
        else:
            messages.error(request, 'نام کاربری یا رمز عبور اشتباه است.')
    
    return render(request, 'contractor_management/auth/contractor_login.html')


@login_required
def contractor_portal_dashboard(request):
    """داشبورد پورتال کاربری پیمانکاران"""
    
    # بررسی دسترسی
    is_contractor = hasattr(request.user, 'contractor_profile')
    is_employee = hasattr(request.user, 'employee_profile')
    if not (is_contractor or is_employee):
        messages.error(request, 'شما مجوز دسترسی به این بخش را ندارید.')
        return redirect('accounts:login')
    
    today = date.today()
    context = {
        'today': today,
    }
    
    if is_employee:
        employee = request.user.employee_profile
        # محاسبه درصد تکمیل پروفایل
        profile_completion = 0
        has_personal_info = False
        has_documents = False
        has_training = False
        
        # بررسی اطلاعات شخصی
        if employee.first_name and employee.last_name and employee.national_id and employee.birth_date:
            has_personal_info = True
            profile_completion += 33
        
        # بررسی مدارک
        documents_uploaded = 0
        if employee.card_national_img:
            documents_uploaded += 1
        if employee.certificate_img:
            documents_uploaded += 1
        if employee.health_certificate:
            documents_uploaded += 1
        if employee.background_check:
            documents_uploaded += 1
        
        if documents_uploaded >= 3:  # حداقل 3 مدرک
            has_documents = True
            profile_completion += 33
        
        # بررسی آموزش‌ها
        if employee.safety_training:
            has_training = True
            profile_completion += 34
        
        # لیست اقدامات مورد نیاز
        required_actions = []
        
        if not has_personal_info:
            required_actions.append({
                'title': 'تکمیل اطلاعات شخصی',
                'description': 'لطفاً اطلاعات شخصی خود را کامل کنید',
                'link': reverse('contractor_management:contractor_profile')
            })
        
        if not has_documents:
            required_actions.append({
                'title': 'بارگذاری مدارک',
                'description': 'لطفاً مدارک هویتی و سلامت خود را بارگذاری کنید',
                'link': reverse('contractor_management:contractor_profile') + '#documents'
            })
        
        if not has_training:
            required_actions.append({
                'title': 'گذراندن دوره ایمنی',
                'description': 'لطفاً گواهی دوره آموزش ایمنی عمومی را بارگذاری کنید',
                'link': reverse('contractor_management:contractor_profile') + '#training'
            })
        
        if employee.entry_permit_expiration and employee.entry_permit_expiration <= today + timedelta(days=30):
            try:
                import jdatetime
                exp_j = jdatetime.date.fromgregorian(date=employee.entry_permit_expiration).strftime('%Y/%m/%d')
            except Exception:
                exp_j = str(employee.entry_permit_expiration)
            required_actions.append({
                'title': 'تمدید مجوز ورود',
                'description': f'مجوز ورود شما در تاریخ {exp_j} منقضی می‌شود',
                'link': reverse('contractor_management:contractor_profile')
            })
        
        context.update({
            'profile_completion': profile_completion,
            'profile': {
                'has_personal_info': has_personal_info,
                'has_documents': has_documents,
                'has_training': has_training,
            },
            'required_actions': required_actions,
            'recent_activity': None,
        })
    else:
        # پیمانکار (مدیر)
        context.update({
            'is_contractor': True,
        })
    
    return render(request, 'contractor_management/auth/contractor_portal.html', context)


@login_required
def contractor_profile_documents(request):
    """صفحه مدیریت پروفایل و مدارک کاربر"""
    
    # بررسی دسترسی
    if not hasattr(request.user, 'employee_profile'):
        messages.error(request, 'شما مجوز دسترسی به این بخش را ندارید.')
        return redirect('accounts:login')
    
    from .forms import EmployeeDocumentForm
    employee = request.user.employee_profile
    
    if request.method == 'POST':
        # اگر فایل آپلود شده باشد، فرم مدارک را ذخیره کن
        if request.FILES:
            doc_form = EmployeeDocumentForm(request.POST or None, request.FILES or None, instance=employee)
            if doc_form.is_valid():
                doc_form.save()
                log_user_activity(
                    user=request.user,
                    activity_type='update',
                    description='بروزرسانی مدارک پرسنل',
                    related_model='Employee',
                    related_object_id=employee.id,
                    url=reverse('contractor_management:contractor_profile'),
                    request=request
                )
                messages.success(request, 'مدارک با موفقیت بارگذاری شد.')
                return redirect('contractor_management:contractor_profile')
        
        # بروزرسانی اطلاعات شخصی
        employee.first_name = request.POST.get('first_name', employee.first_name)
        employee.last_name = request.POST.get('last_name', employee.last_name)
        employee.national_id = request.POST.get('national_id', employee.national_id)
        employee.phone_number = request.POST.get('phone_number', employee.phone_number)
        employee.education = request.POST.get('education', employee.education)
        
        # تاریخ تولد (ورودی جلالی)
        birth_date = request.POST.get('birth_date')
        if birth_date:
            try:
                from .forms import persian_to_english_numbers
                import jdatetime
                s = persian_to_english_numbers(birth_date).strip()
                if s:
                    parts = s.replace('/', '-').split('-')
                    if len(parts) == 3:
                        y, m, d = map(int, parts)
                        employee.birth_date = jdatetime.date(y, m, d).togregorian()
            except Exception:
                pass
        
        employee.save()
        
        # ثبت فعالیت
        log_user_activity(
            user=request.user,
            activity_type='update',
            description=f'بروزرسانی پروفایل کارمند',
            related_model='Employee',
            related_object_id=employee.id,
            url=reverse('contractor_management:contractor_profile'),
            request=request
        )
        
        messages.success(request, 'اطلاعات با موفقیت بروزرسانی شد.')
        return redirect('contractor_management:contractor_profile')
    
    # محاسبه درصد تکمیل پروفایل
    profile_completion = 0
    
    # بررسی اطلاعات شخصی
    if employee.first_name and employee.last_name and employee.national_id and employee.birth_date:
        profile_completion += 33
    
    # بررسی مدارک
    documents_uploaded = 0
    if employee.card_national_img:
        documents_uploaded += 1
    if employee.certificate_img:
        documents_uploaded += 1
    if employee.health_certificate:
        documents_uploaded += 1
    if employee.background_check:
        documents_uploaded += 1
    
    if documents_uploaded >= 3:
        profile_completion += 33
    
    # بررسی آموزش‌ها
    if employee.safety_training:
        profile_completion += 34
    
    # ثبت فعالیت مشاهده پروفایل
    log_user_activity(
        user=request.user,
        activity_type='view',
        description='مشاهده صفحه پروفایل و مدارک',
        related_model='Employee',
        related_object_id=employee.id,
        url=reverse('contractor_management:contractor_profile'),
        request=request
    )
    
    context = {
        'employee': employee,
        'profile_completion': profile_completion,
    }
    
    return render(request, 'contractor_management/auth/contractor_profile.html', context)


# =====================================
# CRUD Operations for Data Management
# =====================================

from .forms import ContractorForm, EmployeeForm, VehicleForm, ContractorEmployeeForm, EmployeeDocumentForm, ContractorUserManagementForm, ContractorVehicleForm
from django.http import JsonResponse

# Contractor CRUD
@permission_required("contractor_dashboard")
@login_required
def contractor_create(request):
    """ایجاد پیمانکار جدید"""
    if request.method == 'POST':
        form = ContractorForm(request.POST)
        if form.is_valid():
            contractor = form.save()
            
            # ثبت فعالیت
            log_user_activity(
                user=request.user,
                activity_type='create',
                description=f'ایجاد پیمانکار جدید: {contractor.company_name}',
                related_model='Contractor',
                related_object_id=contractor.id,
                url=reverse('contractor_management:data_management'),
                request=request
            )
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'پیمانکار با موفقیت ایجاد شد',
                    'contractor': {
                        'id': contractor.id,
                        'company_name': contractor.company_name,
                        'manager_name': contractor.manager_name,
                        'activity_field': contractor.activity_field,
                        'manager_phone': contractor.manager_phone,
                    }
                })
            
            messages.success(request, 'پیمانکار با موفقیت ایجاد شد')
            return redirect('contractor_management:data_management')
        else:
            logger.error(f"Contractor form validation errors: {form.errors.as_json()}")
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors
                }, status=400)
    else:
        form = ContractorForm()
    
    return render(request, 'contractor_management/contractor_form.html', {'form': form})


@permission_required("contractor_dashboard")
@login_required
def contractor_edit(request, pk):
    """ویرایش پیمانکار"""
    contractor = get_object_or_404(Contractor, pk=pk)
    
    if request.method == 'POST':
        form = ContractorForm(request.POST, instance=contractor)
        if form.is_valid():
            contractor = form.save()
            
            # ثبت فعالیت
            log_user_activity(
                user=request.user,
                activity_type='update',
                description=f'ویرایش پیمانکار: {contractor.company_name}',
                related_model='Contractor',
                related_object_id=contractor.id,
                url=reverse('contractor_management:data_management'),
                request=request
            )
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'پیمانکار با موفقیت بروزرسانی شد',
                    'contractor': {
                        'id': contractor.id,
                        'company_name': contractor.company_name,
                        'manager_name': contractor.manager_name,
                        'activity_field': contractor.activity_field,
                        'manager_phone': contractor.manager_phone,
                    }
                })
            
            messages.success(request, 'پیمانکار با موفقیت بروزرسانی شد')
            return redirect('contractor_management:data_management')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors
                }, status=400)
    else:
        form = ContractorForm(instance=contractor)
    
    return render(request, 'contractor_management/contractor_form.html', {'form': form, 'contractor': contractor})


@permission_required("contractor_dashboard")
@login_required
def contractor_delete(request, pk):
    """حذف پیمانکار"""
    contractor = get_object_or_404(Contractor, pk=pk)
    
    if request.method == 'POST':
        company_name = contractor.company_name
        
        # ثبت فعالیت
        log_user_activity(
            user=request.user,
            activity_type='delete',
            description=f'حذف پیمانکار: {company_name}',
            related_model='Contractor',
            related_object_id=contractor.id,
            url=reverse('contractor_management:data_management'),
            request=request
        )
        
        contractor.delete()
        
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'success': True,
                'message': f'پیمانکار {company_name} با موفقیت حذف شد'
            })
        
        messages.success(request, f'پیمانکار {company_name} با موفقیت حذف شد')
        return redirect('contractor_management:data_management')
    
    # GET request - show confirmation page
    return render(request, 'contractor_management/confirm_delete.html', {
        'object': contractor,
        'object_type': 'پیمانکار',
        'object_name': contractor.company_name,
        'cancel_url': reverse('contractor_management:data_management')
    })


# Employee CRUD
@permission_required("contractor_dashboard")
@login_required
def employee_create(request):
    """ایجاد کارمند جدید"""
    if request.method == 'POST':
        form = EmployeeForm(request.POST)
        if form.is_valid():
            employee = form.save()
            
            # ثبت فعالیت
            log_user_activity(
                user=request.user,
                activity_type='create',
                description=f'ایجاد کارمند جدید: {employee.first_name} {employee.last_name}',
                related_model='Employee',
                related_object_id=employee.id,
                url=reverse('contractor_management:data_management'),
                request=request
            )
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'کارمند با موفقیت ایجاد شد',
                    'employee': {
                        'id': employee.id,
                        'first_name': employee.first_name,
                        'last_name': employee.last_name,
                        'national_id': employee.national_id,
                        'position': employee.position,
                        'contractor_name': employee.contractor.company_name,
                    }
                })
            
            messages.success(request, 'کارمند با موفقیت ایجاد شد')
            return redirect('contractor_management:data_management')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors
                }, status=400)
    else:
        form = EmployeeForm()
    
    return render(request, 'contractor_management/employee_form.html', {'form': form})


@permission_required("contractor_dashboard")
@login_required
def employee_edit(request, pk):
    """ویرایش کارمند"""
    employee = get_object_or_404(Employee, pk=pk)
    
    if request.method == 'POST':
        form = EmployeeForm(request.POST, instance=employee)
        if form.is_valid():
            employee = form.save()
            
            # ثبت فعالیت
            log_user_activity(
                user=request.user,
                activity_type='update',
                description=f'ویرایش کارمند: {employee.first_name} {employee.last_name}',
                related_model='Employee',
                related_object_id=employee.id,
                url=reverse('contractor_management:data_management'),
                request=request
            )
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'کارمند با موفقیت بروزرسانی شد',
                    'employee': {
                        'id': employee.id,
                        'first_name': employee.first_name,
                        'last_name': employee.last_name,
                        'national_id': employee.national_id,
                        'position': employee.position,
                        'contractor_name': employee.contractor.company_name,
                    }
                })
            
            messages.success(request, 'کارمند با موفقیت بروزرسانی شد')
            return redirect('contractor_management:data_management')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors
                }, status=400)
    else:
        form = EmployeeForm(instance=employee)
    
    return render(request, 'contractor_management/employee_form.html', {'form': form, 'employee': employee})


@permission_required("contractor_dashboard")
@login_required
def employee_delete(request, pk):
    """حذف کارمند"""
    employee = get_object_or_404(Employee, pk=pk)
    
    if request.method == 'POST':
        employee_name = f"{employee.first_name} {employee.last_name}"
        
        # ثبت فعالیت
        log_user_activity(
            user=request.user,
            activity_type='delete',
            description=f'حذف کارمند: {employee_name}',
            related_model='Employee',
            related_object_id=employee.id,
            url=reverse('contractor_management:data_management'),
            request=request
        )
        
        employee.delete()
        
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'success': True,
                'message': f'کارمند {employee_name} با موفقیت حذف شد'
            })
        
        messages.success(request, f'کارمند {employee_name} با موفقیت حذف شد')
        return redirect('contractor_management:data_management')
    
    # GET request - show confirmation page
    return render(request, 'contractor_management/confirm_delete.html', {
        'object': employee,
        'object_type': 'کارمند',
        'object_name': f"{employee.first_name} {employee.last_name}",
        'cancel_url': reverse('contractor_management:data_management')
    })


# Vehicle CRUD
@permission_required("contractor_dashboard")
@login_required
def vehicle_create(request):
    """ایجاد خودرو جدید"""
    if request.method == 'POST':
        form = VehicleForm(request.POST)
        if form.is_valid():
            vehicle = form.save()
            
            # ثبت فعالیت
            log_user_activity(
                user=request.user,
                activity_type='create',
                description=f'ایجاد خودرو جدید: {vehicle.license_plate}',
                related_model='Vehicle',
                related_object_id=vehicle.id,
                url=reverse('contractor_management:data_management'),
                request=request
            )
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'خودرو با موفقیت ایجاد شد',
                    'vehicle': {
                        'id': vehicle.id,
                        'license_plate': vehicle.license_plate,
                        'vehicle_type': vehicle.vehicle_type,
                        'vehicle_category': vehicle.get_vehicle_category_display(),
                        'driver_name': vehicle.driver_name,
                        'contractor_name': vehicle.contractor.company_name,
                    }
                })
            
            messages.success(request, 'خودرو با موفقیت ایجاد شد')
            return redirect('contractor_management:data_management')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors
                }, status=400)
    else:
        form = VehicleForm()
    
    return render(request, 'contractor_management/vehicle_form.html', {'form': form})


@permission_required("contractor_dashboard")
@login_required
def vehicle_edit(request, pk):
    """ویرایش خودرو"""
    vehicle = get_object_or_404(Vehicle, pk=pk)
    
    if request.method == 'POST':
        form = VehicleForm(request.POST, instance=vehicle)
        if form.is_valid():
            vehicle = form.save()
            
            # ثبت فعالیت
            log_user_activity(
                user=request.user,
                activity_type='update',
                description=f'ویرایش خودرو: {vehicle.license_plate}',
                related_model='Vehicle',
                related_object_id=vehicle.id,
                url=reverse('contractor_management:data_management'),
                request=request
            )
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'خودرو با موفقیت بروزرسانی شد',
                    'vehicle': {
                        'id': vehicle.id,
                        'license_plate': vehicle.license_plate,
                        'vehicle_type': vehicle.vehicle_type,
                        'vehicle_category': vehicle.get_vehicle_category_display(),
                        'driver_name': vehicle.driver_name,
                        'contractor_name': vehicle.contractor.company_name,
                    }
                })
            
            messages.success(request, 'خودرو با موفقیت بروزرسانی شد')
            return redirect('contractor_management:data_management')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'errors': form.errors
                }, status=400)
    else:
        form = VehicleForm(instance=vehicle)
    
    return render(request, 'contractor_management/vehicle_form.html', {'form': form, 'vehicle': vehicle})


@permission_required("contractor_dashboard")
@login_required
def vehicle_delete(request, pk):
    """حذف خودرو"""
    vehicle = get_object_or_404(Vehicle, pk=pk)
    
    if request.method == 'POST':
        license_plate = vehicle.license_plate
        
        # ثبت فعالیت
        log_user_activity(
            user=request.user,
            activity_type='delete',
            description=f'حذف خودرو: {license_plate}',
            related_model='Vehicle',
            related_object_id=vehicle.id,
            url=reverse('contractor_management:data_management'),
            request=request
        )
        
        vehicle.delete()
        
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'success': True,
                'message': f'خودرو {license_plate} با موفقیت حذف شد'
            })
        
        messages.success(request, f'خودرو {license_plate} با موفقیت حذف شد')
        return redirect('contractor_management:data_management')
    
    # GET request - show confirmation page
    return render(request, 'contractor_management/confirm_delete.html', {
        'object': vehicle,
        'object_type': 'خودرو',
        'object_name': f"{vehicle.license_plate} - {vehicle.driver_name}",
        'cancel_url': reverse('contractor_management:data_management')
    })


# =====================================
# Contractor self-service: manage vehicles
# =====================================

@login_required
def manage_vehicles(request):
    if not hasattr(request.user, 'contractor_profile'):
        return HttpResponseForbidden('دسترسی غیرمجاز')
    contractor = request.user.contractor_profile
    vehicles = Vehicle.objects.filter(contractor=contractor).order_by('license_plate')

    if request.method == 'POST':
        form = ContractorVehicleForm(request.POST)
        if form.is_valid():
            v = form.save(commit=False)
            v.contractor = contractor
            v.save()
            messages.success(request, 'خودرو با موفقیت اضافه شد')
            return redirect('contractor_management:manage_vehicles')
        else:
            messages.error(request, 'لطفاً خطاهای فرم را برطرف کنید')
    else:
        form = ContractorVehicleForm()

    return render(request, 'contractor_management/auth/manage_vehicles.html', {
        'vehicles': vehicles,
        'form': form,
        'contractor': contractor,
    })

# =====================================
# Import/Export Functions
# =====================================

import pandas as pd
from io import BytesIO

@permission_required("contractor_dashboard")
@login_required
def export_data(request, model_type):
    """خروجی Excel برای پیمانکاران، کارکنان یا خودروها"""
    
    wb = Workbook()
    ws = wb.active
    ws.sheet_properties.rightToLeft = True
    
    # استایل‌ها
    header_font = Font(name='B Nazanin', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    if model_type == 'contractors':
        ws.title = "پیمانکاران"
        
        headers = ['نام شرکت', 'مدیرعامل', 'حوزه فعالیت', 'شماره تماس', 
                   'تاریخ انقضای بیمه مسئولیت', 'تاریخ انقضای بیمه آتش‌سوزی']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        contractors = Contractor.objects.all()
        for row, contractor in enumerate(contractors, 2):
            ws.cell(row=row, column=1, value=contractor.company_name)
            ws.cell(row=row, column=2, value=contractor.manager_name)
            ws.cell(row=row, column=3, value=contractor.activity_field)
            ws.cell(row=row, column=4, value=contractor.manager_phone)
            ws.cell(row=row, column=5, value=str(contractor.liability_insurance_expiry) if contractor.liability_insurance_expiry else '')
            ws.cell(row=row, column=6, value=str(contractor.fire_insurance_expiry) if contractor.fire_insurance_expiry else '')
        
        filename = f"contractors_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    elif model_type == 'employees':
        ws.title = "کارکنان"
        
        headers = ['نام', 'نام خانوادگی', 'کد ملی', 'تاریخ تولد', 'سمت', 
                   'شماره تماس', 'مدرک تحصیلی', 'پیمانکار']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        employees = Employee.objects.select_related('contractor').all()
        for row, employee in enumerate(employees, 2):
            ws.cell(row=row, column=1, value=employee.first_name)
            ws.cell(row=row, column=2, value=employee.last_name)
            ws.cell(row=row, column=3, value=employee.national_id)
            ws.cell(row=row, column=4, value=str(employee.birth_date) if employee.birth_date else '')
            ws.cell(row=row, column=5, value=employee.position)
            ws.cell(row=row, column=6, value=employee.phone_number)
            ws.cell(row=row, column=7, value=employee.education)
            ws.cell(row=row, column=8, value=employee.contractor.company_name)
        
        filename = f"employees_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    elif model_type == 'vehicles':
        ws.title = "خودروها"
        
        headers = ['پلاک', 'نوع خودرو', 'دسته‌بندی', 'نام راننده', 'پیمانکار',
                   'تاریخ انقضای بیمه', 'تاریخ انقضای معاینه فنی']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        vehicles = Vehicle.objects.select_related('contractor').all()
        for row, vehicle in enumerate(vehicles, 2):
            ws.cell(row=row, column=1, value=vehicle.license_plate)
            ws.cell(row=row, column=2, value=vehicle.vehicle_type)
            ws.cell(row=row, column=3, value=vehicle.get_vehicle_category_display())
            ws.cell(row=row, column=4, value=vehicle.driver_name)
            ws.cell(row=row, column=5, value=vehicle.contractor.company_name)
            ws.cell(row=row, column=6, value=str(vehicle.insurance_expiry) if vehicle.insurance_expiry else '')
            ws.cell(row=row, column=7, value=str(vehicle.technical_inspection_expiry) if vehicle.technical_inspection_expiry else '')
        
        filename = f"vehicles_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    else:
        return JsonResponse({'success': False, 'message': 'Invalid model type'}, status=400)
    
    # ذخیره فایل
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    # ثبت فعالیت
    log_user_activity(
        user=request.user,
        activity_type='export',
        description=f'خروجی Excel {model_type}',
        related_model=model_type,
        related_object_id=None,
        url=request.get_full_path(),
        request=request
    )
    
    return response


@permission_required("contractor_dashboard")
@login_required
def download_template(request, model_type):
    """دانلود فایل الگو برای Import"""
    
    wb = Workbook()
    ws = wb.active
    ws.sheet_properties.rightToLeft = True
    
    # استایل‌ها
    header_font = Font(name='B Nazanin', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='28A745', end_color='28A745', fill_type='solid')
    center_alignment = Alignment(horizontal='center', vertical='center')
    example_fill = PatternFill(start_color='E7F3E7', end_color='E7F3E7', fill_type='solid')
    
    if model_type == 'contractors':
        ws.title = "الگو پیمانکاران"
        
        headers = ['نام شرکت*', 'مدیرعامل*', 'حوزه فعالیت*', 'شماره تماس*', 
                   'تاریخ انقضای بیمه مسئولیت', 'تاریخ انقضای بیمه آتش‌سوزی']
        example_data = ['شرکت نمونه', 'علی احمدی', 'ساختمانی', '09123456789', '2024-12-31', '2024-12-31']
        filename = "template_contractors.xlsx"
    
    elif model_type == 'employees':
        ws.title = "الگو کارکنان"
        
        headers = ['نام*', 'نام خانوادگی*', 'کد ملی*', 'تاریخ تولد*', 'سمت*', 
                   'شماره تماس*', 'مدرک تحصیلی*', 'نام شرکت پیمانکار*']
        example_data = ['علی', 'احمدی', '1234567890', '1980-01-01', 'مهندس', '09123456789', 'کارشناسی', 'شرکت نمونه']
        filename = "template_employees.xlsx"
    
    elif model_type == 'vehicles':
        ws.title = "الگو خودروها"
        
        headers = ['پلاک*', 'نوع خودرو*', 'دسته‌بندی* (mining/light/transportation)', 
                   'نام راننده*', 'نام شرکت پیمانکار*',
                   'تاریخ انقضای بیمه', 'تاریخ انقضای معاینه فنی']
        example_data = ['12ا456', 'کامیون', 'mining', 'علی رضایی', 'شرکت نمونه', '2024-12-31', '2024-12-31']
        filename = "template_vehicles.xlsx"
    
    else:
        return JsonResponse({'success': False, 'message': 'Invalid model type'}, status=400)
    
    # اضافه کردن هدرها
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
    
    # اضافه کردن ردیف نمونه
    for col, data in enumerate(example_data, 1):
        cell = ws.cell(row=2, column=col, value=data)
        cell.fill = example_fill
        cell.alignment = center_alignment
    
    # ذخیره فایل
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response


@permission_required("contractor_dashboard")
@login_required
def import_data(request, model_type):
    """ورود Excel برای پیمانکاران، کارکنان یا خودروها"""
    
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'Invalid request method'}, status=400)
    
    if 'file' not in request.FILES:
        return JsonResponse({'success': False, 'message': 'لطفاً فایل را انتخاب کنید'}, status=400)
    
    excel_file = request.FILES['file']
    
    try:
        df = pd.read_excel(excel_file)
        success_count = 0
        error_count = 0
        errors = []
        
        if model_type == 'contractors':
            for index, row in df.iterrows():
                try:
                    Contractor.objects.create(
                        company_name=row['نام شرکت*'] if 'نام شرکت*' in row else row.iloc[0],
                        manager_name=row['مدیرعامل*'] if 'مدیرعامل*' in row else row.iloc[1],
                        activity_field=row['حوزه فعالیت*'] if 'حوزه فعالیت*' in row else row.iloc[2],
                        manager_phone=str(row['شماره تماس*'] if 'شماره تماس*' in row else row.iloc[3]),
                        number_of_social_insurance=''
                    )
                    success_count += 1
                except Exception as e:
                    error_count += 1
                    errors.append(f'ردیف {index + 2}: {str(e)}')
        
        elif model_type == 'employees':
            for index, row in df.iterrows():
                try:
                    # پیدا کردن پیمانکار
                    contractor_name = row['نام شرکت پیمانکار*'] if 'نام شرکت پیمانکار*' in row else row.iloc[7]
                    contractor = Contractor.objects.filter(company_name=contractor_name).first()
                    
                    if not contractor:
                        raise ValueError(f'پیمانکار {contractor_name} یافت نشد')
                    
                    Employee.objects.create(
                        contractor=contractor,
                        first_name=row['نام*'] if 'نام*' in row else row.iloc[0],
                        last_name=row['نام خانوادگی*'] if 'نام خانوادگی*' in row else row.iloc[1],
                        national_id=str(row['کد ملی*'] if 'کد ملی*' in row else row.iloc[2]),
                        birth_date=pd.to_datetime(row['تاریخ تولد*'] if 'تاریخ تولد*' in row else row.iloc[3]).date(),
                        position=row['سمت*'] if 'سمت*' in row else row.iloc[4],
                        phone_number=str(row['شماره تماس*'] if 'شماره تماس*' in row else row.iloc[5]),
                        education=row['مدرک تحصیلی*'] if 'مدرک تحصیلی*' in row else row.iloc[6],
                        number_of_insurance=''
                    )
                    success_count += 1
                except Exception as e:
                    error_count += 1
                    errors.append(f'ردیف {index + 2}: {str(e)}')
        
        elif model_type == 'vehicles':
            for index, row in df.iterrows():
                try:
                    # پیدا کردن پیمانکار
                    contractor_name = row['نام شرکت پیمانکار*'] if 'نام شرکت پیمانکار*' in row else row.iloc[4]
                    contractor = Contractor.objects.filter(company_name=contractor_name).first()
                    
                    if not contractor:
                        raise ValueError(f'پیمانکار {contractor_name} یافت نشد')
                    
                    category_col = 'دسته‌بندی* (mining/light/transportation)' if 'دسته‌بندی* (mining/light/transportation)' in row else 2
                    
                    Vehicle.objects.create(
                        contractor=contractor,
                        license_plate=row['پلاک*'] if 'پلاک*' in row else row.iloc[0],
                        vehicle_type=row['نوع خودرو*'] if 'نوع خودرو*' in row else row.iloc[1],
                        vehicle_category=row[category_col] if isinstance(category_col, str) else row.iloc[category_col],
                        driver_name=row['نام راننده*'] if 'نام راننده*' in row else row.iloc[3],
                    )
                    success_count += 1
                except Exception as e:
                    error_count += 1
                    errors.append(f'ردیف {index + 2}: {str(e)}')
        
        # ثبت فعالیت
        log_user_activity(
            user=request.user,
            activity_type='import',
            description=f'ورودی Excel {model_type}: {success_count} موفق, {error_count} خطا',
            related_model=model_type,
            related_object_id=None,
            url=request.get_full_path(),
            request=request
        )
        
        message = f'ورودی کامل شد: {success_count} ردیف موفق'
        if error_count > 0:
            message += f', {error_count} ردیف با خطا'
        
        return JsonResponse({
            'success': True,
            'message': message,
            'success_count': success_count,
            'error_count': error_count,
            'errors': errors[:10]  # فقط 10 خطای اول
        })
    
    except Exception as e:
        return JsonResponse({
            'success': False,
            'message': f'خطا در پردازش فایل: {str(e)}'
        }, status=400)
