from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse, JsonResponse
from django.db.models import Count, Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.generic import TemplateView
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from permissions.utils import class_permission_required, permission_required

from .forms import (
    DumpCountEventForm,
    DumpCountSessionForm,
    MachineActivityForm,
    format_jalali_date,
    parse_jalali_date,
)
from accounts.models import UserProfile
from .models import DumpCountEvent, DumpCountSession, LoadingHaulingReport, MachineActivity


def _set_sheet_rtl(sheet):
    sheet.sheet_view.rightToLeft = True


def _apply_table_style(sheet, start_row, end_row, start_col, end_col):
    thin = Side(style='thin', color='999999')
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row=row, column=col)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border


@class_permission_required('machine_activity_create')
class MachineActivityCreateView(LoginRequiredMixin, TemplateView):
    template_name = 'mining_operations/machine_activity_form.html'

    def get(self, request, *args, **kwargs):
        form = MachineActivityForm(initial={
            'date': request.GET.get('date') or None,
            'shift': request.GET.get('shift') or None,
        })
        return render(request, self.template_name, {'form': form})

    def post(self, request, *args, **kwargs):
        form = MachineActivityForm(request.POST)
        if form.is_valid():
            instance = form.save(commit=False)
            instance.created_by = request.user
            instance.save()

            totals = MachineActivity.calculate_shift_totals(instance.date, instance.shift)
            messages.success(
                request,
                (
                    'رکورد با موفقیت ثبت شد. '
                    f'مجموع ساعت کارکرد شیفت: {totals["total_work_hours"]} | '
                    f'مجموع آماده به کاری شیفت: {totals["total_ready_hours"]}'
                ),
            )

            return redirect(
                f"{reverse('mining_operations:machine_activity_create')}?date={instance.date}&shift={instance.shift}"
            )

        return render(request, self.template_name, {'form': form})


@class_permission_required('loading_hauling_list')
class LoadingHaulingListView(LoginRequiredMixin, TemplateView):
    template_name = 'mining_operations/loading_hauling_list.html'

    def get(self, request, *args, **kwargs):
        date_value = request.GET.get('date')
        shift = request.GET.get('shift')
        parsed_date = parse_jalali_date(date_value) if date_value else None

        qs = DumpCountSession.objects.all()
        if parsed_date:
            qs = qs.filter(date=parsed_date)
        if shift in {'A', 'B', 'C', 'D'}:
            qs = qs.filter(shift=shift)

        qs = (
            qs.values('date', 'shift')
            .annotate(
                total_sessions=Count('id', distinct=True),
                total_loads=Count('dump_events', distinct=True),
            )
            .order_by('-date', 'shift')
        )

        context = {
            'reports': qs[:200],
            'filter_date': date_value or '',
            'filter_shift': shift or '',
        }
        return render(request, self.template_name, context)


class LoadingHaulingDetailView(LoginRequiredMixin, TemplateView):
    template_name = 'mining_operations/loading_hauling_detail.html'

    def get(self, request, *args, **kwargs):
        date_value = request.GET.get('date')
        shift = request.GET.get('shift')
        parsed_date = parse_jalali_date(date_value) if date_value else None

        if not parsed_date or shift not in {'A', 'B', 'C', 'D'}:
            messages.error(request, 'لطفا تاریخ و شیفت معتبر انتخاب کنید.')
            return redirect(reverse('mining_operations:loading_hauling_list'))

        event_qs = DumpCountEvent.objects.select_related(
            'dumper_machine',
            'session',
            'session__block',
            'session__dump',
            'session__loader_machine',
            'session__mineral_type',
        ).filter(session__date=parsed_date, session__shift=shift)

        mineral_types = list(
            event_qs.values_list('session__mineral_type__name', flat=True)
            .distinct()
            .order_by('session__mineral_type__name')
        )

        hauling_rows = {}
        for event in event_qs:
            if not event.dumper_machine_id or not event.session_id:
                continue
            driver_label = (event.driver_name or '').strip()
            key = (
                event.dumper_machine_id,
                driver_label,
                event.session.block_id,
                event.session.loader_machine_id,
                event.session.dump_id,
            )
            row = hauling_rows.get(key)
            if not row:
                driver_name, driver_code = _split_personnel_label(driver_label)
                row = {
                    'machine_id': event.dumper_machine_id,
                    'code': event.dumper_machine.workshop_code,
                    'driver_name': driver_name,
                    'driver_code': driver_code,
                    'block': event.session.block.block_name if event.session.block else '',
                    'loader_code': event.session.loader_machine.workshop_code if event.session.loader_machine else '',
                    'dump': event.session.dump.dump_name if event.session.dump else '',
                    'counts': {name: 0 for name in mineral_types},
                    'total': 0,
                }
                hauling_rows[key] = row

            mineral_name = event.session.mineral_type.name if event.session.mineral_type else ''
            if mineral_name in row['counts']:
                row['counts'][mineral_name] += 1
            row['total'] += 1

        dumper_filter = (
            Q(machine__machine_type__name__icontains='دامپ')
            | Q(machine__machine_type__name__icontains='دامپتراک')
            | Q(machine__machine_type__name__icontains='تراک')
            | Q(machine__machine_workgroup__name__icontains='حمل')
            | Q(machine__machine_type__machine_workgroup__name__icontains='حمل')
        )
        dumper_activity_qs = (
            MachineActivity.objects
            .select_related('machine', 'machine__machine_type', 'machine__machine_workgroup')
            .filter(machine__is_active=True, date=parsed_date, shift=shift)
            .filter(dumper_filter)
        )
        dumper_activity_map = {}
        for act in dumper_activity_qs:
            key = (act.machine_id, (act.operator_name or '').strip())
            if key not in dumper_activity_map:
                dumper_activity_map[key] = act

        hauling_data = []
        for row in hauling_rows.values():
            activity = None
            if row['machine_id']:
                activity = dumper_activity_map.get((row['machine_id'], row['driver_name']))
                if not activity:
                    activity = dumper_activity_map.get((row['machine_id'], ''))
            hauling_data.append({
                **row,
                'ready_hours': float(activity.ready_hours) if activity else '',
                'work_hours': float(activity.work_hours) if activity else '',
                'stop_hours': float(activity.stop_hours) if activity else '',
                'stop_reason': (
                    activity.stop_description
                    if activity and activity.stop_description
                    else (activity.get_stop_reason_display() if activity and activity.stop_reason else '')
                ),
            })

        loader_filter = (
            Q(machine__machine_type__name__icontains='لودر')
            | Q(machine__machine_type__name__icontains='بیل')
            | Q(machine__machine_type__name__icontains='بارکننده')
            | Q(machine__machine_workgroup__name__icontains='بارکننده')
            | Q(machine__machine_type__machine_workgroup__name__icontains='بارکننده')
        )
        loader_activities = (
            MachineActivity.objects
            .select_related('machine', 'machine__machine_type', 'machine__machine_workgroup')
            .filter(machine__is_active=True, date=parsed_date, shift=shift)
            .filter(loader_filter)
            .order_by('machine__workshop_code')
        )

        loader_rows = []
        for item in loader_activities:
            name, code = _split_personnel_label(item.operator_name)
            loader_rows.append({
                'machine_code': item.machine.workshop_code,
                'operator_name': name,
                'operator_code': code,
                'start_hour': item.start_hour,
                'end_hour': item.end_hour,
                'ready_hours': item.ready_hours,
                'work_hours': item.work_hours,
                'stop_hours': item.stop_hours,
                'stop_reason': item.stop_description or (item.get_stop_reason_display() if item.stop_reason else ''),
                'date': item.date,
            })

        context = {
            'report_date': parsed_date,
            'report_date_label': date_value or format_jalali_date(parsed_date),
            'report_shift': shift,
            'mineral_types': mineral_types,
            'hauling_rows': hauling_data,
            'loader_rows': loader_rows,
        }
        return render(request, self.template_name, context)


def _split_personnel_label(value):
    if not value:
        return '', ''
    raw = str(value).strip()
    if raw.endswith(')') and '(' in raw:
        name, code = raw.rsplit('(', 1)
        return name.strip(), code[:-1].strip()
    return raw, ''


@class_permission_required('count_sheet_create')
class DumpCountSheetView(LoginRequiredMixin, TemplateView):
    template_name = 'mining_operations/dump_count_sheet.html'

    def get(self, request, *args, **kwargs):
        selected_session = self._get_selected_session(request)
        default_shift = request.GET.get('shift')
        if not default_shift:
            try:
                from shift_manager.utils import get_current_shift_and_group
                _, group = get_current_shift_and_group(request.user)
                if group in {'A', 'B', 'C', 'D'}:
                    default_shift = group
            except Exception:
                default_shift = None
        session_form = DumpCountSessionForm(
            user=request.user,
            shift=default_shift,
        )
        event_form = DumpCountEventForm(
            session=selected_session,
            user=request.user,
        )
        context = self._build_context(selected_session, session_form, event_form)
        context['current_user_label'] = self._get_current_user_label(request.user)
        if request.headers.get('x-requested-with') == 'XMLHttpRequest':
            html = render_to_string(
                'mining_operations/_dump_count_right.html',
                context,
                request=request,
            )
            return JsonResponse({'html': html})
        return render(request, self.template_name, context)

    def post(self, request, *args, **kwargs):
        action = request.POST.get('action')

        if action == 'create_session':
            return self._create_session(request)
        if action == 'add_event':
            return self._add_event(request)
        if action == 'add_activity':
            return self._add_activity(request)
        if action == 'close_session':
            return self._close_session(request)

        messages.error(request, 'درخواست نامعتبر است.')
        return redirect(reverse('mining_operations:count_sheet_create'))

    def _create_session(self, request):
        session_form = DumpCountSessionForm(
            request.POST,
            user=request.user,
            shift=request.POST.get('shift'),
        )
        event_form = DumpCountEventForm(
            session=None,
            user=request.user,
        )

        if not session_form.is_valid():
            messages.error(request, f"خطای فرم: {session_form.errors.as_text()}")
            context = self._build_context(None, session_form, event_form)
            context['current_user_label'] = self._get_current_user_label(request.user)
            return render(request, self.template_name, context)

        session = session_form.save(commit=False)
        from django.utils import timezone
        session.date = timezone.localdate()
        session.start_time = timezone.localtime().time().replace(microsecond=0)
        session.created_by = request.user
        session.save()
        messages.success(request, 'شیفت شمارش با موفقیت آغاز شد. حالا سرویس ها را ثبت کنید.')
        return redirect(f"{reverse('mining_operations:count_sheet_create')}?session_id={session.id}")

    def _add_event(self, request):
        session = get_object_or_404(DumpCountSession, pk=request.POST.get('session_id'))
        event_form = DumpCountEventForm(
            request.POST,
            session=session,
            user=request.user,
        )
        session_form = DumpCountSessionForm(
            instance=session,
            user=request.user,
            shift=session.shift,
        )

        if session.end_time:
            messages.error(request, 'این شیفت بسته شده است و امکان ثبت سرویس جدید ندارد.')
            return redirect(f"{reverse('mining_operations:count_sheet_create')}?session_id={session.id}")

        if not event_form.is_valid():
            messages.error(request, f"خطای فرم: {event_form.errors.as_text()}")
            context = self._build_context(session, session_form, event_form)
            context['current_user_label'] = self._get_current_user_label(request.user)
            return render(request, self.template_name, context)

        event = event_form.save(commit=False)
        event.session = session
        event.created_by = request.user
        event.save()

        messages.success(
            request,
            f'سرویس شماره {session.dump_events.count()} ثبت شد. ساعت ثبت: {event.load_time.strftime("%H:%M:%S")}'
        )
        return redirect(f"{reverse('mining_operations:count_sheet_create')}?session_id={session.id}")

    def _close_session(self, request):
        session = get_object_or_404(DumpCountSession, pk=request.POST.get('session_id'))
        if not session.end_time:
            from django.utils import timezone
            session.end_time = timezone.localtime().time().replace(microsecond=0)
            session.save(update_fields=['end_time'])
            created_count = 0
            updated_count = 0
            skipped_events = 0
            total_events = session.dump_events.count()

            event_counts = (
                session.dump_events
                .exclude(dumper_machine__isnull=True)
                .values('dumper_machine')
                .annotate(total=Count('id'))
            )

            for row in event_counts:
                dumper_id = row['dumper_machine']
                service_total = row['total'] or 0

                if not session.loader_machine_id or not session.mineral_type_id:
                    skipped_events += service_total
                    continue

                latest_driver = (
                    session.dump_events
                    .filter(dumper_machine_id=dumper_id)
                    .exclude(driver_name__isnull=True)
                    .exclude(driver_name__exact='')
                    .order_by('-load_time')
                    .values_list('driver_name', flat=True)
                    .first()
                )

                report, created = LoadingHaulingReport.objects.get_or_create(
                    date=session.date,
                    shift=session.shift,
                    loader_machine=session.loader_machine,
                    dumper_machine_id=dumper_id,
                    material_type=session.mineral_type,
                    block=session.block,
                    dump=session.dump,
                    defaults={
                        'service_count': service_total,
                        'loader_operator_name': session.loader_operator_name or '',
                        'dumper_operator_name': latest_driver or '',
                        'created_by': session.created_by or request.user,
                    },
                )
                if created:
                    created_count += 1
                else:
                    report.service_count = service_total
                    report.loader_operator_name = session.loader_operator_name or report.loader_operator_name
                    report.dumper_operator_name = latest_driver or report.dumper_operator_name
                    if not report.created_by:
                        report.created_by = session.created_by or request.user
                    report.block = session.block
                    report.dump = session.dump
                    report.material_type = session.mineral_type
                    report.loader_machine = session.loader_machine
                    report.save(
                        update_fields=[
                            'service_count',
                            'loader_operator_name',
                            'dumper_operator_name',
                            'created_by',
                            'block',
                            'dump',
                            'material_type',
                            'loader_machine',
                            'updated_at',
                        ]
                    )
                    updated_count += 1

            skipped_events += total_events - sum(row['total'] or 0 for row in event_counts)

            message_parts = [
                f'شیفت بسته شد. جمع کل سرویس ثبت شده: {total_events}',
                f'گزارش جدید: {created_count}',
            ]
            if updated_count:
                message_parts.append(f'به‌روزرسانی شده: {updated_count}')
            if skipped_events:
                message_parts.append(f'سرویس‌های بدون دامپ یا ناقص: {skipped_events}')

            messages.success(request, ' | '.join(message_parts))
        else:
            messages.info(request, 'این شیفت قبلا بسته شده است.')

        return redirect(reverse('mining_operations:count_sheet_create'))

    def _add_activity(self, request):
        session = get_object_or_404(DumpCountSession, pk=request.POST.get('session_id'))
        form = MachineActivityForm(request.POST)
        if not form.is_valid():
            messages.error(request, f"خطای فرم توقفات: {form.errors.as_text()}")
            session_form = DumpCountSessionForm(instance=session, user=request.user, shift=session.shift)
            event_form = DumpCountEventForm(session=session, user=request.user)
            context = self._build_context(session, session_form, event_form)
            context['activity_form'] = form
            context['current_user_label'] = self._get_current_user_label(request.user)
            return render(request, self.template_name, context)

        instance = form.save(commit=False)
        instance.date = session.date
        instance.shift = session.shift
        instance.created_by = request.user
        instance.save()
        messages.success(request, 'توقف/کارکرد دستگاه ثبت شد.')
        return redirect(f"{reverse('mining_operations:count_sheet_create')}?session_id={session.id}")

    def _get_selected_session(self, request):
        session_id = request.GET.get('session_id')
        if session_id:
            try:
                return DumpCountSession.objects.filter(end_time__isnull=True).get(pk=session_id)
            except DumpCountSession.DoesNotExist:
                return None
        return DumpCountSession.objects.filter(end_time__isnull=True).order_by('-created_at').first()

    def _build_context(self, selected_session, session_form, event_form):
        events = DumpCountEvent.objects.none()
        if selected_session:
            events = selected_session.dump_events.select_related('dumper_machine').all()

        sessions = DumpCountSession.objects.filter(end_time__isnull=True).order_by('-created_at')[:20]
        activities = MachineActivity.objects.none()
        activity_form = MachineActivityForm()
        if selected_session:
            activities = (
                MachineActivity.objects
                .select_related('machine')
                .filter(date=selected_session.date, shift=selected_session.shift)
                .order_by('machine__workshop_code')
            )
            activity_form = MachineActivityForm(initial={
                'date': format_jalali_date(selected_session.date),
                'shift': selected_session.shift,
            })

        return {
            'selected_session': selected_session,
            'sessions': sessions,
            'events': events,
            'session_form': session_form,
            'event_form': event_form,
            'activities': activities,
            'activity_form': activity_form,
        }

    @staticmethod
    def _get_current_user_label(user):
        if not user or not hasattr(user, 'userprofile'):
            return ''
        profile = user.userprofile
        name = profile.user.get_full_name() or profile.user.username
        code = profile.personnel_code or ''
        return f'{name} ({code})' if code else name


@permission_required('count_sheet_create')
@login_required
def shift_personnel_options(request):
    shift = request.GET.get('shift')
    role = request.GET.get('role')
    current_user_label = DumpCountSheetView._get_current_user_label(request.user)

    # Reuse form filtering logic for consistency
    personnel_choices = DumpCountSessionForm._get_personnel_choices(shift, request.user, role=role)
    results = []
    for value, label in personnel_choices:
        if not value:
            continue
        results.append({'id': value, 'text': label})

    return JsonResponse({'results': results, 'current_user': current_user_label})


@permission_required('machine_activity_export_excel')
@login_required
def export_machine_activity_excel(request):
    date = request.GET.get('date')
    shift = request.GET.get('shift')

    queryset = MachineActivity.objects.select_related('machine').all().order_by('machine__workshop_code')
    if date:
        queryset = queryset.filter(date=date)
    if shift:
        queryset = queryset.filter(shift=shift)

    wb = Workbook()
    ws = wb.active
    ws.title = 'کارکرد ماشین آلات'
    _set_sheet_rtl(ws)

    ws.merge_cells('A1:L1')
    ws['A1'] = 'لیست کارکرد ماشین آلات عملیات معدنی'
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws['A2'] = 'تاریخ:'
    ws['B2'] = date or '-'
    ws['D2'] = 'شیفت:'
    ws['E2'] = shift or '-'

    headers = [
        'ردیف',
        'کد دستگاه',
        'نام اپراتور',
        'ساعت شروع',
        'ساعت پایان',
        'آماده به کاری',
        'کارکرد مفید',
        'توقف',
        'علت توقف',
        'تاریخ',
    ]

    header_row = 4
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    data_row = header_row + 1
    for i, item in enumerate(queryset, start=1):
        ws.cell(row=data_row, column=1, value=i)
        ws.cell(row=data_row, column=2, value=item.machine.workshop_code)
        ws.cell(row=data_row, column=3, value=item.operator_name)
        ws.cell(row=data_row, column=4, value=item.start_hour.strftime('%H:%M') if item.start_hour else '')
        ws.cell(row=data_row, column=5, value=item.end_hour.strftime('%H:%M') if item.end_hour else '')
        ws.cell(row=data_row, column=6, value=float(item.ready_hours))
        ws.cell(row=data_row, column=7, value=float(item.work_hours))
        ws.cell(row=data_row, column=8, value=float(item.stop_hours))
        ws.cell(row=data_row, column=9, value=item.get_stop_reason_display() if item.stop_reason else '')
        ws.cell(row=data_row, column=10, value=str(item.date))
        data_row += 1

    if date and shift:
        totals = MachineActivity.calculate_shift_totals(date, shift)
        ws.merge_cells(start_row=data_row, start_column=1, end_row=data_row, end_column=5)
        ws.cell(row=data_row, column=1, value='جمع کل شیفت').font = Font(bold=True)
        ws.cell(row=data_row, column=6, value=float(totals['total_ready_hours'])).font = Font(bold=True)
        ws.cell(row=data_row, column=7, value=float(totals['total_work_hours'])).font = Font(bold=True)
        ws.cell(row=data_row, column=8, value=float(totals['total_stop_hours'])).font = Font(bold=True)

    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 16

    _apply_table_style(ws, header_row, max(data_row, header_row + 1), 1, 10)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=machine_activity_report.xlsx'
    wb.save(response)
    return response


@permission_required('loading_hauling_export_excel')
@login_required
def export_loading_hauling_excel(request):
    date = request.GET.get('date')
    shift = request.GET.get('shift')
    parsed_date = parse_jalali_date(date) if date else None

    queryset = LoadingHaulingReport.objects.none()
    if parsed_date and shift in {'A', 'B', 'C', 'D'}:
        queryset = (
            LoadingHaulingReport.objects
            .select_related('loader_machine', 'dumper_machine', 'material_type', 'block', 'dump')
            .filter(date=parsed_date, shift=shift)
            .order_by('dumper_machine__workshop_code', 'loader_machine__workshop_code')
        )

    wb = Workbook()
    ws_hauling = wb.active
    ws_hauling.title = 'گزارش حمل کننده ها'
    _set_sheet_rtl(ws_hauling)

    def _apply_header(ws, title, end_col):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws.cell(row=1, column=1, value=title).font = Font(size=14, bold=True)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')
        ws.cell(row=2, column=1, value='تاریخ:')
        ws.cell(row=2, column=2, value=date or '-')
        ws.cell(row=2, column=4, value='شیفت:')
        ws.cell(row=2, column=5, value=shift or '-')

    event_qs = DumpCountEvent.objects.select_related(
        'dumper_machine',
        'session',
        'session__block',
        'session__dump',
        'session__loader_machine',
        'session__mineral_type',
    )
    if parsed_date and shift in {'A', 'B', 'C', 'D'}:
        event_qs = event_qs.filter(session__date=parsed_date, session__shift=shift)

    mineral_types = list(
        event_qs.values_list('session__mineral_type__name', flat=True)
        .distinct()
        .order_by('session__mineral_type__name')
    )

    hauling_headers = [
        'ردیف',
        'کد دستگاه',
        'نام راننده',
        'کد پرسنلی',
        'شماره بلوک',
        'کد بارکننده',
        'نام دامپ',
        'آماده به کاری',
        'کارکرد مفید',
        'توقف',
        'علت توقف',
    ] + mineral_types + ['جمع کل']

    header_row = 4
    _apply_header(ws_hauling, 'گزارش حمل کننده های معدن سنگ آهن جلال آباد', len(hauling_headers))
    for idx, header in enumerate(hauling_headers, start=1):
        cell = ws_hauling.cell(row=header_row, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    hauling_rows = {}
    for event in event_qs:
        if not event.dumper_machine_id or not event.session_id:
            continue
        driver_label = (event.driver_name or '').strip()
        key = (
            event.dumper_machine_id,
            driver_label,
            event.session.block_id,
            event.session.loader_machine_id,
            event.session.dump_id,
        )
        row = hauling_rows.get(key)
        if not row:
            driver_name, driver_code = _split_personnel_label(driver_label)
            row = {
                'machine_id': event.dumper_machine_id,
                'code': event.dumper_machine.workshop_code,
                'driver_name': driver_name,
                'driver_code': driver_code,
                'block': event.session.block.block_name if event.session.block else '',
                'loader_code': event.session.loader_machine.workshop_code if event.session.loader_machine else '',
                'dump': event.session.dump.dump_name if event.session.dump else '',
                'counts': {name: 0 for name in mineral_types},
                'total': 0,
            }
            hauling_rows[key] = row

        mineral_name = event.session.mineral_type.name if event.session.mineral_type else ''
        if mineral_name in row['counts']:
            row['counts'][mineral_name] += 1
        row['total'] += 1

    dumper_filter = (
        Q(machine__machine_type__name__icontains='دامپ')
        | Q(machine__machine_type__name__icontains='دامپتراک')
        | Q(machine__machine_type__name__icontains='تراک')
        | Q(machine__machine_workgroup__name__icontains='حمل')
        | Q(machine__machine_type__machine_workgroup__name__icontains='حمل')
    )
    dumper_activity_qs = (
        MachineActivity.objects
        .select_related('machine', 'machine__machine_type', 'machine__machine_workgroup')
        .filter(machine__is_active=True)
        .filter(dumper_filter)
    )
    if parsed_date and shift in {'A', 'B', 'C', 'D'}:
        dumper_activity_qs = dumper_activity_qs.filter(date=parsed_date, shift=shift)

    dumper_activity_map = {}
    for act in dumper_activity_qs:
        key = (act.machine_id, (act.operator_name or '').strip())
        if key not in dumper_activity_map:
            dumper_activity_map[key] = act

    data_row = header_row + 1
    for i, row in enumerate(hauling_rows.values(), start=1):
        activity = None
        if row['machine_id']:
            activity = dumper_activity_map.get((row['machine_id'], row['driver_name']))
            if not activity:
                activity = dumper_activity_map.get((row['machine_id'], ''))
        ws_hauling.cell(row=data_row, column=1, value=i)
        ws_hauling.cell(row=data_row, column=2, value=row['code'])
        ws_hauling.cell(row=data_row, column=3, value=row['driver_name'])
        ws_hauling.cell(row=data_row, column=4, value=row['driver_code'])
        ws_hauling.cell(row=data_row, column=5, value=row['block'])
        ws_hauling.cell(row=data_row, column=6, value=row['loader_code'])
        ws_hauling.cell(row=data_row, column=7, value=row['dump'])
        ws_hauling.cell(row=data_row, column=8, value=float(activity.ready_hours) if activity else '')
        ws_hauling.cell(row=data_row, column=9, value=float(activity.work_hours) if activity else '')
        ws_hauling.cell(row=data_row, column=10, value=float(activity.stop_hours) if activity else '')
        ws_hauling.cell(
            row=data_row,
            column=11,
            value=(
                activity.stop_description
                if activity and activity.stop_description
                else (activity.get_stop_reason_display() if activity and activity.stop_reason else '')
            ),
        )
        for idx, mineral_name in enumerate(mineral_types, start=12):
            ws_hauling.cell(row=data_row, column=idx, value=row['counts'].get(mineral_name, 0))
        ws_hauling.cell(row=data_row, column=12 + len(mineral_types), value=row['total'])
        data_row += 1

    for col in range(1, len(hauling_headers) + 1):
        ws_hauling.column_dimensions[get_column_letter(col)].width = 16

    _apply_table_style(ws_hauling, header_row, max(data_row, header_row + 1), 1, len(hauling_headers))

    ws_loader = wb.create_sheet('گزارش بارکننده ها')
    _set_sheet_rtl(ws_loader)

    loader_headers = [
        'ردیف',
        'کد دستگاه',
        'نام اپراتور',
        'کد پرسنلی',
        'ساعت شروع',
        'ساعت پایان',
        'آماده به کاری',
        'کارکرد مفید',
        'توقف',
        'علت توقف',
        'تاریخ',
    ]

    _apply_header(ws_loader, 'گزارش بارکننده های معدن سنگ آهن جلال آباد', len(loader_headers))
    for idx, header in enumerate(loader_headers, start=1):
        cell = ws_loader.cell(row=header_row, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')

    loader_filter = (
        Q(machine__machine_type__name__icontains='لودر')
        | Q(machine__machine_type__name__icontains='بیل')
        | Q(machine__machine_type__name__icontains='بارکننده')
        | Q(machine__machine_workgroup__name__icontains='بارکننده')
        | Q(machine__machine_type__machine_workgroup__name__icontains='بارکننده')
    )
    loader_qs = (
        MachineActivity.objects
        .select_related('machine', 'machine__machine_type', 'machine__machine_workgroup')
        .filter(machine__is_active=True)
        .filter(loader_filter)
        .order_by('machine__workshop_code')
    )
    if parsed_date:
        loader_qs = loader_qs.filter(date=parsed_date)
    if shift in {'A', 'B', 'C', 'D'}:
        loader_qs = loader_qs.filter(shift=shift)

    data_row = header_row + 1
    for i, item in enumerate(loader_qs, start=1):
        name, code = _split_personnel_label(item.operator_name)
        ws_loader.cell(row=data_row, column=1, value=i)
        ws_loader.cell(row=data_row, column=2, value=item.machine.workshop_code)
        ws_loader.cell(row=data_row, column=3, value=name)
        ws_loader.cell(row=data_row, column=4, value=code)
        ws_loader.cell(row=data_row, column=5, value=item.start_hour.strftime('%H:%M') if item.start_hour else '')
        ws_loader.cell(row=data_row, column=6, value=item.end_hour.strftime('%H:%M') if item.end_hour else '')
        ws_loader.cell(row=data_row, column=7, value=float(item.ready_hours))
        ws_loader.cell(row=data_row, column=8, value=float(item.work_hours))
        ws_loader.cell(row=data_row, column=9, value=float(item.stop_hours))
        ws_loader.cell(
            row=data_row,
            column=10,
            value=item.stop_description or (item.get_stop_reason_display() if item.stop_reason else ''),
        )
        ws_loader.cell(row=data_row, column=11, value=str(item.date))
        data_row += 1

    for col in range(1, len(loader_headers) + 1):
        ws_loader.column_dimensions[get_column_letter(col)].width = 16

    _apply_table_style(ws_loader, header_row, max(data_row, header_row + 1), 1, len(loader_headers))

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=loading_hauling_report.xlsx'
    wb.save(response)
    return response
