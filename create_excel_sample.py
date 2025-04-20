import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ایجاد یک فایل اکسل جدید
wb = openpyxl.Workbook()
ws = wb.active

# داده‌های نمونه
headers = [
    'متن سوال', 'محدوده', 'نوع سوال', 'نوع ماشین', 'بخش مکانی', 'دسته‌بندی خودرو',
    'گزینه‌ها', 'گزینه‌های غیرقابل قبول', 'نوع آنومالی', 'حوزه HSE',
    'شرح آنومالی پیش‌فرض', 'اولویت', 'اقدام اصلاحی پیش‌فرض'
]

sample_data = [
    [
        'آیا سیستم روشنایی ماشین سالم است؟', 'ماشین', 'گزینه‌ای', 'دامپتراک', '', '',
        'سالم است,نیاز به تعمیر دارد,خراب است', 'نیاز به تعمیر دارد,خراب است',
        'نقص فنی', 'S', 'نقص در سیستم روشنایی ماشین', 'متوسط', 'تعمیر یا تعویض سیستم روشنایی'
    ],
    [
        'آیا کپسول آتش‌نشانی در محل نصب شده است؟', 'مکان', 'گزینه‌ای', '', 'انبار مواد شیمیایی', '',
        'نصب شده است,نصب نشده است,نیاز به شارژ دارد', 'نصب نشده است,نیاز به شارژ دارد',
        'ایمنی', 'S', 'عدم وجود یا نقص در کپسول آتش‌نشانی', 'فوری', 'نصب یا شارژ کپسول آتش‌نشانی'
    ],
    [
        'وضعیت ترمز ماشین چگونه است؟', 'ماشین‌آلات پیمانکار', 'گزینه‌ای', '', '', 'ماشین‌آلات معدنی',
        'سالم است,نیاز به تنظیم دارد,خراب است', 'نیاز به تنظیم دارد,خراب است',
        'نقص فنی', 'S', 'نقص در سیستم ترمز', 'فوری', 'تعمیر یا تنظیم سیستم ترمز'
    ],
    [
        'میزان آلودگی صوتی را ثبت کنید', 'مکان', 'متنی', '', 'سالن تولید', '',
        '', '', '', 'H', 'آلودگی صوتی بالاتر از حد مجاز', 'متوسط', 'نصب عایق صوتی و تجهیزات کاهش صدا'
    ],
    [
        'آیا کمربند ایمنی راننده سالم است؟', 'ماشین‌آلات پیمانکار', 'گزینه‌ای', '', '', 'خودروهای سبک',
        'سالم است,نیاز به تعمیر دارد,خراب است', 'نیاز به تعمیر دارد,خراب است',
        'ایمنی', 'S', 'نقص در کمربند ایمنی', 'فوری', 'تعمیر یا تعویض کمربند ایمنی'
    ]
]

# تنظیم استایل‌ها
header_font = Font(bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='009EF7', end_color='009EF7', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# اضافه کردن هدرها
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = border
    # تنظیم عرض ستون
    ws.column_dimensions[get_column_letter(col)].width = 20

# اضافه کردن داده‌ها
for row_idx, row_data in enumerate(sample_data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = value
        cell.alignment = cell_alignment
        cell.border = border

# تنظیم ارتفاع سطرها
ws.row_dimensions[1].height = 40
for i in range(2, len(sample_data) + 2):
    ws.row_dimensions[i].height = 30

# ذخیره فایل
wb.save('checklist_app/static/checklist_app/sample_import_questions.xlsx') 