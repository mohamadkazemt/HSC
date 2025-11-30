from django.views.decorators.csrf import csrf_exempt
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.core.paginator import Paginator
from django.db.models import Q
from django.utils import timezone
from datetime import timedelta, datetime
from django.contrib.auth.decorators import login_required
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import FireExtinguisherType, FireExtinguisher, ServiceRecord
from dashboard.models import Notification
from .forms import (
    FireExtinguisherTypeForm, FireExtinguisherForm, ServiceRecordForm,
    FireExtinguisherReplacementForm, FireExtinguisherLocationForm, ExcelImportForm
)

@login_required
@csrf_exempt
def extinguisher_create_ajax(request):
    if request.method == 'POST':
        form = FireExtinguisherForm(request.POST)
        if form.is_valid():
            extinguisher = form.save()
            return JsonResponse({'success': True, 'id': extinguisher.pk}, status=201)
        else:
            # Return structured errors and 400 status so client can handle per-field errors
            errors = {k: [e['message'] for e in v] for k, v in form.errors.get_json_data().items()}
            return JsonResponse({'success': False, 'errors': errors, 'message': 'فرم معتبر نیست'}, status=400)
    else:
        return JsonResponse({'success': False, 'message': 'درخواست نامعتبر است'}, status=400)

@login_required
def dashboard(request):
    # Get counts for different statuses
    status_counts = {
        'operational': FireExtinguisher.objects.filter(status='operational').count(),
        'needs_maintenance': FireExtinguisher.objects.filter(status='needs_maintenance').count(),
        'expired': FireExtinguisher.objects.filter(status='expired').count(),
        'reserved': FireExtinguisher.objects.filter(status='reserved').count(),
    }

    # Get upcoming service dates
    today = timezone.now().date()
    upcoming_extinguishers = FireExtinguisher.objects.filter(
        Q(next_scheduled_service_date__lte=today + timedelta(days=30)) |
        Q(pressure_test_due_date__lte=today + timedelta(days=30))
    ).select_related('extinguisher_type')
    
    # Prepare upcoming services data with calculated fields
    upcoming_services = []
    for extinguisher in upcoming_extinguishers:
        # Calculate days remaining for next service
        days_remaining = None
        service_type = None
        
        if extinguisher.next_scheduled_service_date:
            days_remaining = (extinguisher.next_scheduled_service_date - today).days
            service_type = 'سرویس دوره‌ای'
        elif extinguisher.pressure_test_due_date:
            days_remaining = (extinguisher.pressure_test_due_date - today).days
            service_type = 'تست فشار'
        
        if days_remaining is not None:
            upcoming_services.append({
                'extinguisher': extinguisher,
                'extinguisher_tag': extinguisher.serial_tag,
                'extinguisher_id': extinguisher.pk,
                'service_type': service_type,
                'days_remaining': days_remaining,
            })

    # Get recent service records
    recent_services = ServiceRecord.objects.select_related(
        'extinguisher', 'performed_by_user'
    ).order_by('-service_date')[:10]

    context = {
        'status_counts': status_counts,
        'upcoming_services': upcoming_services,
        'recent_services': recent_services,
    }
    return render(request, 'fire_extinguisher_management/dashboard.html', context)

@login_required
def extinguisher_list(request):
    extinguishers = FireExtinguisher.objects.select_related('extinguisher_type')
    
    # Filtering
    status = request.GET.get('status')
    extinguisher_type = request.GET.get('type')
    search = request.GET.get('search')
    
    if status:
        extinguishers = extinguishers.filter(status=status)
    if extinguisher_type:
        extinguishers = extinguishers.filter(extinguisher_type_id=extinguisher_type)
    if search:
        extinguishers = extinguishers.filter(
            Q(serial_tag__icontains=search) |
            Q(manufacturer__icontains=search) |
            Q(model_number__icontains=search)
        )

    # Pagination
    paginator = Paginator(extinguishers, 20)
    page = request.GET.get('page')
    extinguishers = paginator.get_page(page)

    types = FireExtinguisherType.objects.all()

    # Additional choices and querysets used by the create modal
    from django.apps import apps
    LocationSection = apps.get_model('anomalis', 'LocationSection')
    MiningMachine = apps.get_model('BaseInfo', 'MiningMachine')

    sections = LocationSection.objects.all()
    machines = MiningMachine.objects.filter(is_active=True)
    existing_extinguishers = FireExtinguisher.objects.all()

    # Add form for the create modal
    create_form = FireExtinguisherForm()

    context = {
        'extinguishers': extinguishers,
        'types': types,
        'sections': sections,
        'machines': machines,
        'existing_extinguishers': existing_extinguishers,
        'current_status': status,
        'current_type': extinguisher_type,
        'current_search': search,
        'form': create_form,  # Add the form to context
    }
    return render(request, 'fire_extinguisher_management/extinguisher_list.html', context)

@login_required
def extinguisher_detail(request, pk):
    extinguisher = get_object_or_404(FireExtinguisher.objects.select_related('extinguisher_type'), pk=pk)
    service_records = extinguisher.servicerecord_set.select_related('performed_by_user').order_by('-service_date')
    # Prepare forms used by modals on the detail page
    service_record_form = ServiceRecordForm(initial={'extinguisher': extinguisher})
    location_form = FireExtinguisherLocationForm(initial={
        'location_type': extinguisher.location_type,
        'location_section': extinguisher.location_section,
        'location_machine': extinguisher.location_machine,
    })
    replacement_form = FireExtinguisherReplacementForm()
    # Additional querysets for selects
    from django.apps import apps
    LocationSection = apps.get_model('anomalis', 'LocationSection')
    MiningMachine = apps.get_model('BaseInfo', 'MiningMachine')
    sections = LocationSection.objects.all()
    machines = MiningMachine.objects.filter(is_active=True)
    reserved_extinguishers = FireExtinguisher.objects.filter(status='reserved', extinguisher_type=extinguisher.extinguisher_type)

    context = {
        'extinguisher': extinguisher,
        'service_records': service_records,
        'service_record_form': service_record_form,
        'location_form': location_form,
        'replacement_form': replacement_form,
        'sections': sections,
        'machines': machines,
        'reserved_extinguishers': reserved_extinguishers,
    }
    return render(request, 'fire_extinguisher_management/extinguisher_detail.html', context)

@login_required
def extinguisher_create(request):
    if request.method == 'POST':
        form = FireExtinguisherForm(request.POST)
        if form.is_valid():
            extinguisher = form.save()
            messages.success(request, 'کپسول با موفقیت ثبت شد.')
            return redirect('extinguisher_detail', pk=extinguisher.pk)
    else:
        form = FireExtinguisherForm()
    
    return render(request, 'fire_extinguisher_management/extinguisher_form.html', {'form': form})

@login_required
def extinguisher_edit(request, pk):
    extinguisher = get_object_or_404(FireExtinguisher, pk=pk)
    if request.method == 'POST':
        form = FireExtinguisherForm(request.POST, instance=extinguisher)
        if form.is_valid():
            form.save()
            messages.success(request, 'اطلاعات کپسول با موفقیت به‌روزرسانی شد.')
            return redirect('extinguisher_detail', pk=extinguisher.pk)
    else:
        form = FireExtinguisherForm(instance=extinguisher)
    
    return render(request, 'fire_extinguisher_management/extinguisher_form.html', {'form': form})

@login_required
def extinguisher_replace(request):
    if request.method == 'POST':
        form = FireExtinguisherReplacementForm(request.POST)
        if form.is_valid():
            old_extinguisher = form.cleaned_data['old_extinguisher']
            new_extinguisher = form.cleaned_data['new_extinguisher']
            notes = form.cleaned_data['replacement_notes']

            # Update old extinguisher
            old_extinguisher.status = 'DISPOSED'
            old_extinguisher.replaced_by_extinguisher = new_extinguisher
            old_extinguisher.save()

            # Update new extinguisher
            new_extinguisher.status = 'OPERATIONAL'
            new_extinguisher.replaces_extinguisher = old_extinguisher
            new_extinguisher.location_type = old_extinguisher.location_type
            new_extinguisher.location_section = old_extinguisher.location_section
            new_extinguisher.location_machine = old_extinguisher.location_machine
            new_extinguisher.save()

            # Create service record
            ServiceRecord.objects.create(
                extinguisher=new_extinguisher,
                service_date=timezone.now().date(),
                service_type='REPLACEMENT',
                performed_by_user=request.user,
                outcome='COMPLETED',
                actions_taken=f'جایگزینی کپسول {old_extinguisher.serial_tag} با {new_extinguisher.serial_tag}',
                notes=notes
            )

            messages.success(request, 'عملیات جایگزینی با موفقیت انجام شد.')
            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'message': 'عملیات جایگزینی با موفقیت انجام شد.', 'new_pk': new_extinguisher.pk}, status=200)
            return redirect('fire_extinguisher_management:extinguisher_detail', pk=new_extinguisher.pk)
    else:
        form = FireExtinguisherReplacementForm()
    
    return render(request, 'fire_extinguisher_management/extinguisher_replace.html', {'form': form})

@login_required
def extinguisher_change_location(request, pk):
    extinguisher = get_object_or_404(FireExtinguisher, pk=pk)
    if request.method == 'POST':
        form = FireExtinguisherLocationForm(request.POST)
        if form.is_valid():
            extinguisher.location_type = form.cleaned_data['location_type']
            if form.cleaned_data['location_type'] == 'section':
                extinguisher.location_section = form.cleaned_data['location_section']
                extinguisher.location_machine = None
            else:
                extinguisher.location_machine = form.cleaned_data['location_machine']
                extinguisher.location_section = None
            extinguisher.save()

            # Create service record
            ServiceRecord.objects.create(
                extinguisher=extinguisher,
                service_date=timezone.now().date(),
                service_type='MAINTENANCE',
                performed_by_user=request.user,
                outcome='COMPLETED',
                actions_taken='تغییر مکان کپسول',
                notes=form.cleaned_data['notes']
            )

            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'message': 'مکان کپسول با موفقیت تغییر کرد.'}, status=200)
            messages.success(request, 'مکان کپسول با موفقیت تغییر کرد.')
            return redirect('fire_extinguisher_management:extinguisher_detail', pk=extinguisher.pk)
    else:
        form = FireExtinguisherLocationForm(initial={
            'location_type': extinguisher.location_type,
            'location_section': extinguisher.location_section,
            'location_machine': extinguisher.location_machine
        })
    
    return render(request, 'fire_extinguisher_management/extinguisher_location.html', {
        'form': form,
        'extinguisher': extinguisher
    })

@login_required
def service_record_create(request, extinguisher_pk):
    extinguisher = get_object_or_404(FireExtinguisher, pk=extinguisher_pk)
    if request.method == 'POST':
        form = ServiceRecordForm(request.POST)
        if form.is_valid():
            service_record = form.save(commit=False)
            service_record.extinguisher = extinguisher
            service_record.performed_by_user = request.user
            service_record.save()

            # Update extinguisher status and dates
            extinguisher.last_serviced_date = service_record.service_date
            extinguisher.next_scheduled_service_date = extinguisher.calculate_next_service_date()
            extinguisher.pressure_test_due_date = extinguisher.calculate_pressure_test_date()

            if service_record.outcome == 'FAIL':
                extinguisher.status = 'NEEDS_MAINTENANCE'
            elif service_record.outcome == 'PASS':
                extinguisher.status = 'OPERATIONAL'

            extinguisher.save()

            # AJAX response on successful creation
            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'message': 'سابقه سرویس با موفقیت ثبت شد.'}, status=201)
            messages.success(request, 'سابقه سرویس با موفقیت ثبت شد.')
            return redirect('fire_extinguisher_management:extinguisher_detail', pk=extinguisher.pk)
        else:
            # Return structured errors for AJAX clients
            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                errors = {k: [e['message'] for e in v] for k, v in form.errors.get_json_data().items()}
                return JsonResponse({'success': False, 'errors': errors, 'message': 'فرم معتبر نیست'}, status=400)
            messages.error(request, 'فرم نامعتبر است.')
            return render(request, 'fire_extinguisher_management/service_record_form.html', {
                'form': form,
                'extinguisher': extinguisher
            })
    else:
        form = ServiceRecordForm(initial={'extinguisher': extinguisher})
    
    return render(request, 'fire_extinguisher_management/service_record_form.html', {
        'form': form,
        'extinguisher': extinguisher
    })

@login_required
def notification_list(request):
    notifications = request.user.notifications.order_by('-created_at')
    return render(request, 'fire_extinguisher_management/notification_list.html', {
        'notifications': notifications
    })

@login_required
def mark_notification_read(request, pk):
    notification = get_object_or_404(Notification, pk=pk, user=request.user)
    notification.is_read = True
    notification.save()
    return JsonResponse({'status': 'success'})

@login_required
def extinguisher_details_api(request, pk):
    extinguisher = get_object_or_404(FireExtinguisher, pk=pk)
    
    # Get available extinguishers for replacement
    available_extinguishers = FireExtinguisher.objects.filter(
        status='RESERVED',
        extinguisher_type=extinguisher.extinguisher_type
    ).values('id', 'serial_tag', 'extinguisher_type__name')
    
    # Format the data for the API response
    available_extinguishers = [
        {
            'id': e['id'],
            'serial_tag': e['serial_tag'],
            'type': e['extinguisher_type__name']
        }
        for e in available_extinguishers
    ]
    
    # Include location information
    location = {
        'content_type': extinguisher.content_type_id,
        'object_id': extinguisher.object_id
    }
    
    return JsonResponse({
        'available_extinguishers': available_extinguishers,
        'location': location
    })

@login_required
def download_excel_template(request):
    """دانلود فایل نمونه Excel برای import کپسول‌ها"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Fire Extinguishers"

    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Headers
    headers = [
        'شماره سریال*',
        'نوع کپسول*',
        'سازنده*',
        'شماره مدل*',
        'ظرفیت (عدد)*',
        'واحد ظرفیت (kg یا L)*',
        'تاریخ خرید (۱۴۰۳/۱۰/۲۵)*',
        'تاریخ ساخت (۱۴۰۳/۱۰/۲۵)',
        'تاریخ بهره‌برداری (۱۴۰۳/۱۰/۲۵)*',
        'عمر مفید (سال)*',
        'وضعیت*',
        'نوع مکان (section یا machine)*',
        'نام بخش',
        'نام دستگاه',
        'توضیحات'
    ]
    
    ws.append(headers)
    
    # Apply header styles
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 18

    # Add sample data
    sample_data = [
        [
            'FE-001',
            'پودر و گاز ABC',
            'شرکت نمونه',
            'ABC-6KG',
            '6',
            'kg',
            '۱۴۰۳/۱۰/۲۵',
            '۱۴۰۳/۰۹/۱۰',
            '۱۴۰۳/۱۱/۱۲',
            '10',
            'operational',
            'section',
            'انبار مواد',
            '',
            'کپسول نمونه'
        ],
        [
            'FE-002',
            'CO2',
            'شرکت نمونه',
            'CO2-5KG',
            '5',
            'kg',
            '۱۴۰۳/۱۰/۳۰',
            '۱۴۰۳/۰۹/۲۰',
            '۱۴۰۳/۱۱/۱۵',
            '10',
            'reserved',
            'machine',
            '',
            'بیل مکانیکی 01',
            ''
        ]
    ]
    
    for row_data in sample_data:
        ws.append(row_data)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=col_num)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Add instructions sheet
    ws2 = wb.create_sheet("راهنما")
    instructions = [
        ["راهنمای استفاده از فایل Excel"],
        [""],
        ["فیلدهای دارای علامت * الزامی هستند"],
        [""],
        ["نوع کپسول: باید دقیقاً با نام یکی از انواع موجود در سیستم مطابقت داشته باشد"],
        [""],
        ["واحد ظرفیت: فقط kg یا L"],
        [""],
        ["تاریخ‌ها: به فرمت جلالی YYYY/MM/DD مانند ۱۴۰۳/۱۰/۲۵ (با اعداد فارسی یا انگلیسی)"],
        [""],
        ["وضعیت: یکی از موارد زیر"],
        ["  - operational: عملیاتی"],
        ["  - needs_maintenance: نیازمند تعمیر"],
        ["  - expired: منقضی شده"],
        ["  - reserved: رزرو شده"],
        ["  - under_test: در انتظار تست"],
        ["  - disposed: مستهلک شده"],
        [""],
        ["نوع مکان: section (بخش) یا machine (دستگاه)"],
        [""],
        ["نام بخش: اگر نوع مکان section باشد، نام دقیق بخش را وارد کنید"],
        [""],
        ["نام دستگاه: اگر نوع مکان machine باشد، نام دقیق دستگاه را وارد کنید"],
    ]
    
    for row in instructions:
        ws2.append(row)
        if row and row[0].startswith("راهنما"):
            cell = ws2.cell(row=ws2.max_row, column=1)
            cell.font = Font(bold=True, size=14)
            cell.fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")

    ws2.column_dimensions['A'].width = 70

    # Prepare response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=fire_extinguishers_template.xlsx'
    wb.save(response)
    
    return response

@login_required
def import_excel(request):
    """صفحه آپلود و import فایل Excel"""
    if request.method == 'POST':
        form = ExcelImportForm(request.POST, request.FILES)
        if form.is_valid():
            excel_file = request.FILES['excel_file']
            
            try:
                wb = load_workbook(excel_file, data_only=True)
                ws = wb.active
                
                errors = []
                success_count = 0
                row_num = 2  # Start from row 2 (skip header)
                
                # Get lookup data
                extinguisher_types = {t.name: t for t in FireExtinguisherType.objects.all()}
                
                from django.apps import apps
                LocationSection = apps.get_model('anomalis', 'LocationSection')
                MiningMachine = apps.get_model('BaseInfo', 'MiningMachine')
                
                sections = {s.name: s for s in LocationSection.objects.all()}
                machines = {m.name: m for m in MiningMachine.objects.all()}
                
                for row in ws.iter_rows(min_row=2, values_only=True):
                    # Skip empty rows
                    if not any(row):
                        continue
                    
                    try:
                        serial_tag = str(row[0]).strip() if row[0] else None
                        type_name = str(row[1]).strip() if row[1] else None
                        manufacturer = str(row[2]).strip() if row[2] else None
                        model_number = str(row[3]).strip() if row[3] else None
                        capacity_value = row[4]
                        capacity_unit = str(row[5]).strip() if row[5] else None
                        purchase_date = row[6]
                        manufacture_date = row[7]
                        commission_date = row[8]
                        lifespan = row[9]
                        status = str(row[10]).strip() if row[10] else None
                        location_type = str(row[11]).strip() if row[11] else None
                        section_name = str(row[12]).strip() if row[12] else None
                        machine_name = str(row[13]).strip() if row[13] else None
                        notes = str(row[14]).strip() if row[14] else ""
                        
                        # Validate required fields
                        if not all([serial_tag, type_name, manufacturer, model_number, 
                                   capacity_value, capacity_unit, purchase_date, 
                                   commission_date, lifespan, status, location_type]):
                            errors.append(f"ردیف {row_num}: فیلدهای الزامی خالی است")
                            row_num += 1
                            continue
                        
                        # Check if serial_tag already exists
                        if FireExtinguisher.objects.filter(serial_tag=serial_tag).exists():
                            errors.append(f"ردیف {row_num}: شماره سریال {serial_tag} تکراری است")
                            row_num += 1
                            continue
                        
                        # Get extinguisher type
                        ext_type = extinguisher_types.get(type_name)
                        if not ext_type:
                            errors.append(f"ردیف {row_num}: نوع کپسول '{type_name}' یافت نشد")
                            row_num += 1
                            continue
                        
                        # Validate capacity_unit
                        if capacity_unit not in ['kg', 'L']:
                            errors.append(f"ردیف {row_num}: واحد ظرفیت باید kg یا L باشد")
                            row_num += 1
                            continue
                        
                        # Validate status
                        valid_statuses = ['operational', 'needs_maintenance', 'expired', 
                                        'reserved', 'under_test', 'disposed']
                        if status not in valid_statuses:
                            errors.append(f"ردیف {row_num}: وضعیت '{status}' نامعتبر است")
                            row_num += 1
                            continue
                        
                        # Parse dates - support both Jalali (Persian) and Gregorian
                        def parse_date(date_val):
                            if isinstance(date_val, datetime):
                                return date_val.date()
                            elif isinstance(date_val, str):
                                date_str = str(date_val).strip()
                                if not date_str:
                                    return None
                                
                                # Convert Persian digits to English
                                persian_to_english = str.maketrans('۰۱۲۳۴۵۶۷۸۹', '0123456789')
                                date_str = date_str.translate(persian_to_english)
                                
                                # Try parsing as Jalali date (YYYY/MM/DD or YYYY-MM-DD)
                                try:
                                    date_str = date_str.replace('/', '-')
                                    parts = date_str.split('-')
                                    if len(parts) == 3:
                                        year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
                                        # Check if it's Jalali (year > 1300 indicates Jalali)
                                        if year > 1300:
                                            # Convert Jalali to Gregorian using jdatetime
                                            try:
                                                import jdatetime
                                                jdate = jdatetime.date(year, month, day)
                                                return jdate.togregorian()
                                            except:
                                                return None
                                        else:
                                            # It's already Gregorian
                                            return datetime(year, month, day).date()
                                except:
                                    pass
                                
                                # Try standard date parsing
                                try:
                                    return datetime.strptime(date_str, '%Y-%m-%d').date()
                                except:
                                    return None
                            return None
                        
                        purchase_date = parse_date(purchase_date)
                        manufacture_date = parse_date(manufacture_date)
                        commission_date = parse_date(commission_date)
                        
                        if not purchase_date or not commission_date:
                            errors.append(f"ردیف {row_num}: فرمت تاریخ نامعتبر است")
                            row_num += 1
                            continue
                        
                        # Handle location
                        location_section = None
                        location_machine = None
                        
                        if location_type == 'section':
                            if not section_name:
                                errors.append(f"ردیف {row_num}: برای مکان از نوع بخش، نام بخش الزامی است")
                                row_num += 1
                                continue
                            location_section = sections.get(section_name)
                            if not location_section:
                                errors.append(f"ردیف {row_num}: بخش '{section_name}' یافت نشد")
                                row_num += 1
                                continue
                        elif location_type == 'machine':
                            if not machine_name:
                                errors.append(f"ردیف {row_num}: برای مکان از نوع دستگاه، نام دستگاه الزامی است")
                                row_num += 1
                                continue
                            location_machine = machines.get(machine_name)
                            if not location_machine:
                                errors.append(f"ردیف {row_num}: دستگاه '{machine_name}' یافت نشد")
                                row_num += 1
                                continue
                        else:
                            errors.append(f"ردیف {row_num}: نوع مکان باید section یا machine باشد")
                            row_num += 1
                            continue
                        
                        # Create extinguisher
                        FireExtinguisher.objects.create(
                            serial_tag=serial_tag,
                            extinguisher_type=ext_type,
                            manufacturer=manufacturer,
                            model_number=model_number,
                            capacity_value=capacity_value,
                            capacity_unit=capacity_unit,
                            purchase_date=purchase_date,
                            manufacture_date=manufacture_date,
                            commission_date=commission_date,
                            expected_lifespan_years=lifespan,
                            status=status,
                            location_type=location_type,
                            location_section=location_section,
                            location_machine=location_machine,
                            notes=notes
                        )
                        success_count += 1
                        
                    except Exception as e:
                        errors.append(f"ردیف {row_num}: خطا در پردازش - {str(e)}")
                    
                    row_num += 1
                
                # Show results
                if success_count > 0:
                    messages.success(request, f"{success_count} کپسول با موفقیت ثبت شد.")
                
                if errors:
                    for error in errors[:10]:  # Show first 10 errors
                        messages.error(request, error)
                    if len(errors) > 10:
                        messages.warning(request, f"و {len(errors) - 10} خطای دیگر...")
                
                if success_count > 0 and not errors:
                    return redirect('fire_extinguisher_management:extinguisher_list')
                    
            except Exception as e:
                messages.error(request, f"خطا در خواندن فایل Excel: {str(e)}")
    else:
        form = ExcelImportForm()
    
    return render(request, 'fire_extinguisher_management/import_excel.html', {'form': form})

@login_required
def import_excel_ajax(request):
    """Import Excel via AJAX"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'درخواست نامعتبر است'}, status=400)
    
    # Debug logging
    print(f"FILES: {request.FILES}")
    print(f"POST: {request.POST}")
    
    if not request.FILES.get('excel_file'):
        return JsonResponse({'success': False, 'message': 'فایل Excel انتخاب نشده است'}, status=400)
    
    excel_file = request.FILES['excel_file']
    print(f"File name: {excel_file.name}")
    print(f"File size: {excel_file.size}")
    
    # Validate file
    if not excel_file.name.endswith('.xlsx'):
        return JsonResponse({'success': False, 'message': 'فقط فایل‌های xlsx پذیرفته می‌شود'}, status=400)
    
    if excel_file.size > 5 * 1024 * 1024:
        return JsonResponse({'success': False, 'message': 'حجم فایل نباید بیشتر از 5 مگابایت باشد'}, status=400)
    
    try:
        print("Starting to load workbook...")
        wb = load_workbook(excel_file, data_only=True)
        print("Workbook loaded successfully")
        ws = wb.active
        print(f"Active sheet: {ws.title}")
        
        errors = []
        success_count = 0
        row_num = 2  # Start from row 2 (skip header)
        
        # Get lookup data
        print("Loading lookup data...")
        extinguisher_types = {t.name: t for t in FireExtinguisherType.objects.all()}
        print(f"Loaded {len(extinguisher_types)} extinguisher types")
        
        from django.apps import apps
        LocationSection = apps.get_model('anomalis', 'LocationSection')
        MiningMachine = apps.get_model('BaseInfo', 'MiningMachine')
        
        sections = {s.section: s for s in LocationSection.objects.all()}
        machines = {m.workshop_code: m for m in MiningMachine.objects.all()}
        print(f"Loaded {len(sections)} sections and {len(machines)} machines")
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            # Skip empty rows
            if not any(row):
                continue
            
            try:
                serial_tag = str(row[0]).strip() if row[0] else None
                type_name = str(row[1]).strip() if row[1] else None
                manufacturer = str(row[2]).strip() if row[2] else None
                model_number = str(row[3]).strip() if row[3] else None
                capacity_value = row[4]
                capacity_unit = str(row[5]).strip() if row[5] else None
                purchase_date = row[6]
                manufacture_date = row[7]
                commission_date = row[8]
                lifespan = row[9]
                status = str(row[10]).strip() if row[10] else None
                location_type = str(row[11]).strip() if row[11] else None
                section_name = str(row[12]).strip() if row[12] else None
                machine_name = str(row[13]).strip() if row[13] else None
                notes = str(row[14]).strip() if row[14] else ""
                
                # Validate required fields
                if not all([serial_tag, type_name, manufacturer, model_number, 
                           capacity_value, capacity_unit, purchase_date, 
                           commission_date, lifespan, status, location_type]):
                    errors.append(f"ردیف {row_num}: فیلدهای الزامی خالی است")
                    row_num += 1
                    continue
                
                # Check if serial_tag already exists
                if FireExtinguisher.objects.filter(serial_tag=serial_tag).exists():
                    errors.append(f"ردیف {row_num}: شماره سریال {serial_tag} تکراری است")
                    row_num += 1
                    continue
                
                # Get or create extinguisher type
                ext_type = extinguisher_types.get(type_name)
                if not ext_type:
                    # Create with default values
                    ext_type = FireExtinguisherType.objects.create(
                        name=type_name,
                        agent=type_name,  # Use type name as agent
                        use_class="ABC",  # Default class
                        inspection_interval_months=6,  # Default 6 months
                        service_interval_years=1,  # Default 1 year
                        pressure_test_interval_years=5,  # Default 5 years
                        notes="ایجاد شده خودکار از ایمپورت اکسل"
                    )
                    extinguisher_types[type_name] = ext_type
                    print(f"Created new extinguisher type: {type_name}")
                
                # Validate capacity_unit
                if capacity_unit not in ['kg', 'L']:
                    errors.append(f"ردیف {row_num}: واحد ظرفیت باید kg یا L باشد")
                    row_num += 1
                    continue
                
                # Validate status
                valid_statuses = ['operational', 'needs_maintenance', 'expired', 
                                'reserved', 'under_test', 'disposed']
                if status not in valid_statuses:
                    errors.append(f"ردیف {row_num}: وضعیت '{status}' نامعتبر است")
                    row_num += 1
                    continue
                
                # Parse dates - support both Jalali (Persian) and Gregorian
                def parse_date(date_val):
                    if isinstance(date_val, datetime):
                        return date_val.date()
                    elif isinstance(date_val, str):
                        date_str = str(date_val).strip()
                        if not date_str:
                            return None
                        
                        # Convert Persian digits to English
                        persian_to_english = str.maketrans('۰۱۲۳۴۵۶۷۸۹', '0123456789')
                        date_str = date_str.translate(persian_to_english)
                        
                        # Try parsing as Jalali date (YYYY/MM/DD or YYYY-MM-DD)
                        try:
                            date_str = date_str.replace('/', '-')
                            parts = date_str.split('-')
                            if len(parts) == 3:
                                year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
                                # Check if it's Jalali (year > 1300 indicates Jalali)
                                if year > 1300:
                                    # Convert Jalali to Gregorian using jdatetime
                                    try:
                                        import jdatetime
                                        jdate = jdatetime.date(year, month, day)
                                        return jdate.togregorian()
                                    except:
                                        return None
                                else:
                                    # It's already Gregorian
                                    return datetime(year, month, day).date()
                        except:
                            pass
                        
                        # Try standard date parsing
                        try:
                            return datetime.strptime(date_str, '%Y-%m-%d').date()
                        except:
                            return None
                    return None
                
                purchase_date = parse_date(purchase_date)
                manufacture_date = parse_date(manufacture_date)
                commission_date = parse_date(commission_date)
                
                if not purchase_date or not commission_date:
                    errors.append(f"ردیف {row_num}: فرمت تاریخ نامعتبر است")
                    row_num += 1
                    continue
                
                # Handle location
                location_section = None
                location_machine = None
                
                if location_type == 'section':
                    if not section_name:
                        errors.append(f"ردیف {row_num}: برای مکان از نوع بخش، نام بخش الزامی است")
                        row_num += 1
                        continue
                    location_section = sections.get(section_name)
                    if not location_section:
                        errors.append(f"ردیف {row_num}: بخش '{section_name}' یافت نشد")
                        row_num += 1
                        continue
                elif location_type == 'machine':
                    if not machine_name:
                        errors.append(f"ردیف {row_num}: برای مکان از نوع دستگاه، نام دستگاه الزامی است")
                        row_num += 1
                        continue
                    location_machine = machines.get(machine_name)
                    if not location_machine:
                        errors.append(f"ردیف {row_num}: دستگاه '{machine_name}' یافت نشد")
                        row_num += 1
                        continue
                else:
                    errors.append(f"ردیف {row_num}: نوع مکان باید section یا machine باشد")
                    row_num += 1
                    continue
                
                # Create extinguisher
                FireExtinguisher.objects.create(
                    serial_tag=serial_tag,
                    extinguisher_type=ext_type,
                    manufacturer=manufacturer,
                    model_number=model_number,
                    capacity_value=capacity_value,
                    capacity_unit=capacity_unit,
                    purchase_date=purchase_date,
                    manufacture_date=manufacture_date,
                    commission_date=commission_date,
                    expected_lifespan_years=lifespan,
                    status=status,
                    location_type=location_type,
                    location_section=location_section,
                    location_machine=location_machine,
                    notes=notes
                )
                success_count += 1
                
            except Exception as e:
                errors.append(f"ردیف {row_num}: خطا در پردازش - {str(e)}")
            
            row_num += 1
        
        return JsonResponse({
            'success': True,
            'success_count': success_count,
            'errors': errors
        })
                
    except Exception as e:
        import traceback
        print("="*50)
        print("ERROR in import_excel_ajax:")
        print(traceback.format_exc())
        print("="*50)
        return JsonResponse({
            'success': False,
            'message': f'خطا در خواندن فایل Excel: {str(e)}'
        }, status=400)
