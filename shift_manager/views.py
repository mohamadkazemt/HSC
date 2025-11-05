from django.shortcuts import render
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.decorators import login_required, user_passes_test
from django.views.decorators.http import require_http_methods
import datetime
from django.utils import timezone
from shift_manager.utils import get_shift_for_date
from shift_manager.forms import InitialShiftSetupForm
from shift_manager.models import InitialShiftSetup, SHIFT_CHOICES
from core.models import SiteSettings
import jdatetime
import logging
import calendar
from django.template.defaultfilters import date
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
try:
    from openpyxl.drawing.image import Image as XLImage
except ImportError:
    XLImage = None

# تنظیمات لاگینگ
logger = logging.getLogger(__name__)


def shift_calendar_view(request):
    today = jdatetime.datetime.now()
    year = request.GET.get('year', today.year)
    month = request.GET.get('month', today.month)

    # اسامی ماه‌های شمسی
    months = {
        1: 'فروردین',
        2: 'اردیبهشت',
        3: 'خرداد',
        4: 'تیر',
        5: 'مرداد',
        6: 'شهریور',
        7: 'مهر',
        8: 'آبان',
        9: 'آذر',
        10: 'دی',
        11: 'بهمن',
        12: 'اسفند'
    }

    try:
        year = int(year)
        month = int(month)

        if year < 1 or month < 1 or month > 12:
            logger.error(f"Invalid year or month: year={year}, month={month}")
            raise ValueError("Year or month is out of valid range.")

        # محاسبه روز آخر ماه شمسی
        if month <= 6:
            last_day = 31
        elif month <= 11:
            last_day = 30
        else:  # month == 12
            if jdatetime.date(year, month, 1).isleap():
                last_day = 30
            else:
                last_day = 29

        # تولید لیست تاریخ‌های شمسی برای ماه انتخاب شده
        calendar_data = []
        first_day_of_month = jdatetime.date(year, month, 1)
        weekday = first_day_of_month.weekday()  # 0: شنبه, 1: یکشنبه, ..., 6: جمعه

        # اضافه کردن روزهای خالی برای روزهای قبل از اول ماه
        first_week = [None] * weekday

        # اضافه کردن روزهای ماه
        days_in_month = []
        for day in range(1, last_day + 1):
            days_in_month.append(jdatetime.date(year, month, day))

        # ترکیب روزهای خالی و روزهای ماه
        first_week.extend(days_in_month[:7 - weekday])
        calendar_data.append(first_week)

        # پر کردن بقیه هفته‌ها
        remaining_days = days_in_month[7 - weekday:]
        while remaining_days:
            week = remaining_days[:7]
            calendar_data.append(week)
            remaining_days = remaining_days[7:]

        # محاسبه شیفت‌ها برای کل روزهای ماه
        shift_data = {}
        for week in calendar_data:
            for j_date in week:
                if j_date:
                    g_date = j_date.togregorian()
                    shifts = get_shift_for_date(g_date)
                    # تبدیل تاریخ به فرمت مناسب
                    formatted_date = j_date.strftime('%Y-%m-%d')
                    shift_data[formatted_date] = shifts

    except ValueError as e:
        logger.error(f"Error processing date: {e}")
        return render(request, 'shift_manager/error.html', {'error': 'تاریخ نامعتبر'})

    # Get current Jalali date for comparison
    today_jalali = jdatetime.datetime.now()
    
    context = {
        'year': year,
        'month': month,
        'months': months,
        'shift_data': shift_data,
        'calendar': calendar_data,
        'SHIFT_CHOICES': SHIFT_CHOICES,
        'today_jalali': today_jalali.date(),
    }

    return render(request, 'shift_manager/shift_calendar.html', context)


@login_required
@user_passes_test(lambda u: u.is_superuser)
@require_http_methods(["POST"])
def initial_shift_setup_ajax(request):
    """AJAX view for creating/updating InitialShiftSetup"""
    try:
        # Get the latest setup if exists
        try:
            latest_setup = InitialShiftSetup.objects.latest('start_date')
            form = InitialShiftSetupForm(request.POST, instance=latest_setup)
        except InitialShiftSetup.DoesNotExist:
            form = InitialShiftSetupForm(request.POST)

        if form.is_valid():
            form.save()
            return JsonResponse({
                'success': True,
                'message': 'تنظیمات اولیه شیفت با موفقیت ذخیره شد. تقویم به‌روزرسانی می‌شود.'
            })
        else:
            # Return form errors
            errors = {}
            for field, field_errors in form.errors.items():
                errors[field] = field_errors
            return JsonResponse({
                'success': False,
                'errors': errors,
                'message': 'لطفاً خطاهای فرم را بررسی کنید.'
            }, status=400)
    except Exception as e:
        logger.error(f"Error in initial_shift_setup_ajax: {e}")
        return JsonResponse({
            'success': False,
            'message': f'خطا در ذخیره تنظیمات: {str(e)}'
        }, status=500)


@login_required
def export_shift_calendar_excel(request):
    """Export shift calendar to Excel with company logo and name in header"""
    today = jdatetime.datetime.now()
    year = request.GET.get('year', today.year)
    month = request.GET.get('month', today.month)

    # اسامی ماه‌های شمسی
    months = {
        1: 'فروردین',
        2: 'اردیبهشت',
        3: 'خرداد',
        4: 'تیر',
        5: 'مرداد',
        6: 'شهریور',
        7: 'مهر',
        8: 'آبان',
        9: 'آذر',
        10: 'دی',
        11: 'بهمن',
        12: 'اسفند'
    }

    try:
        year = int(year)
        month = int(month)

        if year < 1 or month < 1 or month > 12:
            raise ValueError("Year or month is out of valid range.")

        # محاسبه روز آخر ماه شمسی
        if month <= 6:
            last_day = 31
        elif month <= 11:
            last_day = 30
        else:  # month == 12
            if jdatetime.date(year, month, 1).isleap():
                last_day = 30
            else:
                last_day = 29

        # تولید لیست تاریخ‌های شمسی برای ماه انتخاب شده
        days_in_month = []
        for day in range(1, last_day + 1):
            days_in_month.append(jdatetime.date(year, month, day))

        # محاسبه شیفت‌ها برای کل روزهای ماه
        shift_data = {}
        for j_date in days_in_month:
            g_date = j_date.togregorian()
            shifts = get_shift_for_date(g_date)
            formatted_date = j_date.strftime('%Y-%m-%d')
            shift_data[formatted_date] = shifts

        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        month_name = months.get(month, 'نامشخص')
        ws.title = f"تقویم شیفت {month_name} {year}"

        # RTL
        ws.sheet_properties.rightToLeft = True

        # Styles
        title_font = Font(name='B Nazanin', size=20, bold=True)
        subtitle_font = Font(name='B Nazanin', size=14, bold=True)
        header_font = Font(name='B Nazanin', size=12, bold=True)
        data_font = Font(name='B Nazanin', size=11)

        title_fill = PatternFill(start_color='E8E8E8', end_color='E8E8E8', fill_type='solid')
        header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
        day_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
        
        # Shift color fills
        shift_colors = {
            'روزکار اول': PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid'),
            'روزکار دوم': PatternFill(start_color='FDE68A', end_color='FDE68A', fill_type='solid'),
            'عصرکار اول': PatternFill(start_color='FED7AA', end_color='FED7AA', fill_type='solid'),
            'عصرکار دوم': PatternFill(start_color='FDBA74', end_color='FDBA74', fill_type='solid'),
            'شب کار اول': PatternFill(start_color='BFDBFE', end_color='BFDBFE', fill_type='solid'),
            'شب کار دوم': PatternFill(start_color='93C5FD', end_color='93C5FD', fill_type='solid'),
            'OFF اول': PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid'),
            'OFF دوم': PatternFill(start_color='D1D5DB', end_color='D1D5DB', fill_type='solid'),
        }

        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        right_alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)

        # Site settings (logo + name)
        site = SiteSettings.objects.first()
        site_name = site.site_name if site and site.site_name else 'شرکت'
        title_text = f"تقویم شیفت {month_name} {year} - {site_name}"

        # Title row (merge across A:H)
        title_cell = ws.cell(row=1, column=1, value=title_text)
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = center_alignment
        ws.merge_cells('A1:H1')
        ws.row_dimensions[1].height = 45

        # Logo (if available)
        if site and getattr(site, 'company_logo', None) and hasattr(site.company_logo, 'path') and XLImage:
            try:
                logo_img = XLImage(site.company_logo.path)
                logo_img.height = 60
                logo_img.width = 60
                # Anchor at A1 (top-right in RTL appearance)
                ws.add_image(logo_img, 'A1')
            except Exception as e:
                logger.error(f"Error adding logo to Excel: {e}")

        # Headers
        current_row = 3
        headers = ['تاریخ', 'روز هفته', 'گروه A', 'گروه B', 'گروه C', 'گروه D']
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

        # Set column widths
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 18

        # Weekday names
        weekday_names = ['شنبه', 'یکشنبه', 'دوشنبه', 'سه‌شنبه', 'چهارشنبه', 'پنج‌شنبه', 'جمعه']

        # Data rows
        current_row = 4
        for j_date in days_in_month:
            formatted_date = j_date.strftime('%Y-%m-%d')
            shifts = shift_data.get(formatted_date, {})
            weekday = weekday_names[j_date.weekday()]

            # Date
            date_cell = ws.cell(row=current_row, column=1, value=f"{j_date.day} {month_name} {year}")
            date_cell.font = data_font
            date_cell.fill = day_fill
            date_cell.alignment = right_alignment
            date_cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Weekday
            weekday_cell = ws.cell(row=current_row, column=2, value=weekday)
            weekday_cell.font = data_font
            weekday_cell.fill = day_fill
            weekday_cell.alignment = center_alignment
            weekday_cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Group shifts
            groups = ['A', 'B', 'C', 'D']
            for group_idx, group in enumerate(groups, start=3):
                shift = shifts.get(group, '')
                shift_cell = ws.cell(row=current_row, column=group_idx, value=shift)
                shift_cell.font = data_font
                shift_cell.alignment = center_alignment
                
                # Apply color based on shift type
                if shift in shift_colors:
                    shift_cell.fill = shift_colors[shift]
                
                shift_cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )

            current_row += 1

        # Create HTTP response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"تقویم_شیفت_{month_name}_{year}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        wb.save(response)
        return response

    except Exception as e:
        logger.error(f"Error exporting shift calendar to Excel: {e}")
        return HttpResponse(f"خطا در تولید فایل Excel: {str(e)}", status=500)