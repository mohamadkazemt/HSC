import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import os

# ایجاد یک فایل اکسل جدید
wb = openpyxl.Workbook()
ws = wb.active

# تعریف داده‌ها
headers = [
    'متن سوال',
    'محدوده',
    'نوع سوال',
    'نوع ماشین',
    'بخش مکانی',
    'گزینه‌ها',
    'گزینه‌های غیرقابل قبول',
    'نوع آنومالی',
    'حوزه HSE',
    'شرح آنومالی پیش‌فرض',
    'اولویت',
    'اقدام اصلاحی پیش‌فرض'
]

data = [
    [
        'آیا دستگاه دارای محافظ‌های ایمنی مناسب است؟',
        'ماشین',
        'گزینه‌ای',
        'دستگاه برش',
        '',
        'کاملاً سالم و مناسب,نیاز به تعمیر جزئی,نیاز به تعویض,فاقد محافظ',
        'نیاز به تعویض,فاقد محافظ',
        'نقص فنی',
        'S',
        'محافظ‌های ایمنی نصب نشده یا معیوب هستند',
        'متوسط',
        'نصب یا تعمیر محافظ‌های ایمنی'
    ],
    [
        'آیا تجهیزات اطفاء حریق در دسترس و قابل استفاده هستند؟',
        'مکان',
        'گزینه‌ای',
        '',
        'سالن تولید',
        'در دسترس و سالم,در دسترس ولی نیاز به شارژ,خارج از دسترس,معیوب,فاقد تجهیزات',
        'خارج از دسترس,معیوب,فاقد تجهیزات',
        'نقص تجهیزات',
        'S',
        'تجهیزات اطفاء حریق در دسترس نیستند یا معیوب هستند',
        'زیاد',
        'تامین و نصب تجهیزات اطفاء حریق'
    ],
    [
        'آیا کارکنان از تجهیزات حفاظت فردی استفاده می‌کنند؟',
        'ماشین',
        'گزینه‌ای',
        'دستگاه جوش',
        '',
        'استفاده کامل و صحیح,استفاده ناقص,عدم استفاده,تجهیزات معیوب',
        'عدم استفاده,تجهیزات معیوب',
        'نقص تجهیزات',
        'H',
        'عدم استفاده از تجهیزات حفاظت فردی',
        'زیاد',
        'استفاده از تجهیزات حفاظت فردی'
    ],
    [
        'وضعیت سیستم تهویه مطبوع چگونه است؟',
        'مکان',
        'گزینه‌ای',
        '',
        'اتاق کنترل',
        'عملکرد عادی,نیاز به سرویس,عملکرد ضعیف,خراب,خاموش',
        'خراب,خاموش',
        'نقص فنی',
        'H',
        'سیستم تهویه مطبوع معیوب است',
        '3',
        'تعمیر سیستم تهویه مطبوع'
    ],
    [
        'وضعیت برچسب‌گذاری مواد شیمیایی چگونه است؟',
        'مکان',
        'گزینه‌ای',
        '',
        'انبار مواد شیمیایی',
        'کاملاً صحیح,نیاز به بروزرسانی,برچسب ناقص,بدون برچسب',
        'برچسب ناقص,بدون برچسب',
        'نقص تجهیزات',
        'E',
        'مواد شیمیایی بدون برچسب یا با برچسب نادرست',
        '2',
        'برچسب‌گذاری صحیح مواد شیمیایی'
    ]
]

# تنظیم عرض ستون‌ها
column_widths = [50, 15, 15, 20, 20, 50, 40, 20, 10, 40, 15, 40]
for i, width in enumerate(column_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

# نوشتن هدرها
header_font = Font(bold=True)
header_fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# نوشتن داده‌ها
for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# ایجاد مسیر فایل
file_path = os.path.join('checklist_app', 'static', 'checklist_app', 'sample_questions.xlsx')

# اطمینان از وجود پوشه‌ها
os.makedirs(os.path.dirname(file_path), exist_ok=True)

# ذخیره فایل
wb.save(file_path)
print(f"فایل نمونه در مسیر {file_path} ذخیره شد.") 