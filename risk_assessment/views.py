from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Count, Q
from django.db.models.functions import TruncMonth
from django.core.paginator import Paginator
from django.utils import timezone
from django.http import HttpResponse, JsonResponse
from datetime import timedelta
from accounts.models import Position, UserProfile
from .models import RiskAssessment
from .forms import RiskAssessmentForm, RiskReEvaluationForm, RiskFilterForm, ExcelImportForm
from .analytics import (
    get_high_frequency_hazards,
    get_incident_based_hazards,
    get_position_risk_suggestions,
    get_critical_locations,
    generate_risk_insights,
)
from .notifications import notify_risk_created, notify_risk_approved, notify_risk_rejected
from django.contrib.auth.decorators import user_passes_test
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from anomalis.models import AnomalyDescription
from hse_incidents.models import InjuryType
import jdatetime


@login_required
def risk_dashboard(request):
    """
    داشبورد اصلی - نمایش لیست مشاغل با تعداد ریسک‌های هر کدام
    """
    # فیلتر بر اساس تأیید (اگر مدیر HSE نیست)
    is_hse_mgr = request.user.groups.filter(name='مدیر HSE').exists()
    risk_filter = Q()
    if not is_hse_mgr:
        risk_filter = Q(risk_assessments__approval_status='approved')
    
    positions = Position.objects.annotate(
        total_risks=Count('risk_assessments', filter=risk_filter),
        high_risks=Count('risk_assessments', filter=Q(risk_assessments__risk_level='High') & risk_filter),
        medium_risks=Count('risk_assessments', filter=Q(risk_assessments__risk_level='Medium') & risk_filter),
        low_risks=Count('risk_assessments', filter=Q(risk_assessments__risk_level='Low') & risk_filter),
    ).filter(total_risks__gt=0).order_by('-high_risks', '-total_risks')
    
    # آمار کلی (فقط ریسک‌های تأیید شده)
    is_hse_mgr = request.user.groups.filter(name='مدیر HSE').exists()
    base_qs = RiskAssessment.objects.all()
    if not is_hse_mgr:
        base_qs = base_qs.filter(approval_status='approved')
    
    total_assessments = base_qs.count()
    high_risk_count = base_qs.filter(risk_level='High').count()
    medium_risk_count = base_qs.filter(risk_level='Medium').count()
    low_risk_count = base_qs.filter(risk_level='Low').count()
    
    # تحلیل‌های هوشمند
    high_frequency_hazards = get_high_frequency_hazards(days=90, min_count=3)
    incident_data = get_incident_based_hazards(months=12)
    critical_locations = get_critical_locations()
    
    context = {
        'positions': positions,
        'total_assessments': total_assessments,
        'high_risk_count': high_risk_count,
        'medium_risk_count': medium_risk_count,
        'low_risk_count': low_risk_count,
        'high_frequency_hazards': high_frequency_hazards[:5],  # 5 خطر پرتکرار
        'incident_stats': incident_data,
        'critical_locations': critical_locations[:5],  # 5 مکان حساس
    }
    
    return render(request, 'risk_assessment/dashboard.html', context)


@login_required
def risk_global_list(request):
    """لیست جامع تمام ریسک‌های ثبت شده با فیلتر و اولویت‌بندی"""
    # اگر کاربر مدیر HSE است، همه ریسک‌ها را ببین، وگرنه فقط تأیید شده‌ها
    is_hse_mgr = request.user.groups.filter(name='مدیر HSE').exists()
    risks = RiskAssessment.objects.select_related(
        'position', 'hazard', 'consequence', 'responsible_person', 'created_by'
    ).prefetch_related('people_at_risk')
    
    if not is_hse_mgr:
        risks = risks.filter(approval_status='approved')

    filter_form = RiskFilterForm(request.GET or None)
    if filter_form.is_valid():
        position = filter_form.cleaned_data.get('position')
        if position:
            risks = risks.filter(position=position)

        risk_level = filter_form.cleaned_data.get('risk_level')
        if risk_level:
            risks = risks.filter(risk_level=risk_level)

        risk_source = filter_form.cleaned_data.get('risk_source')
        if risk_source:
            risks = risks.filter(risk_source=risk_source)

        is_routine = filter_form.cleaned_data.get('is_routine')
        if is_routine == 'true':
            risks = risks.filter(is_routine=True)
        elif is_routine == 'false':
            risks = risks.filter(is_routine=False)

        corrective_required = filter_form.cleaned_data.get('corrective_action_required')
        if corrective_required == 'true':
            risks = risks.filter(corrective_action_required=True)
        elif corrective_required == 'false':
            risks = risks.filter(corrective_action_required=False)

    search_query = request.GET.get('q', '').strip()
    if search_query:
        risks = risks.filter(
            Q(activity_component__icontains=search_query)
            | Q(hazard__description__icontains=search_query)
            | Q(potential_event__icontains=search_query)
            | Q(causes__icontains=search_query)
            | Q(position__name__icontains=search_query)
            | Q(notes__icontains=search_query)
        )

    sort_option = request.GET.get('sort', 'critical')
    if sort_option == 'newest':
        risks = risks.order_by('-created_at')
    elif sort_option == 'oldest':
        risks = risks.order_by('created_at')
    elif sort_option == 'deadline':
        risks = risks.order_by('action_deadline', '-risk_number')
    else:
        risks = risks.order_by('-risk_number', '-created_at')

    aggregates = risks.aggregate(
        total=Count('id'),
        high=Count('id', filter=Q(risk_level='High')),
        medium=Count('id', filter=Q(risk_level='Medium')),
        low=Count('id', filter=Q(risk_level='Low')),
        corrective_open=Count('id', filter=Q(corrective_action_required=True))
    )
    stats = {key: aggregates.get(key, 0) or 0 for key in aggregates}

    today = timezone.now().date()
    lookahead = today + timedelta(days=14)
    overdue_count = risks.filter(
        corrective_action_required=True,
        action_deadline__lt=today
    ).count()
    upcoming_actions = risks.filter(
        corrective_action_required=True,
        action_deadline__isnull=False,
        action_deadline__gte=today,
        action_deadline__lte=lookahead
    ).order_by('action_deadline')[:5]

    high_risk_alerts = risks.filter(risk_level='High').order_by('-risk_number', '-created_at')[:5]
    recent_updates = risks.order_by('-updated_at')[:6]

    try:
        per_page = int(request.GET.get('per_page', 25))
    except (TypeError, ValueError):
        per_page = 25
    per_page = max(5, min(per_page, 100))

    paginator = Paginator(risks, per_page)
    page_number = request.GET.get('page')
    risks_page = paginator.get_page(page_number)

    query_params = request.GET.copy()
    query_params.pop('page', None)
    query_string = query_params.urlencode()
    per_page_options = sorted({25, 50, 75, 100, per_page})

    context = {
        'risks': risks_page,
        'filter_form': filter_form,
        'search_query': search_query,
        'sort_option': sort_option,
        'per_page': per_page,
        'per_page_options': per_page_options,
        'stats': stats,
        'overdue_count': overdue_count,
        'upcoming_actions': upcoming_actions,
        'high_risk_alerts': high_risk_alerts,
        'recent_updates': recent_updates,
        'today': today,
        'query_string': query_string,
    }

    return render(request, 'risk_assessment/risk_list.html', context)


@login_required
def position_risk_list(request, position_id):
    """
    نمایش جدول ریسک‌های یک شغل خاص (مشابه اکسل)
    """
    position = get_object_or_404(Position, id=position_id)
    
    # اگر کاربر مدیر HSE است، همه ریسک‌ها را ببین، وگرنه فقط تأیید شده‌ها
    is_hse_mgr = request.user.groups.filter(name='مدیر HSE').exists()
    
    # فیلتر کردن
    risks = RiskAssessment.objects.filter(position=position).select_related(
        'hazard', 'consequence', 'responsible_person', 'created_by'
    ).prefetch_related('people_at_risk')
    
    if not is_hse_mgr:
        risks = risks.filter(approval_status='approved')
    
    filter_form = RiskFilterForm(request.GET)
    
    if filter_form.is_valid():
        if filter_form.cleaned_data.get('risk_level'):
            risks = risks.filter(risk_level=filter_form.cleaned_data['risk_level'])
        
        if filter_form.cleaned_data.get('risk_source'):
            risks = risks.filter(risk_source=filter_form.cleaned_data['risk_source'])
        
        if filter_form.cleaned_data.get('is_routine') == 'true':
            risks = risks.filter(is_routine=True)
        elif filter_form.cleaned_data.get('is_routine') == 'false':
            risks = risks.filter(is_routine=False)
        
        if filter_form.cleaned_data.get('corrective_action_required') == 'true':
            risks = risks.filter(corrective_action_required=True)
        elif filter_form.cleaned_data.get('corrective_action_required') == 'false':
            risks = risks.filter(corrective_action_required=False)
    
    # مرتب‌سازی بر اساس عدد ریسک (بالاترین اول)
    risks = risks.order_by('-risk_number', '-created_at')

    # آمار سطح ریسک برای نمایش کارت‌های خلاصه
    total_count = risks.count()
    high_count = risks.filter(risk_level='High').count()
    medium_count = risks.filter(risk_level='Medium').count()
    low_count = risks.filter(risk_level='Low').count()
    corrective_open = risks.filter(corrective_action_required=True).count()
    reevaluated_count = risks.filter(residual_risk_number__isnull=False).count()
    
    # اقدام‌های اصلاحی با مهلت گذشته (deadline تاریخ گذشته و هنوز اقدام لازم است)
    from django.utils import timezone
    today = timezone.now().date()
    overdue_count = risks.filter(corrective_action_required=True, action_deadline__lt=today).count()

    # درصد کاهش میانگین ریسک در ارزیابی‌های مجدد (در صورت وجود داده باقی‌مانده)
    residual_with_initial = risks.filter(residual_risk_number__isnull=False)
    avg_initial = residual_with_initial.aggregate(c=Count('id'))['c'] and sum(
        r.risk_number for r in residual_with_initial
    ) / residual_with_initial.count() if residual_with_initial.count() else 0
    avg_residual = residual_with_initial.aggregate(c=Count('id'))['c'] and sum(
        r.residual_risk_number for r in residual_with_initial
    ) / residual_with_initial.count() if residual_with_initial.count() else 0
    reduction_percent = 0
    if avg_initial and avg_residual:
        reduction_percent = round(((avg_initial - avg_residual) / avg_initial) * 100, 1)

    # روند ۶ ماهه ایجاد ریسک برای این شغل بر اساس سطح ریسک
    from datetime import datetime, timedelta
    six_months_ago = (timezone.now() - timedelta(days=180))
    trend_qs = risks.filter(created_at__gte=six_months_ago)
    trend = trend_qs.annotate(m=TruncMonth('created_at')).values('m', 'risk_level').annotate(cnt=Count('id')).order_by('m')
    # ساخت ساختار داده قابل استفاده در نمودار
    trend_labels = []
    trend_high = []
    trend_medium = []
    trend_low = []
    # ایندکس‌سازی: ماه -> سطح -> تعداد
    from collections import defaultdict
    month_map = defaultdict(lambda: {'High':0,'Medium':0,'Low':0})
    for row in trend:
        month_map[row['m'].strftime('%Y-%m')] [row['risk_level']] = row['cnt']
    for month in sorted(month_map.keys()):
        trend_labels.append(month)
        trend_high.append(month_map[month]['High'])
        trend_medium.append(month_map[month]['Medium'])
        trend_low.append(month_map[month]['Low'])
    
    # Pagination
    paginator = Paginator(risks, 20)  # 20 ریسک در هر صفحه
    page_number = request.GET.get('page')
    risks_page = paginator.get_page(page_number)
    
    context = {
        'position': position,
        'risks': risks_page,
        'filter_form': filter_form,
        'total_count': total_count,
        'high_count': high_count,
        'medium_count': medium_count,
        'low_count': low_count,
        'corrective_open': corrective_open,
        'reevaluated_count': reevaluated_count,
        'overdue_count': overdue_count,
        'reduction_percent': reduction_percent,
        'trend_labels': trend_labels,
        'trend_high': trend_high,
        'trend_medium': trend_medium,
        'trend_low': trend_low,
    }
    
    return render(request, 'risk_assessment/risk_matrix.html', context)


@login_required
def risk_create_start(request):
    """
    صفحه شروع - انتخاب شغل برای ثبت ریسک
    """
    positions = Position.objects.all().order_by('name')
    
    context = {
        'positions': positions,
    }
    
    return render(request, 'risk_assessment/select_position.html', context)


@login_required
def risk_create_for_position(request, position_id):
    """
    ایجاد ریسک جدید برای یک شغل
    """
    position = get_object_or_404(Position, id=position_id)
    
    # پیشنهادات هوشمند
    risk_suggestions = get_position_risk_suggestions(position)
    
    if request.method == 'POST':
        form = RiskAssessmentForm(request.POST)
        if form.is_valid():
            risk = form.save(commit=False)
            risk.position = position
            risk.approval_status = 'pending'  # وضعیت پیش‌فرض: در انتظار تأیید
            if hasattr(request.user, 'userprofile'):
                risk.created_by = request.user.userprofile
            risk.save()
            form.save_m2m()  # ذخیره ManyToMany fields
            
            # ارسال اعلان به مدیر HSE
            notify_risk_created(risk, actor=request.user)
            
            messages.success(
                request, 
                f'ریسک با موفقیت ثبت شد. عدد ریسک: {risk.risk_number}. '
                f'ریسک در انتظار تأیید مدیر HSE است.'
            )
            return redirect('risk_assessment:position_risks', position_id=position.id)
        else:
            # لاگ خطاها برای دیباگ
            print("Form errors:", form.errors)
            print("Form data:", request.POST)
    else:
        form = RiskAssessmentForm(initial={'position': position})
    
    context = {
        'form': form,
        'position': position,
        'action': 'create',
        'risk_suggestions': risk_suggestions,
    }
    
    return render(request, 'risk_assessment/risk_form.html', context)


@login_required
def risk_create(request, position_id):
    """
    [DEPRECATED] استفاده از risk_create_for_position
    """
    return risk_create_for_position(request, position_id)


@login_required
def risk_update(request, risk_id):
    """
    ویرایش ریسک موجود
    """
    risk = get_object_or_404(RiskAssessment, id=risk_id)
    
    if request.method == 'POST':
        form = RiskAssessmentForm(request.POST, instance=risk)
        if form.is_valid():
            form.save()
            messages.success(request, 'ریسک با موفقیت بروزرسانی شد.')
            return redirect('risk_assessment:position_risks', position_id=risk.position.id)
    else:
        form = RiskAssessmentForm(instance=risk)
    
    context = {
        'form': form,
        'risk': risk,
        'position': risk.position,
        'action': 'update',
    }
    
    return render(request, 'risk_assessment/risk_form.html', context)


@login_required
def risk_detail(request, risk_id):
    """
    نمایش جزئیات کامل یک ریسک
    """
    risk = get_object_or_404(
        RiskAssessment.objects.select_related(
            'position', 'hazard', 'consequence', 'responsible_person', 'created_by'
        ).prefetch_related('people_at_risk'),
        id=risk_id
    )
    
    # تولید بینش‌های هوشمند
    insights = generate_risk_insights(risk)
    
    # محاسبه تفاوت ریسک
    risk_difference = None
    if risk.residual_risk_number:
        risk_difference = risk.residual_risk_number - risk.risk_number
    
    # دریافت اقدامات اصلاحی مرتبط با این ریسک
    related_corrective_actions = []
    try:
        from corrective_actions.models import CorrectiveAction
        related_corrective_actions = CorrectiveAction.objects.filter(related_risk=risk).select_related('requester', 'receiver').order_by('-created_at')
    except ImportError:
        pass
    
    context = {
        'risk': risk,
        'insights': insights,
        'related_anomalies': insights['related_anomalies'],
        'related_incidents': insights['related_incidents'],
        'trend_data': insights['anomaly_trend'],
        'suggestions': insights['suggestions'],
        'risk_difference': risk_difference,
        'related_corrective_actions': related_corrective_actions,
    }
    
    return render(request, 'risk_assessment/risk_detail.html', context)


@login_required
def risk_re_evaluate(request, risk_id):
    """
    ارزیابی مجدد ریسک (پس از اقدامات اصلاحی)
    """
    risk = get_object_or_404(RiskAssessment, id=risk_id)
    
    if request.method == 'POST':
        form = RiskReEvaluationForm(request.POST, instance=risk)
        if form.is_valid():
            form.save()
            messages.success(
                request, 
                f'ارزیابی مجدد انجام شد. ریسک باقی‌مانده: {risk.residual_risk_number}'
            )
            return redirect('risk_assessment:risk_detail', risk_id=risk.id)
    else:
        form = RiskReEvaluationForm(instance=risk)
    
    context = {
        'form': form,
        'risk': risk,
    }
    
    return render(request, 'risk_assessment/risk_re_evaluate.html', context)


@login_required
def risk_delete(request, risk_id):
    """
    حذف ریسک
    """
    risk = get_object_or_404(RiskAssessment, id=risk_id)
    position_id = risk.position.id
    
    if request.method == 'POST':
        risk.delete()
        messages.success(request, 'ریسک با موفقیت حذف شد.')
        return redirect('risk_assessment:position_risks', position_id=position_id)
    
    context = {
        'risk': risk,
    }
    
    return render(request, 'risk_assessment/risk_confirm_delete.html', context)


@login_required
def risk_matrix_view(request):
    """
    نمایش ماتریس ۵×۵ ریسک (برای تمام ریسک‌ها)
    """
    # ساخت ساختار سطرها برای استفاده ساده در قالب
    rows = []
    for severity in range(1, 6):
        cells = []
        for probability in range(1, 6):
            rpn = severity * probability
            count = RiskAssessment.objects.filter(
                severity=severity,
                probability=probability,
                approval_status='approved'  # فقط ریسک‌های تأیید شده
            ).count()
            # منطق رنگ‌بندی بر اساس ماتریس JHA
            # قرمز: RPN >= 15 (15, 16, 20, 25)
            # زرد: RPN در [5, 6, 8, 9, 10, 12] (شامل دو سلول با RPN=5 و دو سلول با RPN=6)
            # آبی/سبز: RPN در [1, 2, 3, 4]
            if rpn >= 15:
                color = 'danger'
            elif rpn in [5, 6, 8, 9, 10, 12]:
                color = 'warning'
            else:
                color = 'success'
            cells.append({
                'severity': severity,
                'probability': probability,
                'rpn': rpn,
                'count': count,
                'color': color,
            })
        rows.append({'severity': severity, 'cells': cells})

    context = {
        'rows': rows,
        'severities': range(1, 6),
        'probabilities': range(1, 6),
    }
    return render(request, 'risk_assessment/matrix_view.html', context)


def is_hse_manager(user):
    """بررسی اینکه آیا کاربر عضو گروه مدیر HSE است"""
    return user.groups.filter(name='مدیر HSE').exists()


@login_required
@user_passes_test(is_hse_manager)
def risk_approve(request, risk_id):
    """
    تأیید ریسک توسط مدیر HSE
    """
    risk = get_object_or_404(RiskAssessment, id=risk_id)
    
    if risk.approval_status == 'approved':
        messages.warning(request, 'این ریسک قبلاً تأیید شده است.')
        return redirect('risk_assessment:risk_detail', risk_id=risk.id)
    
    if request.method == 'POST':
        risk.approval_status = 'approved'
        if hasattr(request.user, 'userprofile'):
            risk.approved_by = request.user.userprofile
        from django.utils import timezone
        risk.approved_at = timezone.now()
        risk.rejected_by = None
        risk.rejected_at = None
        risk.rejection_reason = ''
        risk.save()
        
        # ارسال اعلان
        notify_risk_approved(risk, actor=request.user)
        
        messages.success(request, f'ریسک با عدد ریسک {risk.risk_number} با موفقیت تأیید شد.')
        return redirect('risk_assessment:risk_detail', risk_id=risk.id)
    
    context = {
        'risk': risk,
        'action': 'approve',
    }
    return render(request, 'risk_assessment/risk_approval.html', context)


@login_required
@user_passes_test(is_hse_manager)
def risk_reject(request, risk_id):
    """
    رد ریسک توسط مدیر HSE
    """
    risk = get_object_or_404(RiskAssessment, id=risk_id)
    
    if risk.approval_status == 'rejected':
        messages.warning(request, 'این ریسک قبلاً رد شده است.')
        return redirect('risk_assessment:risk_detail', risk_id=risk.id)
    
    if request.method == 'POST':
        rejection_reason = request.POST.get('rejection_reason', '').strip()
        if not rejection_reason:
            messages.error(request, 'لطفاً دلیل رد را وارد کنید.')
            context = {
                'risk': risk,
                'action': 'reject',
            }
            return render(request, 'risk_assessment/risk_approval.html', context)
        
        risk.approval_status = 'rejected'
        if hasattr(request.user, 'userprofile'):
            risk.rejected_by = request.user.userprofile
        from django.utils import timezone
        risk.rejected_at = timezone.now()
        risk.rejection_reason = rejection_reason
        risk.approved_by = None
        risk.approved_at = None
        risk.save()
        
        # ارسال اعلان
        notify_risk_rejected(risk, actor=request.user)
        
        messages.success(request, f'ریسک با عدد ریسک {risk.risk_number} رد شد.')
        return redirect('risk_assessment:risk_detail', risk_id=risk.id)
    
    context = {
        'risk': risk,
        'action': 'reject',
    }
    return render(request, 'risk_assessment/risk_approval.html', context)


@login_required
@user_passes_test(is_hse_manager)
def risk_pending_list(request):
    """
    لیست ریسک‌های در انتظار تأیید (فقط برای مدیر HSE)
    """
    risks = RiskAssessment.objects.filter(
        approval_status='pending'
    ).select_related(
        'position', 'hazard', 'consequence', 'responsible_person', 'created_by'
    ).prefetch_related('people_at_risk').order_by('-created_at')
    
    # Pagination
    paginator = Paginator(risks, 20)
    page_number = request.GET.get('page')
    risks_page = paginator.get_page(page_number)
    
    context = {
        'risks': risks_page,
        'pending_count': risks.count(),
    }
    return render(request, 'risk_assessment/risk_pending_list.html', context)


@login_required
def risk_export_excel(request):
    """Export تمام ریسک‌ها به Excel"""
    # فیلتر بر اساس دسترسی کاربر
    is_hse_mgr = request.user.groups.filter(name='مدیر HSE').exists()
    risks = RiskAssessment.objects.select_related(
        'position', 'hazard', 'consequence', 'responsible_person', 'created_by'
    ).prefetch_related('people_at_risk')
    
    if not is_hse_mgr:
        risks = risks.filter(approval_status='approved')
    
    # اعمال فیلترهای موجود در query string
    filter_form = RiskFilterForm(request.GET or None)
    if filter_form.is_valid():
        position = filter_form.cleaned_data.get('position')
        if position:
            risks = risks.filter(position=position)
        
        risk_level = filter_form.cleaned_data.get('risk_level')
        if risk_level:
            risks = risks.filter(risk_level=risk_level)
        
        risk_source = filter_form.cleaned_data.get('risk_source')
        if risk_source:
            risks = risks.filter(risk_source=risk_source)
        
        is_routine = filter_form.cleaned_data.get('is_routine')
        if is_routine == 'true':
            risks = risks.filter(is_routine=True)
        elif is_routine == 'false':
            risks = risks.filter(is_routine=False)
        
        corrective_required = filter_form.cleaned_data.get('corrective_action_required')
        if corrective_required == 'true':
            risks = risks.filter(corrective_action_required=True)
        elif corrective_required == 'false':
            risks = risks.filter(corrective_action_required=False)
    
    search_query = request.GET.get('q', '').strip()
    if search_query:
        risks = risks.filter(
            Q(activity_component__icontains=search_query)
            | Q(hazard__description__icontains=search_query)
            | Q(potential_event__icontains=search_query)
            | Q(causes__icontains=search_query)
            | Q(position__name__icontains=search_query)
            | Q(notes__icontains=search_query)
        )
    
    risks = risks.order_by('-risk_number', '-created_at')
    
    # ایجاد فایل Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Assessments"
    
    # استایل‌ها
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # هدرها
    headers = [
        'شغل/سمت',
        'منشا شناسایی',
        'منشا دیگر',
        'اجزا فعالیت',
        'روتین',
        'خطر',
        'افراد در معرض خطر',
        'رویداد محتمل',
        'علل وقوع',
        'نوع پیامد',
        'کنترل‌های موجود',
        'علل شکست کنترل',
        'الزام قانونی',
        'شرح الزام',
        'رعایت شده',
        'احتمال (P)',
        'شدت (S)',
        'عدد ریسک (RPN)',
        'سطح ریسک',
        'حذف خطر',
        'جایگزینی',
        'کنترل مهندسی',
        'کنترل اداری',
        'PPE',
        'نیاز به اقدام',
        'شماره اقدام',
        'تاریخ اقدام',
        'مهلت اقدام',
        'مسئول',
        'MUE',
        'کد MUE',
        'شرایط اضطراری',
        'کد اضطراری',
        'تاریخ ارزیابی مجدد',
        'احتمال باقی‌مانده',
        'شدت باقی‌مانده',
        'ریسک باقی‌مانده',
        'سطح ریسک باقی‌مانده',
        'وضعیت تأیید',
        'یادداشت‌ها',
        'ایجاد شده توسط',
        'تاریخ ایجاد',
        'آخرین بروزرسانی',
    ]
    
    ws.append(headers)
    
    # اعمال استایل به هدر
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 20
    
    # داده‌ها
    for risk in risks:
        people_at_risk_names = ', '.join([p.name for p in risk.people_at_risk.all()])
        
        row_data = [
            risk.position.name if risk.position else '',
            risk.get_risk_source_display() if risk.risk_source else '',
            risk.risk_source_other or '',
            risk.activity_component or '',
            'بله' if risk.is_routine else 'خیر',
            risk.hazard.description if risk.hazard else '',
            people_at_risk_names,
            risk.potential_event or '',
            risk.causes or '',
            risk.consequence.name if risk.consequence else '',
            risk.existing_controls or '',
            risk.control_failure_causes or '',
            'بله' if risk.has_legal_requirement else 'خیر',
            risk.legal_requirement_desc or '',
            'بله' if risk.is_legal_compliant else ('خیر' if risk.is_legal_compliant is False else ''),
            risk.probability or '',
            risk.severity or '',
            risk.risk_number or '',
            risk.get_risk_level_display() if risk.risk_level else '',
            risk.control_elimination or '',
            risk.control_substitution or '',
            risk.control_engineering or '',
            risk.control_admin or '',
            risk.control_ppe or '',
            'بله' if risk.corrective_action_required else 'خیر',
            risk.action_number or '',
            jdatetime.date.fromgregorian(date=risk.action_date).strftime('%Y/%m/%d') if risk.action_date else '',
            jdatetime.date.fromgregorian(date=risk.action_deadline).strftime('%Y/%m/%d') if risk.action_deadline else '',
            risk.responsible_person.user.get_full_name() if risk.responsible_person else '',
            'بله' if risk.is_mue else 'خیر',
            risk.mue_code or '',
            'بله' if risk.is_emergency else 'خیر',
            risk.emergency_code or '',
            jdatetime.date.fromgregorian(date=risk.re_evaluation_date).strftime('%Y/%m/%d') if risk.re_evaluation_date else '',
            risk.residual_probability or '',
            risk.residual_severity or '',
            risk.residual_risk_number or '',
            risk.get_residual_risk_level_display() if risk.residual_risk_level else '',
            risk.get_approval_status_display() if risk.approval_status else '',
            risk.notes or '',
            risk.created_by.user.get_full_name() if risk.created_by else '',
            jdatetime.datetime.fromgregorian(datetime=risk.created_at).strftime('%Y/%m/%d %H:%M') if risk.created_at else '',
            jdatetime.datetime.fromgregorian(datetime=risk.updated_at).strftime('%Y/%m/%d %H:%M') if risk.updated_at else '',
        ]
        
        ws.append(row_data)
        
        # اعمال استایل به ردیف
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=col_num)
            cell.border = border
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    # پاسخ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'risk_assessments_export_{jdatetime.datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename={filename}'
    wb.save(response)
    
    return response


@login_required
def risk_download_template(request):
    """دانلود الگوی Excel برای import ریسک‌ها"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Assessments"
    
    # استایل‌ها
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # هدرها
    headers = [
        'شغل/سمت*',
        'منشا شناسایی*',
        'منشا دیگر',
        'اجزا فعالیت*',
        'روتین (بله/خیر)*',
        'خطر (نام دقیق)*',
        'افراد در معرض خطر (با کاما جدا کنید)',
        'رویداد محتمل*',
        'علل وقوع*',
        'نوع پیامد (نام دقیق)*',
        'کنترل‌های موجود*',
        'علل شکست کنترل*',
        'الزام قانونی (بله/خیر)',
        'شرح الزام',
        'رعایت شده (بله/خیر)',
        'احتمال (P) (1-5)*',
        'شدت (S) (1-5)*',
        'حذف خطر',
        'جایگزینی',
        'کنترل مهندسی',
        'کنترل اداری',
        'PPE',
        'نیاز به اقدام (بله/خیر)',
        'شماره اقدام',
        'تاریخ اقدام (YYYY/MM/DD)',
        'مهلت اقدام (YYYY/MM/DD)',
        'مسئول (نام کاربری)',
        'MUE (بله/خیر)',
        'کد MUE',
        'شرایط اضطراری (بله/خیر)',
        'کد اضطراری',
        'یادداشت‌ها',
    ]
    
    ws.append(headers)
    
    # اعمال استایل به هدر
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 25
    
    # داده نمونه
    sample_data = [
        [
            'راهبر لودر',
            'site_visit',
            '',
            'راهبری لودر در معدن',
            'بله',
            'سقوط از ارتفاع',
            'راهبر لودر,کارگر معدن',
            'سقوط از لودر',
            'عدم استفاده از کمربند ایمنی',
            'آسیب شدید',
            'تابلو هشدار',
            'عدم توجه به تابلو',
            'خیر',
            '',
            '',
            '3',
            '4',
            '',
            '',
            '',
            '',
            '',
            'بله',
            'ACT-001',
            '1403/10/15',
            '1403/11/15',
            '',
            'خیر',
            '',
            'خیر',
            '',
            'یادداشت نمونه',
        ],
    ]
    
    for row_data in sample_data:
        ws.append(row_data)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=col_num)
            cell.border = border
            cell.alignment = Alignment(horizontal="right", vertical="center")
    
    # برگه راهنما
    ws2 = wb.create_sheet("راهنما")
    instructions = [
        ["راهنمای استفاده از فایل Excel برای Import ریسک‌ها"],
        [""],
        ["فیلدهای دارای علامت * الزامی هستند"],
        [""],
        ["شغل/سمت: باید دقیقاً با نام یکی از مشاغل موجود در سیستم مطابقت داشته باشد"],
        [""],
        ["منشا شناسایی: یکی از موارد زیر"],
        ["  - site_visit: بازدیدهای میدانی و آنومالی ریپورت"],
        ["  - consultation: مشارکت و مشاوره کارکنان"],
        ["  - audit: نتایج ممیزی‌ها"],
        ["  - incidents: حوادث و شبه حوادث"],
        ["  - changes: تغییرات"],
        ["  - physical: عوامل فیزیکی"],
        ["  - chemical: عوامل شیمیایی"],
        ["  - biological: عوامل محیط زیستی و بیولوژیکی"],
        ["  - ergonomic: عوامل ارگونومیک و انسانی"],
        ["  - activities: فعالیت‌ها و عملیات کاری"],
        ["  - energy: انرژی‌های خطرناک"],
        ["  - environmental: وضعیت‌های محیطی"],
        ["  - other: سایر"],
        [""],
        ["خطر: باید دقیقاً با نام یکی از خطرات موجود در سیستم مطابقت داشته باشد"],
        [""],
        ["نوع پیامد: باید دقیقاً با نام یکی از انواع آسیب موجود در سیستم مطابقت داشته باشد"],
        [""],
        ["احتمال (P): عدد بین 1 تا 5"],
        ["  1: خیلی کم (کمتر از یک بار در ۵ سال)"],
        ["  2: کم (۵ سال یکبار)"],
        ["  3: متوسط (وقوع سالیانه)"],
        ["  4: زیاد (چند بار در سال)"],
        ["  5: خیلی زیاد (چند بار در ماه)"],
        [""],
        ["شدت (S): عدد بین 1 تا 5"],
        ["  1: آسیب جزئی (بدون استراحت)"],
        ["  2: آسیب شدید (1 تا 14 روز استراحت)"],
        ["  3: آسیب خیلی شدید (قطع عضو/بیش از 15 روز)"],
        ["  4: ناتوانی و از کار افتادگی دائم"],
        ["  5: مرگ"],
        [""],
        ["تاریخ‌ها: به فرمت جلالی YYYY/MM/DD مانند 1403/10/25"],
        [""],
        ["روتین، الزام قانونی، رعایت شده، نیاز به اقدام، MUE، شرایط اضطراری: بله یا خیر"],
        [""],
        ["افراد در معرض خطر: نام مشاغل را با کاما جدا کنید"],
        [""],
        ["مسئول: نام کاربری (username) شخص مسئول را وارد کنید"],
    ]
    
    for row in instructions:
        ws2.append(row)
        if row and row[0] and ("راهنما" in row[0] or "فیلدهای" in row[0]):
            cell = ws2.cell(row=ws2.max_row, column=1)
            cell.font = Font(bold=True, size=14)
            cell.fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    
    ws2.column_dimensions['A'].width = 80
    
    # پاسخ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=risk_assessment_template.xlsx'
    wb.save(response)
    
    return response


@login_required
def risk_import_excel(request):
    """صفحه import ریسک‌ها از Excel"""
    form = ExcelImportForm()
    
    if request.method == 'POST':
        form = ExcelImportForm(request.POST, request.FILES)
        if form.is_valid():
            excel_file = request.FILES['excel_file']
            
            try:
                wb = load_workbook(excel_file, data_only=True)
                ws = wb.active
                
                errors = []
                success_count = 0
                row_num = 2  # شروع از ردیف 2 (ردیف 1 هدر است)
                
                # داده‌های lookup
                positions = {p.name: p for p in Position.objects.all()}
                hazards = {h.description: h for h in AnomalyDescription.objects.all()}
                consequences = {c.name: c for c in InjuryType.objects.all()}
                users = {u.user.username: u for u in UserProfile.objects.all() if u.user}
                
                # نقشه منشاها
                source_map = {
                    'site_visit': 'site_visit',
                    'بازدیدهای میدانی': 'site_visit',
                    'consultation': 'consultation',
                    'مشارکت و مشاوره': 'consultation',
                    'audit': 'audit',
                    'ممیزی': 'audit',
                    'incidents': 'incidents',
                    'حوادث': 'incidents',
                    'changes': 'changes',
                    'تغییرات': 'changes',
                    'physical': 'physical',
                    'عوامل فیزیکی': 'physical',
                    'chemical': 'chemical',
                    'عوامل شیمیایی': 'chemical',
                    'biological': 'biological',
                    'عوامل بیولوژیکی': 'biological',
                    'ergonomic': 'ergonomic',
                    'عوامل ارگونومیک': 'ergonomic',
                    'activities': 'activities',
                    'فعالیت‌ها': 'activities',
                    'energy': 'energy',
                    'انرژی': 'energy',
                    'environmental': 'environmental',
                    'محیطی': 'environmental',
                    'other': 'other',
                    'سایر': 'other',
                }
                
                for row in ws.iter_rows(min_row=2, values_only=True):
                    # رد کردن ردیف‌های خالی
                    if not any(row):
                        continue
                    
                    try:
                        # خواندن داده‌ها
                        position_name = str(row[0]).strip() if row[0] else None
                        risk_source_str = str(row[1]).strip() if row[1] else None
                        risk_source_other = str(row[2]).strip() if row[2] else ''
                        activity_component = str(row[3]).strip() if row[3] else None
                        is_routine_str = str(row[4]).strip().lower() if row[4] else None
                        hazard_name = str(row[5]).strip() if row[5] else None
                        people_at_risk_str = str(row[6]).strip() if row[6] else ''
                        potential_event = str(row[7]).strip() if row[7] else None
                        causes = str(row[8]).strip() if row[8] else None
                        consequence_name = str(row[9]).strip() if row[9] else None
                        existing_controls = str(row[10]).strip() if row[10] else None
                        control_failure_causes = str(row[11]).strip() if row[11] else None
                        has_legal_str = str(row[12]).strip().lower() if row[12] else 'خیر'
                        legal_requirement_desc = str(row[13]).strip() if row[13] else ''
                        is_legal_compliant_str = str(row[14]).strip().lower() if row[14] else ''
                        probability = int(row[15]) if row[15] else None
                        severity = int(row[16]) if row[16] else None
                        control_elimination = str(row[17]).strip() if row[17] else ''
                        control_substitution = str(row[18]).strip() if row[18] else ''
                        control_engineering = str(row[19]).strip() if row[19] else ''
                        control_admin = str(row[20]).strip() if row[20] else ''
                        control_ppe = str(row[21]).strip() if row[21] else ''
                        corrective_required_str = str(row[22]).strip().lower() if row[22] else 'خیر'
                        action_number = str(row[23]).strip() if row[23] else ''
                        action_date_str = str(row[24]).strip() if row[24] else ''
                        action_deadline_str = str(row[25]).strip() if row[25] else ''
                        responsible_username = str(row[26]).strip() if row[26] else ''
                        is_mue_str = str(row[27]).strip().lower() if row[27] else 'خیر'
                        mue_code = str(row[28]).strip() if row[28] else ''
                        is_emergency_str = str(row[29]).strip().lower() if row[29] else 'خیر'
                        emergency_code = str(row[30]).strip() if row[30] else ''
                        notes = str(row[31]).strip() if row[31] else ''
                        
                        # اعتبارسنجی فیلدهای الزامی
                        if not all([position_name, risk_source_str, activity_component, 
                                   hazard_name, potential_event, causes, consequence_name,
                                   existing_controls, control_failure_causes, probability, severity]):
                            errors.append(f"ردیف {row_num}: فیلدهای الزامی خالی است")
                            row_num += 1
                            continue
                        
                        # تبدیل منشا
                        risk_source = source_map.get(risk_source_str, risk_source_str)
                        if risk_source not in dict(RiskAssessment.SOURCE_CHOICES):
                            errors.append(f"ردیف {row_num}: منشا شناسایی نامعتبر: {risk_source_str}")
                            row_num += 1
                            continue
                        
                        # بررسی position
                        if position_name not in positions:
                            errors.append(f"ردیف {row_num}: شغل '{position_name}' یافت نشد")
                            row_num += 1
                            continue
                        position = positions[position_name]
                        
                        # بررسی hazard
                        if hazard_name not in hazards:
                            errors.append(f"ردیف {row_num}: خطر '{hazard_name}' یافت نشد")
                            row_num += 1
                            continue
                        hazard = hazards[hazard_name]
                        
                        # بررسی consequence
                        if consequence_name not in consequences:
                            errors.append(f"ردیف {row_num}: نوع پیامد '{consequence_name}' یافت نشد")
                            row_num += 1
                            continue
                        consequence = consequences[consequence_name]
                        
                        # تبدیل boolean
                        is_routine = is_routine_str in ['بله', 'yes', 'true', '1', 'y']
                        has_legal = has_legal_str in ['بله', 'yes', 'true', '1', 'y']
                        is_legal_compliant = None
                        if is_legal_compliant_str in ['بله', 'yes', 'true', '1', 'y']:
                            is_legal_compliant = True
                        elif is_legal_compliant_str in ['خیر', 'no', 'false', '0', 'n']:
                            is_legal_compliant = False
                        
                        corrective_required = corrective_required_str in ['بله', 'yes', 'true', '1', 'y']
                        is_mue = is_mue_str in ['بله', 'yes', 'true', '1', 'y']
                        is_emergency = is_emergency_str in ['بله', 'yes', 'true', '1', 'y']
                        
                        # بررسی probability و severity
                        if probability not in [1, 2, 3, 4, 5]:
                            errors.append(f"ردیف {row_num}: احتمال باید بین 1 تا 5 باشد")
                            row_num += 1
                            continue
                        
                        if severity not in [1, 2, 3, 4, 5]:
                            errors.append(f"ردیف {row_num}: شدت باید بین 1 تا 5 باشد")
                            row_num += 1
                            continue
                        
                        # تبدیل تاریخ‌ها
                        action_date = None
                        action_deadline = None
                        if action_date_str:
                            try:
                                parts = action_date_str.split('/')
                                if len(parts) == 3:
                                    jdate = jdatetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
                                    action_date = jdate.togregorian()
                            except:
                                pass
                        
                        if action_deadline_str:
                            try:
                                parts = action_deadline_str.split('/')
                                if len(parts) == 3:
                                    jdate = jdatetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
                                    action_deadline = jdate.togregorian()
                            except:
                                pass
                        
                        # بررسی مسئول
                        responsible_person = None
                        if responsible_username:
                            if responsible_username in users:
                                responsible_person = users[responsible_username]
                            else:
                                errors.append(f"ردیف {row_num}: کاربر '{responsible_username}' یافت نشد")
                                row_num += 1
                                continue
                        
                        # بررسی افراد در معرض خطر
                        people_at_risk_list = []
                        if people_at_risk_str:
                            people_names = [p.strip() for p in people_at_risk_str.split(',')]
                            for p_name in people_names:
                                if p_name in positions:
                                    people_at_risk_list.append(positions[p_name])
                                else:
                                    errors.append(f"ردیف {row_num}: شغل '{p_name}' در افراد در معرض خطر یافت نشد")
                        
                        # ایجاد ریسک
                        risk = RiskAssessment(
                            position=position,
                            risk_source=risk_source,
                            risk_source_other=risk_source_other if risk_source == 'other' else '',
                            activity_component=activity_component,
                            is_routine=is_routine,
                            hazard=hazard,
                            potential_event=potential_event,
                            causes=causes,
                            consequence=consequence,
                            existing_controls=existing_controls,
                            control_failure_causes=control_failure_causes,
                            has_legal_requirement=has_legal,
                            legal_requirement_desc=legal_requirement_desc if has_legal else '',
                            is_legal_compliant=is_legal_compliant,
                            probability=probability,
                            severity=severity,
                            control_elimination=control_elimination,
                            control_substitution=control_substitution,
                            control_engineering=control_engineering,
                            control_admin=control_admin,
                            control_ppe=control_ppe,
                            corrective_action_required=corrective_required,
                            action_number=action_number if corrective_required else '',
                            action_date=action_date if corrective_required else None,
                            action_deadline=action_deadline if corrective_required else None,
                            responsible_person=responsible_person if corrective_required else None,
                            is_mue=is_mue,
                            mue_code=mue_code if is_mue else '',
                            is_emergency=is_emergency,
                            emergency_code=emergency_code if is_emergency else '',
                            notes=notes,
                            approval_status='pending',
                        )
                        
                        if hasattr(request.user, 'userprofile'):
                            risk.created_by = request.user.userprofile
                        
                        risk.save()
                        risk.people_at_risk.set(people_at_risk_list)
                        
                        success_count += 1
                        row_num += 1
                        
                    except Exception as e:
                        errors.append(f"ردیف {row_num}: خطا - {str(e)}")
                        row_num += 1
                        continue
                
                if errors:
                    messages.warning(request, f'{success_count} ریسک با موفقیت وارد شد. {len(errors)} خطا رخ داد.')
                    for error in errors[:10]:  # نمایش 10 خطای اول
                        messages.error(request, error)
                else:
                    messages.success(request, f'{success_count} ریسک با موفقیت وارد شد.')
                
                return redirect('risk_assessment:risk_list')
                
            except Exception as e:
                messages.error(request, f'خطا در خواندن فایل Excel: {str(e)}')
    
    context = {
        'form': form,
    }
    return render(request, 'risk_assessment/risk_import.html', context)
